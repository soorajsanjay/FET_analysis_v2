"""Modular readers that expose pipeline artifacts as dashboard datasets."""
from __future__ import annotations

import csv
import json
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Protocol

from fet_analyzer.path_utils import safe_path


@dataclass(frozen=True)
class DatasetRef:
    id: str
    report: str
    name: str
    path: Path
    kind: str
    sheet: str | None = None


class ReportAdapter(Protocol):
    suffixes: tuple[str, ...]

    def inspect(self, path: Path, root: Path) -> list[DatasetRef]: ...
    def read(self, ref: DatasetRef, limit: int | None = None) -> dict[str, Any]: ...


def _clean(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _table(rows: list[list[Any]], limit: int | None = None) -> dict[str, Any]:
    if not rows:
        return {"columns": [], "rows": [], "row_count": 0}
    width = max(len(row) for row in rows)
    first = list(rows[0]) + [None] * (width - len(rows[0]))
    columns = [str(v).strip() if v not in (None, "") else f"column_{i + 1}" for i, v in enumerate(first)]
    seen: dict[str, int] = {}
    for i, name in enumerate(columns):
        seen[name] = seen.get(name, 0) + 1
        if seen[name] > 1:
            columns[i] = f"{name}_{seen[name]}"
    body = rows[1:]
    total = len(body)
    if limit is not None:
        body = body[:limit]
    return {
        "columns": columns,
        "rows": [dict(zip(columns, [_clean(v) for v in row] + [None] * (width - len(row)))) for row in body],
        "row_count": total,
    }


class CsvAdapter:
    suffixes = (".csv", ".tsv")

    def inspect(self, path: Path, root: Path) -> list[DatasetRef]:
        rel = path.relative_to(root).as_posix()
        return [DatasetRef(f"file:{rel}", rel, path.stem, path, "table")]

    def read(self, ref: DatasetRef, limit: int | None = None) -> dict[str, Any]:
        delimiter = "\t" if ref.path.suffix.lower() == ".tsv" else ","
        with safe_path(ref.path).open(newline="", encoding="utf-8-sig", errors="replace") as handle:
            return _table([list(row) for row in csv.reader(handle, delimiter=delimiter)], limit)


class ExcelAdapter:
    suffixes = (".xlsx", ".xlsm")

    def inspect(self, path: Path, root: Path) -> list[DatasetRef]:
        from openpyxl import load_workbook
        rel = path.relative_to(root).as_posix()
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            return [DatasetRef(f"xlsx:{rel}:{sheet}", rel, sheet, path, "table", sheet) for sheet in workbook.sheetnames]
        finally:
            workbook.close()

    def read(self, ref: DatasetRef, limit: int | None = None) -> dict[str, Any]:
        from openpyxl import load_workbook
        workbook = load_workbook(ref.path, read_only=True, data_only=True)
        try:
            sheet = workbook[ref.sheet or workbook.sheetnames[0]]
            max_rows = None if limit is None else limit + 1
            rows = [list(row) for row in sheet.iter_rows(values_only=True, max_row=max_rows)]
            result = _table(rows, limit)
            result["row_count"] = max(sheet.max_row - 1, 0)
            return result
        finally:
            workbook.close()


class JsonAdapter:
    suffixes = (".json",)

    def inspect(self, path: Path, root: Path) -> list[DatasetRef]:
        rel = path.relative_to(root).as_posix()
        return [DatasetRef(f"json:{rel}", rel, path.stem, path, "json")]

    def read(self, ref: DatasetRef, limit: int | None = None) -> dict[str, Any]:
        payload = json.loads(safe_path(ref.path).read_text(encoding="utf-8"))
        rows: list[dict[str, Any]] = []

        def flatten(value: Any, prefix: str = "") -> None:
            if isinstance(value, dict):
                for key, child in value.items():
                    flatten(child, f"{prefix}.{key}" if prefix else str(key))
            elif isinstance(value, list):
                if value and all(isinstance(item, dict) for item in value):
                    for index, item in enumerate(value):
                        flatten(item, f"{prefix}[{index}]")
                else:
                    rows.append({"field": prefix, "value": json.dumps(value, default=str)})
            else:
                rows.append({"field": prefix, "value": _clean(value)})

        flatten(payload)
        total = len(rows)
        return {"columns": ["field", "value"], "rows": rows[:limit] if limit else rows, "row_count": total}


class ReportCatalog:
    """Registry-based artifact catalog reusable by other pipelines."""

    def __init__(self, adapters: list[ReportAdapter] | None = None):
        self.adapters = adapters or [ExcelAdapter(), CsvAdapter(), JsonAdapter()]
        self.last_errors: list[dict[str, str]] = []

    def scan(self, root: Path) -> list[DatasetRef]:
        refs: list[DatasetRef] = []
        self.last_errors = []
        ignored = {"decompressed_ztr", "__pycache__"}
        for path in sorted(root.rglob("*")):
            if not path.is_file() or any(part in ignored for part in path.parts):
                continue
            adapter = next((item for item in self.adapters if path.suffix.lower() in item.suffixes), None)
            if adapter:
                try:
                    refs.extend(adapter.inspect(path, root))
                except Exception as exc:
                    self.last_errors.append({
                        "path": str(path),
                        "error": f"{type(exc).__name__}: {exc}",
                    })
                    continue
        return refs

    def read(self, ref: DatasetRef, limit: int | None = None) -> dict[str, Any]:
        adapter = next(item for item in self.adapters if ref.path.suffix.lower() in item.suffixes)
        return adapter.read(ref, limit)
