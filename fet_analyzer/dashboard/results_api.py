"""Read-only adapters for the interactive Results dashboard.

This module never runs extraction or fitting code.  It presents canonical CSV
and workbook artifacts that have already been written by the analysis pipeline.
"""
from __future__ import annotations

import base64
import csv
import math
import re
from pathlib import Path
from typing import Any

from fet_analyzer.path_utils import safe_path


def _value(value: Any) -> Any:
    if value in (None, ""):
        return None
    if isinstance(value, (int, float, bool)):
        return value
    try:
        number = float(value)
        return number if math.isfinite(number) else None
    except (TypeError, ValueError):
        return str(value)


def _csv_rows(path: Path) -> list[dict[str, Any]]:
    if not path.is_file():
        return []
    with safe_path(path).open(newline="", encoding="utf-8-sig", errors="replace") as handle:
        return [{key: _value(value) for key, value in row.items()} for row in csv.DictReader(handle)]


def _sheet_rows(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    if not path.is_file():
        return []
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            return []
        values = workbook[sheet_name].iter_rows(values_only=True)
        headers = [str(item or f"column_{index + 1}") for index, item in enumerate(next(values, ()))]
        return [
            {header: _value(value) for header, value in zip(headers, row)}
            for row in values
        ]
    finally:
        workbook.close()


def artifact_id(output_root: Path, path: Path | None) -> str | None:
    if path is None or not path.is_file():
        return None
    try:
        relative = path.resolve().relative_to(output_root.resolve()).as_posix()
    except ValueError:
        return None
    return base64.urlsafe_b64encode(relative.encode("utf-8")).decode("ascii").rstrip("=")


def resolve_artifact(output_root: Path, identifier: str) -> Path:
    try:
        padded = identifier + "=" * (-len(identifier) % 4)
        relative = base64.urlsafe_b64decode(padded.encode("ascii")).decode("utf-8")
    except Exception as exc:
        raise ValueError("Invalid report identifier") from exc
    root = output_root.resolve()
    candidate = (root / relative).resolve()
    if candidate != root and root not in candidate.parents:
        raise ValueError("Report path is outside the selected output directory")
    if not candidate.is_file() or candidate.suffix.lower() not in {".html", ".xlsx", ".xlsm"}:
        raise ValueError("Report artifact is unavailable")
    return candidate


def _normal_name(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", Path(str(value or "")).stem.casefold())


def _device_artifacts(output_root: Path, row: dict[str, Any]) -> tuple[str | None, str | None]:
    workbook_value = row.get("source_workbook")
    source_stem = Path(str(row.get("source_file") or "")).stem
    candidates = [
        Path(str(workbook_value)) if workbook_value else None,
        output_root / f"{source_stem}.xlsx" if source_stem else None,
        output_root / f"{row.get('device', '')}.xlsx",
    ]
    workbook = next((item for item in candidates if item and item.is_file()), candidates[-1])
    report_candidates = [
        workbook.with_name(f"{workbook.stem}_report.html"),
        output_root / f"{source_stem}_report.html" if source_stem else None,
        output_root / f"{row.get('device', '')}_report.html",
    ]
    report = next((item for item in report_candidates if item and item.is_file()), report_candidates[0])
    return artifact_id(output_root, report), artifact_id(output_root, workbook)


def sample_results_payload(output_root: Path, definitions: list[dict[str, str]]) -> dict[str, Any]:
    rows = _csv_rows(output_root / "batch_summary" / "master_summary.csv")
    statistics = _csv_rows(output_root / "batch_summary" / "metric_statistics.csv")
    for row in rows:
        row["report_id"], row["workbook_id"] = _device_artifacts(output_root, row)
    ion_conditions = [
        {"id": "configured", "label": "Configured Ion", "ion": "ion_configured_ua_per_um", "ioff": "ioff_configured_ua_per_um", "ratio": "ion_ioff", "log_ratio": "ion_ioff_log10"},
    ]
    if any(row.get("ion_const_vg_ua_per_um") is not None for row in rows):
        ion_conditions.append({"id": "fixed_vg", "label": "Constant Vg / gate field", "ion": "ion_const_vg_ua_per_um", "ioff": "ioff_configured_ua_per_um", "ratio": "ion_ioff_const_vg", "log_ratio": "ion_ioff_const_vg_log10"})
    if any(row.get("ion_const_vov_ua_per_um") is not None for row in rows):
        ion_conditions.append({"id": "fixed_vov", "label": "Constant Vov / overdrive field", "ion": "ion_const_vov_ua_per_um", "ioff": "ioff_configured_ua_per_um", "ratio": "ion_ioff_const_vov", "log_ratio": "ion_ioff_const_vov_log10"})
    if any(row.get("ion_max_ua_per_um") is not None for row in rows):
        ion_conditions.append({"id": "maximum", "label": "Maximum measured Ion", "ion": "ion_max_ua_per_um", "ioff": "ioff_configured_ua_per_um", "ratio": "ion_ioff_max", "log_ratio": "ion_ioff_max_log10"})
    return {
        "available": bool(rows), "rows": rows, "statistics": statistics,
        "definitions": definitions, "ion_conditions": ion_conditions,
        "statistics_provenance": "canonical metric_statistics.csv",
    }


def tlm_results_payload(output_root: Path, batch_rows: list[dict[str, Any]]) -> dict[str, Any]:
    tlm_root = output_root / "TLM"
    summaries = _csv_rows(tlm_root / "tlm_master_summary.csv")
    device_lookup: dict[str, dict[str, Any]] = {}
    for row in batch_rows:
        for candidate in (row.get("device"), row.get("source_file")):
            if candidate:
                device_lookup[_normal_name(candidate)] = row
    structures: list[dict[str, Any]] = []
    if tlm_root.is_dir():
        for workbook in sorted(tlm_root.glob("*/*/*.xlsx")):
            if workbook.name == "tlm_master_summary.xlsx":
                continue
            sample, tlm_id = workbook.parent.parent.name, workbook.parent.name
            level = "master" if tlm_id.casefold() == "master" else "individual"
            report = workbook.with_name(f"{workbook.stem}_report.html")
            metrics = _sheet_rows(workbook, "Metrics vs Lch")
            for row in metrics:
                canonical = device_lookup.get(_normal_name(row.get("source_file")), {})
                for key in (
                    "dibl_mv_v", "ion_configured_ua_per_um", "ioff_configured_ua_per_um",
                    "quality_status", "warning_count", "preferred_bias_v", "selection_reason",
                ):
                    if row.get(key) is None and canonical.get(key) is not None:
                        row[key] = canonical[key]
                row["device_report_id"] = canonical.get("report_id")
                row["device_workbook_id"] = canonical.get("workbook_id")
            fits = _sheet_rows(workbook, "Aggregate Fits" if level == "master" else "TLM Parameters vs Vg")
            if level == "master":
                fits = [row for row in fits if row.get("mode") in {"transfer", "ungated"}]
            if level == "individual":
                for row in fits:
                    row.setdefault("mode", "transfer")
                ungated_fits = _sheet_rows(workbook, "Ungated LTLM Fits")
                for row in ungated_fits:
                    row.setdefault("mode", "ungated")
                fits += ungated_fits
            for row in fits:
                if row.get("rcw_kohm_um") is None and isinstance(row.get("rcw_ohm_um"), (int, float)):
                    row["rcw_kohm_um"] = row["rcw_ohm_um"] / 1000.0
            length_stats = _sheet_rows(workbook, "Length Statistics")
            diagnostics = _sheet_rows(workbook, "Fit Diagnostics" if level == "master" else "Transfer Fit Diagnostics")
            if level == "individual":
                for row in diagnostics:
                    row.setdefault("mode", "transfer")
            structures.append({
                "sample_id": sample, "tlm_id": tlm_id, "analysis_level": level,
                "metrics_vs_lch": metrics, "fits": fits,
                "length_statistics": length_stats, "diagnostics": diagnostics,
                "report_id": artifact_id(output_root, report),
                "workbook_id": artifact_id(output_root, workbook),
            })
    return {"available": bool(summaries or structures), "summaries": summaries, "structures": structures}
