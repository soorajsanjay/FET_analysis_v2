"""
Batch summary: collect metrics from all processed devices and produce
master_summary.xlsx, master_summary.csv, and statistical overview plots.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

import numpy as np
from collections import defaultdict

from fet_analyzer.utils.logging import LOGGER
from fet_analyzer.analysis.metric_summary import (
    METRIC_DEFINITIONS, STATISTIC_METRICS, build_device_summary,
)
from fet_analyzer.path_utils import (
    prepare_write_path, safe_mkdir, safe_path, safe_rmtree, safe_unlink,
)

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    HAS_MPL = True
except ImportError:
    HAS_MPL = False


BATCH_COLUMNS: list[tuple[str, str, int]] = [
    ("sample_id", "Sample ID", 24), ("device", "Device", 30),
    ("type", "Type", 10), ("measurement_type", "Measurement", 16),
    ("channel_length_um", "Lch (μm)", 14),
    ("channel_width_um", "Wch (μm)", 14),
    ("preferred_direction", "Preferred direction", 18), ("preferred_bias_v", "Selected measured Vds (V)", 23),
    ("measured_vd_v", "Measured Vd (V)", 16), ("measured_vs_v", "Measured Vs (V)", 16),
    ("vds_reference", "Vds reference", 28),
    ("selection_reason", "Selection reason", 38),
    ("cox_f_per_cm2", "Effective Cox (F/cm2)", 20), ("cox_source", "Cox source", 22),
    ("oxide_thickness_nm", "Gate dielectric thickness (nm)", 24),
    ("film_thickness_nm", "tch / film thickness (nm)", 22),
    ("gate_dielectric", "Dielectric stack", 24), ("dielectric_constant", "Dielectric constant", 18),
    ("ion_ioff", "Ion/Ioff", 14), ("ion_ioff_log10", "Ion/Ioff (log10)", 16),
    ("ion_ioff_const_vg", "Ion/Ioff const Vg/field", 20), ("ion_ioff_const_vg_log10", "log10 Ion/Ioff const Vg/field", 24),
    ("ion_ioff_const_vov", "Ion/Ioff const Vov/field", 22), ("ion_ioff_const_vov_log10", "log10 Ion/Ioff const Vov/field", 25),
    ("ion_ioff_max", "Ion/Ioff maximum Ion", 20), ("ion_ioff_max_log10", "log10 Ion/Ioff maximum Ion", 23),
    ("ion_configured_a", "Ion configured (A)", 18), ("ion_max_a", "Ion maximum (A)", 17),
    ("ion_configured_ua_per_um", "Ion configured (uA/um)", 21),
    ("ion_max_ua_per_um", "Ion maximum (uA/um)", 20),
    ("ion_const_vov_a", "Ion at const Vov (A)", 19),
    ("ion_const_vov_ua_per_um", "Ion at const Vov (uA/um)", 23),
    ("ion_const_vov_requested_v", "Requested Vov (V)", 18), ("ion_const_vov_v", "Actual Vov used (V)", 18),
    ("ion_const_vov_target_vg_v", "Requested Vg=Vth+Vov (V)", 24), ("ion_const_vov_actual_vg_v", "Actual Vg used for Vov (V)", 23),
    ("ion_overdrive_field_requested_mv_cm", "Requested overdrive field (MV/cm)", 27), ("ion_overdrive_field_actual_mv_cm", "Actual overdrive field (MV/cm)", 25),
    ("ion_const_vg_a", "Ion at const Vg (A)", 18),
    ("ion_const_vg_ua_per_um", "Ion at const Vg (uA/um)", 22),
    ("ion_const_vg_requested_v", "Requested Vg (V)", 17), ("ion_const_vg_v", "Actual Vg used (V)", 17),
    ("ion_gate_field_requested_mv_cm", "Requested gate field (MV/cm)", 24), ("ion_gate_field_actual_mv_cm", "Actual gate field (MV/cm)", 22),
    ("ioff_configured_a", "Ioff configured (A)", 18), ("ioff_min_a", "Ioff minimum (A)", 17),
    ("ioff_configured_ua_per_um", "Ioff configured (uA/um)", 22),
    ("ioff_min_ua_per_um", "Ioff minimum (uA/um)", 21),
    ("ioff_above_ig_a", "Ioff above Ig (A)", 18), ("vth_v", "Vth (V)", 12),
    ("ioff_above_ig_ua_per_um", "Ioff above Ig (uA/um)", 22),
    ("ss_min_mv_dec", "SS min >=1 decade (mV/dec)", 22), ("gm_max_s", "gm,max (S)", 14),
    ("gm_smoothing_mode", "gm smoothing mode", 20),
    ("gm_smoothing_window_points", "gm peak window (points)", 22),
    ("gm_smoothing_status", "gm smoothing status", 20),
    ("mobility_cm2_vs", "Mobility (cm2/V.s)", 18),
    ("hysteresis_abs_v", "Hysteresis abs (V)", 18), ("hysteresis_signed_v", "Hysteresis signed (V)", 20),
    ("dibl_mv_v", "DIBL (mV/V)", 15), ("rout_ohm", "Rout (ohm)", 16),
    ("rout_fit_r2", "Rout fit R2", 14), ("rc_ohm", "Rc group-derived (ohm)", 20),
    ("rcw_ohm_um", "RcW (ohm.um)", 18), ("rho_film_ohm_cm", "Film resistivity (ohm.cm)", 22),
    ("rcw_kohm_um", "Contact resistance (kOhm.um)", 25),
    ("rhoc_ohm_cm2", "Contact resistivity (ohm.cm2)", 24), ("tlm_fit_r2", "TLM fit R2", 14),
    ("tlm_zero_vg_v", "TLM near-zero actual Vg (V)", 25),
    ("tlm_zero_vds_v", "TLM near-zero Vds (V)", 22),
    ("tlm_zero_rc_ohm", "TLM near-zero Rc (ohm)", 22),
    ("tlm_zero_rcw_ohm_um", "TLM near-zero RcW (ohm.um)", 27),
    ("tlm_zero_rcw_kohm_um", "TLM near-zero RcW (kOhm.um)", 28),
    ("tlm_zero_rsh_ohm_sq", "TLM near-zero Rsh (ohm/sq)", 27),
    ("tlm_zero_rho_film_ohm_cm", "TLM near-zero film resistivity (ohm.cm)", 33),
    ("tlm_zero_rhoc_ohm_cm2", "TLM near-zero contact resistivity (ohm.cm2)", 36),
    ("tlm_zero_lt_um", "TLM near-zero LT (um)", 21),
    ("tlm_zero_r2", "TLM near-zero fit R2", 20),
    ("tlm_zero_acceptance_status", "TLM near-zero acceptance", 24),
    ("tlm_zero_acceptance_reasons", "TLM near-zero acceptance reasons", 34),
    ("tlm_zero_analysis_level", "TLM near-zero analysis level", 27),
    ("tlm_zero_tlm_id", "TLM near-zero structure", 24),
    ("tlm_zero_selection_reason", "TLM near-zero selection reason", 55),
    ("tlm_fixed_source", "TLM fixed condition source", 25),
    ("tlm_fixed_requested_vg_v", "TLM fixed requested Vg (V)", 27),
    ("tlm_fixed_requested_field_mv_cm", "TLM fixed requested Eg (MV/cm)", 31),
    ("tlm_fixed_actual_field_mv_cm", "TLM fixed actual Eg (MV/cm)", 28),
    ("tlm_fixed_field_oxide_thickness_nm", "TLM fixed Eg oxide thickness (nm)", 33),
    ("tlm_fixed_vg_v", "TLM fixed actual Vg (V)", 24),
    ("tlm_fixed_vg_delta_v", "TLM fixed actual-requested Vg (V)", 33),
    ("tlm_fixed_vds_v", "TLM fixed Vds (V)", 20),
    ("tlm_fixed_rc_ohm", "TLM fixed Rc (ohm)", 20),
    ("tlm_fixed_rcw_ohm_um", "TLM fixed RcW (ohm.um)", 25),
    ("tlm_fixed_rcw_kohm_um", "TLM fixed RcW (kOhm.um)", 26),
    ("tlm_fixed_rsh_ohm_sq", "TLM fixed Rsh (ohm/sq)", 25),
    ("tlm_fixed_rho_film_ohm_cm", "TLM fixed film resistivity (ohm.cm)", 31),
    ("tlm_fixed_rhoc_ohm_cm2", "TLM fixed contact resistivity (ohm.cm2)", 34),
    ("tlm_fixed_lt_um", "TLM fixed LT (um)", 19),
    ("tlm_fixed_r2", "TLM fixed fit R2", 18),
    ("tlm_fixed_acceptance_status", "TLM fixed acceptance", 22),
    ("tlm_fixed_acceptance_reasons", "TLM fixed acceptance reasons", 32),
    ("tlm_fixed_analysis_level", "TLM fixed analysis level", 25),
    ("tlm_fixed_tlm_id", "TLM fixed structure", 22),
    ("tlm_fixed_selection_reason", "TLM fixed selection reason", 60),
    ("tlm_max_vg_v", "TLM max-|Vg| actual Vg (V)", 27),
    ("tlm_max_vds_v", "TLM max-|Vg| Vds (V)", 23),
    ("tlm_max_rc_ohm", "TLM max-|Vg| Rc (ohm)", 23),
    ("tlm_max_rcw_ohm_um", "TLM max-|Vg| RcW (ohm.um)", 28),
    ("tlm_max_rcw_kohm_um", "TLM max-|Vg| RcW (kOhm.um)", 29),
    ("tlm_max_rsh_ohm_sq", "TLM max-|Vg| Rsh (ohm/sq)", 28),
    ("tlm_max_rho_film_ohm_cm", "TLM max-|Vg| film resistivity (ohm.cm)", 34),
    ("tlm_max_rhoc_ohm_cm2", "TLM max-|Vg| contact resistivity (ohm.cm2)", 37),
    ("tlm_max_lt_um", "TLM max-|Vg| LT (um)", 22),
    ("tlm_max_r2", "TLM max-|Vg| fit R2", 21),
    ("tlm_max_acceptance_status", "TLM max-|Vg| acceptance", 25),
    ("tlm_max_acceptance_reasons", "TLM max-|Vg| acceptance reasons", 35),
    ("tlm_max_analysis_level", "TLM max-|Vg| analysis level", 28),
    ("tlm_max_tlm_id", "TLM max-|Vg| structure", 25),
    ("tlm_max_selection_reason", "TLM max-|Vg| selection reason", 62),
    ("gds_max_s", "gds,max (S)", 14), ("leakage_ratio_max", "Max |Ig|/|Id|", 16),
    ("quality_status", "Quality", 14), ("warning_count", "Warnings", 12),
    ("warning_codes", "Warning codes", 30), ("warning_summary", "Warning summary", 55),
    ("best_ion_configured_a", "Best Ion configured (A)", 20),
    ("best_ioff_configured_a", "Best Ioff configured (A)", 20),
    ("best_ion_ioff", "Best Ion/Ioff", 18), ("best_mobility_cm2_vs", "Best mobility (cm2/V.s)", 22),
    ("best_ss_min_mv_dec", "Best SS (mV/dec)", 18), ("best_gm_max_s", "Best gm (S)", 16),
    ("best_hysteresis_abs_v", "Best abs hysteresis (V)", 22),
    ("best_results_remarks", "Best results remarks", 80),
    ("ion_method_used", "Ion method used", 24), ("ioff_method", "Ioff method used", 26),
    ("ioff_leakage_factor", "Ioff leakage factor", 18), ("source_workbook", "Source workbook", 45),
]

EXTENDED_BATCH_COLUMNS: list[tuple[str, str, int]] = [
    *BATCH_COLUMNS[:4],
    ("sweep_id", "Sweep ID", 12),
    ("vds_v", "Vds (V)", 14),
    *BATCH_COLUMNS[4:],
]


def _summary_metadata(metrics: dict[str, Any], summary: dict[str, Any]) -> dict[str, Any]:
    geometry = metrics.get("device_geometry") or {}
    sources = metrics.get("geometry_sources") or {}
    cox = geometry.get("cox_f_per_cm2")
    cox_source = sources.get("cox_f_per_cm2") or "direct override"
    if cox is None:
        tox, epsilon_r = geometry.get("oxide_thickness_nm"), geometry.get("dielectric_constant")
        if isinstance(tox, (int, float)) and tox > 0 and isinstance(epsilon_r, (int, float)):
            cox = 8.854187817e-14 * float(epsilon_r) / (float(tox) * 1e-7)
            cox_source = "calculated from dielectric constant and oxide thickness"
    warnings: list[tuple[str, str]] = []
    for item in metrics.get("parameter_warnings", []):
        if isinstance(item, dict):
            warnings.append((str(item.get("code") or "parameter"), str(item.get("message") or item)))
        else:
            warnings.append(("parameter", str(item)))
    for item in (metrics.get("quality") or {}).get("flags", []):
        warnings.append((str(item.get("code") or "quality"), str(item.get("message") or item)))
    unique = list(dict.fromkeys(warnings))
    messages = [message for _, message in unique]
    warning_summary = "; ".join(messages[:3])
    if len(messages) > 3:
        warning_summary += f"; +{len(messages) - 3} more"
    best = summary.get("best") or {}
    def best_value(key: str) -> Any:
        item = best.get(key) or {}
        return item.get("value") if isinstance(item, dict) else None
    remarks = " | ".join(
        f"{key}: {item.get('remarks')}" for key, item in best.items()
        if isinstance(item, dict) and item.get("remarks")
    )
    return {
        "channel_length_um": geometry.get("channel_length_um"),
        "channel_width_um": geometry.get("channel_width_um"),
        "cox_f_per_cm2": cox, "cox_source": cox_source,
        "oxide_thickness_nm": geometry.get("oxide_thickness_nm"),
        "film_thickness_nm": geometry.get("film_thickness_nm"),
        "gate_dielectric": geometry.get("gate_dielectric"),
        "dielectric_constant": geometry.get("dielectric_constant"),
        "warning_codes": "; ".join(dict.fromkeys(code for code, _ in unique)),
        "warning_summary": warning_summary,
        "best_ion_configured_a": best_value("ion_configured_a"),
        "best_ioff_configured_a": best_value("ioff_configured_a"),
        "best_ion_ioff": best_value("ion_ioff"),
        "best_mobility_cm2_vs": best_value("mobility_cm2_vs"),
        "best_ss_min_mv_dec": best_value("ss_min_mv_dec"),
        "best_gm_max_s": best_value("gm_max_s"),
        "best_hysteresis_abs_v": best_value("hysteresis_abs_v"),
        "best_results_remarks": remarks,
        "_best_results": best,
    }


def _selected_tlm_summary_rows(path: Path) -> list[dict[str, str]]:
    """Choose one sample-level TLM result, preferring the statistical MASTER."""
    with safe_path(path).open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    grouped: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        grouped[str(row.get("sample_id", ""))].append(row)
    selected = []
    def score(value: Any) -> float:
        try: return float(value)
        except (TypeError, ValueError): return -1.0
    for sample_rows in grouped.values():
        master = next((row for row in sample_rows if row.get("analysis_level") == "master_aggregate" or row.get("tlm_id") == "MASTER"), None)
        if master is not None:
            selected.append(master)
            continue
        rank = {"accepted": 2, "review": 1, "rejected_nonphysical": 0}
        selected.append(max(sample_rows, key=lambda row: (
            rank.get(row.get("transfer_acceptance_status", ""), -1),
            score(row.get("transfer_best_r2")),
        )))
    return selected


def _add_tlm_summary_metrics(row: dict[str, Any], tlm_row: dict[str, str]) -> None:
    """Attach best-fit and selected operating-point TLM values to one batch row."""
    def number(name: str) -> float | None:
        try:
            return float(tlm_row[name]) if tlm_row.get(name) not in (None, "") else None
        except (TypeError, ValueError):
            return None
    prefix = "transfer" if row.get("type") == "transfer" else "output"
    row["rc_ohm"] = number(f"{prefix}_rc_ohm")
    row["rcw_ohm_um"] = number(f"{prefix}_rcw_ohm_um")
    row["rho_film_ohm_cm"] = number(f"{prefix}_rho_film_ohm_cm")
    row["rhoc_ohm_cm2"] = number(f"{prefix}_rhoc_ohm_cm2")
    row["tlm_fit_r2"] = number(f"{prefix}_best_r2")
    # These gated transfer-TLM views are intentionally additive, even when the
    # device row itself came from another measurement type.
    numeric_suffixes = (
        "requested_vg_v", "requested_field_mv_cm", "actual_field_mv_cm",
        "field_oxide_thickness_nm", "vg_v", "vg_delta_v", "vds_v",
        "rc_ohm", "rcw_ohm_um", "rsh_ohm_sq",
        "rho_film_ohm_cm", "rhoc_ohm_cm2", "lt_um", "r2",
    )
    text_suffixes = ("source", "acceptance_status", "acceptance_reasons", "selection_reason")
    for label in ("zero", "fixed", "max"):
        for suffix in numeric_suffixes:
            row[f"tlm_{label}_{suffix}"] = number(f"transfer_{label}_{suffix}")
        rcw = row.get(f"tlm_{label}_rcw_ohm_um")
        row[f"tlm_{label}_rcw_kohm_um"] = rcw / 1000.0 if rcw is not None else None
        for suffix in text_suffixes:
            row[f"tlm_{label}_{suffix}"] = tlm_row.get(f"transfer_{label}_{suffix}") or None
        row[f"tlm_{label}_analysis_level"] = tlm_row.get("analysis_level") or None
        row[f"tlm_{label}_tlm_id"] = tlm_row.get("tlm_id") or None


def _collect_flat_metrics(output_root: Path, metric_files: list[Path]) -> tuple[list[dict[str, Any]], dict[str, list[dict[str, Any]]]]:
    """Read the v2 flat ``<input>_metrics.json`` layout."""
    rows: list[dict[str, Any]] = []
    by_sample: dict[str, list[dict[str, Any]]] = {}
    for metrics_file in metric_files:
        metrics = json.loads(safe_path(metrics_file).read_text(encoding="utf-8"))
        device_name = metrics_file.name.removesuffix("_metrics.json")
        measurement = str(metrics.get("analysis_type") or metrics.get("measurement_type", "")).lower()
        if not measurement:
            measurement = str((metrics.get("classification") or {}).get("type", "")).lower()
        if "output" in measurement:
            kind = "output"
        elif "transfer" in measurement or "idvg" in measurement or "id-vg" in measurement:
            kind = "transfer"
        else:
            continue
        settings = metrics.get("analysis_settings", {})
        summary = metrics.get("device_summary") or build_device_summary(metrics, settings.get("summary_config", {}))
        preferred = summary.get("preferred", {})
        sample_id = metrics.get("sample_id", device_name)
        workbook = output_root / f"{device_name}.xlsx"
        common = {
            **_summary_metadata(metrics, summary),
            "device": device_name, "sample_id": sample_id, "type": kind,
            "measurement_type": metrics.get("measurement_type", kind),
            "source_workbook": str(workbook),
            "quality_status": (metrics.get("quality") or {}).get("status", "unknown"),
            "warning_count": len(metrics.get("parameter_warnings", [])) + len((metrics.get("quality") or {}).get("flags", [])),
            "preferred_direction": preferred.get("preferred_direction") or preferred.get("direction"),
            "preferred_bias_v": preferred.get("preferred_bias_v") if "preferred_bias_v" in preferred else preferred.get("bias_value_v"),
            "measured_vd_v": preferred.get("measured_vd_v"), "measured_vs_v": preferred.get("measured_vs_v"),
            "vds_reference": preferred.get("vds_reference"),
            "selection_reason": preferred.get("selection_reason"),
        }
        if kind == "output":
            row = {**common,
                "gds_max": preferred.get("gds_max_s", metrics.get("gds_max")),
                "gds_max_s": preferred.get("gds_max_s", metrics.get("gds_max")),
                "rout_ohm": preferred.get("rout_ohm"), "rout_fit_r2": preferred.get("rout_fit_r2"),
                "gds_method_used": settings.get("gds_method_used"),
                "resistance_method_used": settings.get("resistance_method_used"),
                "resistance_fit_vd_max_v": settings.get("resistance_fit_vd_max_v"),
            }
        else:
            segments = metrics.get("segments", {})
            def values(key: str) -> list[float]:
                return [float(item[key]) for item in segments.values() if isinstance(item, dict) and isinstance(item.get(key), (int, float)) and np.isfinite(item[key])]
            ion_vg, ion_vov, ioff = values("ion_const_vg_a"), values("ion_const_overdrive_a"), values("ioff_a")
            vg_values, vov_values = values("ion_const_vg_v"), values("ion_const_overdrive_v")
            row = {**common,
                "vth_v": preferred.get("vth_v"), "vth_v_mean": preferred.get("vth_v"),
                "ss_min_mv_dec": preferred.get("ss_min_mv_dec"), "ion_ioff": preferred.get("ion_ioff"),
                "ion_ioff_log10": preferred.get("ion_ioff_log10"), "ion_configured_a": preferred.get("ion_configured_a"),
                "ion_ioff_const_vg": preferred.get("ion_ioff_const_vg"), "ion_ioff_const_vg_log10": preferred.get("ion_ioff_const_vg_log10"),
                "ion_ioff_const_vov": preferred.get("ion_ioff_const_vov"), "ion_ioff_const_vov_log10": preferred.get("ion_ioff_const_vov_log10"),
                "ion_ioff_max": preferred.get("ion_ioff_max"), "ion_ioff_max_log10": preferred.get("ion_ioff_max_log10"),
                "ion_max_a": preferred.get("ion_max_a"),
                "ion_configured_ua_per_um": preferred.get("ion_configured_ua_per_um"),
                "ion_max_ua_per_um": preferred.get("ion_max_ua_per_um"),
                "ion_const_vov_a": preferred.get("ion_const_vov_a") if preferred.get("ion_const_vov_a") is not None else (max(ion_vov) if ion_vov else None),
                "ion_const_vov_ua_per_um": preferred.get("ion_const_vov_ua_per_um"),
                "ion_const_vov_requested_v": preferred.get("ion_const_vov_requested_v"),
                "ion_const_vov_v": preferred.get("ion_const_vov_v") if preferred.get("ion_const_vov_v") is not None else (vov_values[0] if vov_values else settings.get("ion_overdrive_v")),
                "ion_const_vov_target_vg_v": preferred.get("ion_const_vov_target_vg_v"),
                "ion_const_vov_actual_vg_v": preferred.get("ion_const_vov_actual_vg_v"),
                "ion_overdrive_field_requested_mv_cm": preferred.get("ion_overdrive_field_requested_mv_cm"),
                "ion_overdrive_field_actual_mv_cm": preferred.get("ion_overdrive_field_actual_mv_cm"),
                "ion_const_vg_a": preferred.get("ion_const_vg_a") if preferred.get("ion_const_vg_a") is not None else (max(ion_vg) if ion_vg else None),
                "ion_const_vg_ua_per_um": preferred.get("ion_const_vg_ua_per_um"),
                "ion_const_vg_requested_v": preferred.get("ion_const_vg_requested_v"),
                "ion_const_vg_v": preferred.get("ion_const_vg_v") if preferred.get("ion_const_vg_v") is not None else (vg_values[0] if vg_values else settings.get("ion_constant_vg_v")),
                "ion_gate_field_requested_mv_cm": preferred.get("ion_gate_field_requested_mv_cm"),
                "ion_gate_field_actual_mv_cm": preferred.get("ion_gate_field_actual_mv_cm"),
                "ioff_min_a": preferred.get("ioff_min_a"),
                "ioff_min_ua_per_um": preferred.get("ioff_min_ua_per_um"),
                "ioff_above_ig_a": preferred.get("ioff_above_ig_a") if preferred.get("ioff_above_ig_a") is not None else (min(ioff) if ioff else None),
                "ioff_above_ig_ua_per_um": preferred.get("ioff_above_ig_ua_per_um"),
                "ioff_configured_a": preferred.get("ioff_configured_a"),
                "ioff_configured_ua_per_um": preferred.get("ioff_configured_ua_per_um"),
                "ioff_leakage_factor": settings.get("ioff_leakage_factor"),
                "ioff_method": preferred.get("ioff_method_used"), "gm_max_s": preferred.get("gm_max_s"), "gm_max": preferred.get("gm_max_s"),
                "gm_smoothing_mode": preferred.get("gm_smoothing_mode"),
                "gm_smoothing_window_points": preferred.get("gm_smoothing_window_points"),
                "gm_smoothing_status": preferred.get("gm_smoothing_status"),
                "mobility_cm2_vs": preferred.get("mobility_cm2_vs"), "mobility_cm2_vs_max": preferred.get("mobility_cm2_vs"),
                "hysteresis_abs_v": preferred.get("hysteresis_abs_v"), "hysteresis_signed_v": preferred.get("hysteresis_signed_v"),
                "dibl_mv_v": preferred.get("dibl_mv_v"), "leakage_ratio_max": preferred.get("leakage_ratio_max"),
                "ion_method_used": preferred.get("ion_method_used"), "ion_method_configured": settings.get("ion_method_configured"),
                "vth_method": settings.get("vth_method"), "ss_method_used": settings.get("ss_method_used"), "noise_floor_a": settings.get("noise_floor_a"),
            }
        rows.append(row)
        by_sample.setdefault(sample_id, []).append(row)
    for row in rows:
        value = row.get("rcw_ohm_um")
        row["rcw_kohm_um"] = float(value) / 1000.0 if isinstance(value, (int, float)) and np.isfinite(value) else None
    return rows, by_sample


def _collect_preferred_metrics(
    output_root: Path,
) -> tuple[list[dict[str, Any]], dict[str, list[dict[str, Any]]]]:
    """Collect legacy preferred rows used to enrich the canonical sweep table.

    Returns (flat_rows, categorized_by_sample).
    """
    rows: list[dict[str, Any]] = []
    by_sample: dict[str, list[dict[str, Any]]] = {}

    flat_files = sorted(output_root.glob("*_metrics.json"))
    if flat_files:
        rows, by_sample = _collect_flat_metrics(output_root, flat_files)
        tlm_summary = output_root / "TLM" / "tlm_master_summary.csv"
        if tlm_summary.exists():
            for tlm_row in _selected_tlm_summary_rows(tlm_summary):
                    sample = str(tlm_row.get("sample_id", ""))
                    for row in rows:
                        if str(row.get("sample_id", "")) != sample:
                            continue
                        _add_tlm_summary_metrics(row, tlm_row)
                        row["tlm_group_id"] = f"{sample}:{tlm_row.get('tlm_id', '')}"
        for row in rows:
            value = row.get("rcw_ohm_um")
            row["rcw_kohm_um"] = float(value) / 1000.0 if isinstance(value, (int, float)) and np.isfinite(value) else None
        return rows, by_sample

    for metrics_file in sorted(output_root.rglob("transfer_metrics.json")):
        device_dir = metrics_file.parent.parent
        device_name = device_dir.name

        with open(safe_path(metrics_file), encoding="utf-8") as f:
            tmetrics = json.load(f)
        enriched_file = device_dir / "extracted_metrics.json"
        if enriched_file.exists():
            with open(safe_path(enriched_file), encoding="utf-8") as f:
                enriched = json.load(f)
            if enriched.get("sweep_results"):
                tmetrics = enriched

        sample_id = tmetrics.get("sample_id", device_name)
        meas_type = tmetrics.get("measurement_type", "transfer")
        settings = tmetrics.get("analysis_settings", {})
        segments = tmetrics.get("segments", {})
        device_summary = tmetrics.get("device_summary") or build_device_summary(
            tmetrics, settings.get("summary_config", {})
        )
        preferred = device_summary.get("preferred", {})

        def segment_values(key: str) -> list[float]:
            return [
                float(segment[key])
                for segment in segments.values()
                if isinstance(segment, dict)
                and isinstance(segment.get(key), (int, float))
                and np.isfinite(segment[key])
            ]

        ion_const_vg = segment_values("ion_const_vg_a")
        ion_const_vov = segment_values("ion_const_overdrive_a")
        ioff_above_ig = segment_values("ioff_a")
        const_vg = segment_values("ion_const_vg_v")
        const_vov = segment_values("ion_const_overdrive_v")

        def preferred_or(key: str, fallback: Any = None) -> Any:
            value = preferred.get(key)
            return fallback if value is None else value

        row: dict[str, Any] = {
            **_summary_metadata(tmetrics, device_summary),
            "device": device_name,
            "sample_id": sample_id,
            "type": "transfer",
            "measurement_type": meas_type,
            "vth_v": preferred.get("vth_v"),
            "vth_v_mean": preferred.get("vth_v"),
            "ss_min_mv_dec": preferred.get("ss_min_mv_dec"),
            "ion_ioff": preferred.get("ion_ioff"),
            "ion_ioff_log10": preferred.get("ion_ioff_log10"),
            "ion_ioff_const_vg": preferred.get("ion_ioff_const_vg"),
            "ion_ioff_const_vg_log10": preferred.get("ion_ioff_const_vg_log10"),
            "ion_ioff_const_vov": preferred.get("ion_ioff_const_vov"),
            "ion_ioff_const_vov_log10": preferred.get("ion_ioff_const_vov_log10"),
            "ion_ioff_max": preferred.get("ion_ioff_max"),
            "ion_ioff_max_log10": preferred.get("ion_ioff_max_log10"),
            "ion_configured_a": preferred.get("ion_configured_a"),
            "ion_configured_ua_per_um": preferred.get("ion_configured_ua_per_um"),
            "ion_max_a": preferred.get("ion_max_a"),
            "ion_max_ua_per_um": preferred.get("ion_max_ua_per_um"),
            "ion_const_vov_a": preferred_or("ion_const_vov_a", max(ion_const_vov) if ion_const_vov else None),
            "ion_const_vov_ua_per_um": preferred.get("ion_const_vov_ua_per_um"),
            "ion_const_vov_requested_v": preferred.get("ion_const_vov_requested_v"),
            "ion_const_vov_v": preferred_or("ion_const_vov_v", const_vov[0] if const_vov else settings.get("ion_overdrive_v")),
            "ion_const_vov_target_vg_v": preferred.get("ion_const_vov_target_vg_v"),
            "ion_const_vov_actual_vg_v": preferred.get("ion_const_vov_actual_vg_v"),
            "ion_overdrive_field_requested_mv_cm": preferred.get("ion_overdrive_field_requested_mv_cm"),
            "ion_overdrive_field_actual_mv_cm": preferred.get("ion_overdrive_field_actual_mv_cm"),
            "ion_const_vg_a": preferred_or("ion_const_vg_a", max(ion_const_vg) if ion_const_vg else None),
            "ion_const_vg_ua_per_um": preferred.get("ion_const_vg_ua_per_um"),
            "ion_const_vg_requested_v": preferred.get("ion_const_vg_requested_v"),
            "ion_const_vg_v": preferred_or("ion_const_vg_v", const_vg[0] if const_vg else settings.get("ion_constant_vg_v")),
            "ion_gate_field_requested_mv_cm": preferred.get("ion_gate_field_requested_mv_cm"),
            "ion_gate_field_actual_mv_cm": preferred.get("ion_gate_field_actual_mv_cm"),
            "ioff_min_a": preferred.get("ioff_min_a"),
            "ioff_min_ua_per_um": preferred.get("ioff_min_ua_per_um"),
            "ioff_above_ig_a": preferred_or("ioff_above_ig_a", min(ioff_above_ig) if ioff_above_ig else None),
            "ioff_above_ig_ua_per_um": preferred.get("ioff_above_ig_ua_per_um"),
            "ioff_configured_a": preferred.get("ioff_configured_a"),
            "ioff_configured_ua_per_um": preferred.get("ioff_configured_ua_per_um"),
            "ioff_leakage_factor": settings.get("ioff_leakage_factor"),
            "ioff_method": preferred.get("ioff_method_used"),
            "gm_max_s": preferred.get("gm_max_s"),
            "gm_max": preferred.get("gm_max_s"),
            "mobility_cm2_vs": preferred.get("mobility_cm2_vs"),
            "mobility_cm2_vs_max": preferred.get("mobility_cm2_vs"),
            "hysteresis_abs_v": preferred.get("hysteresis_abs_v"),
            "hysteresis_signed_v": preferred.get("hysteresis_signed_v"),
            "dibl_mv_v": preferred.get("dibl_mv_v"),
            "leakage_ratio_max": preferred.get("leakage_ratio_max"),
            "preferred_direction": preferred.get("preferred_direction") or preferred.get("direction"),
            "preferred_bias_v": preferred.get("preferred_bias_v") if "preferred_bias_v" in preferred else preferred.get("bias_value_v"),
            "measured_vd_v": preferred.get("measured_vd_v"), "measured_vs_v": preferred.get("measured_vs_v"),
            "vds_reference": preferred.get("vds_reference"),
            "selection_reason": preferred.get("selection_reason"),
            "source_workbook": str(device_dir / "analysis.xlsx"),
            "quality_status": (tmetrics.get("quality") or {}).get("status", "unknown"),
            "warning_count": len(tmetrics.get("parameter_warnings", [])) + len((tmetrics.get("quality") or {}).get("flags", [])),
            "ion_method_used": preferred.get("ion_method_used"),
            "ion_method_configured": settings.get("ion_method_configured"),
            "vth_method": settings.get("vth_method"),
            "ss_method_used": settings.get("ss_method_used"),
            "noise_floor_a": settings.get("noise_floor_a"),
        }

        rows.append(row)
        by_sample.setdefault(sample_id, []).append(row)

        # Also check output metrics
        output_file = device_dir / "plots" / "output_metrics.json"
        if output_file.exists():
            with open(safe_path(output_file), encoding="utf-8") as f:
                ometrics = json.load(f)
            enriched_file = device_dir / "extracted_metrics.json"
            if enriched_file.exists():
                with open(safe_path(enriched_file), encoding="utf-8") as f:
                    enriched = json.load(f)
                if enriched.get("sweep_results"):
                    ometrics = enriched
            output_settings = ometrics.get("analysis_settings", {})
            output_summary = ometrics.get("device_summary") or build_device_summary(ometrics, {})
            output_preferred = output_summary.get("preferred", {})
            row_out = {
                **_summary_metadata(ometrics, output_summary),
                "device": device_name,
                "sample_id": ometrics.get("sample_id", sample_id),
                "type": "output",
                "measurement_type": ometrics.get("measurement_type", "output"),
                "gds_max": output_preferred.get("gds_max_s", ometrics.get("gds_max")),
                "gds_max_s": output_preferred.get("gds_max_s", ometrics.get("gds_max")),
                "rout_ohm": output_preferred.get("rout_ohm"),
                "rout_fit_r2": output_preferred.get("rout_fit_r2"),
                "preferred_direction": output_preferred.get("direction"),
                "preferred_bias_v": output_preferred.get("bias_value_v"),
                "selection_reason": output_preferred.get("selection_reason"),
                "source_workbook": str(device_dir / "analysis.xlsx"),
                "quality_status": (ometrics.get("quality") or {}).get("status", "unknown"),
                "warning_count": len(ometrics.get("parameter_warnings", [])) + len((ometrics.get("quality") or {}).get("flags", [])),
                "gds_method_used": output_settings.get("gds_method_used"),
                "resistance_method_used": output_settings.get("resistance_method_used"),
                "resistance_fit_vd_max_v": output_settings.get("resistance_fit_vd_max_v"),
            }
            rows.append(row_out)
            by_sample.setdefault(row_out["sample_id"], []).append(row_out)

    # Also collect standalone output metrics (devices with only output data)
    for metrics_file in sorted(output_root.rglob("output_metrics.json")):
        device_dir = metrics_file.parent.parent
        device_name = device_dir.name

        # Skip if already picked up via transfer_metrics
        if any(r["device"] == device_name and r.get("type") == "output" for r in rows):
            continue

        with open(safe_path(metrics_file), encoding="utf-8") as f:
            ometrics = json.load(f)
        enriched_file = device_dir / "extracted_metrics.json"
        if enriched_file.exists():
            with open(safe_path(enriched_file), encoding="utf-8") as f:
                enriched = json.load(f)
            if enriched.get("sweep_results"):
                ometrics = enriched
        output_settings = ometrics.get("analysis_settings", {})
        output_summary = ometrics.get("device_summary") or build_device_summary(ometrics, {})
        output_preferred = output_summary.get("preferred", {})

        row_out = {
            **_summary_metadata(ometrics, output_summary),
            "device": device_name,
            "sample_id": ometrics.get("sample_id", device_name),
            "type": "output",
            "measurement_type": ometrics.get("measurement_type", "output"),
            "gds_max": output_preferred.get("gds_max_s", ometrics.get("gds_max")),
            "gds_max_s": output_preferred.get("gds_max_s", ometrics.get("gds_max")),
            "rout_ohm": output_preferred.get("rout_ohm"),
            "rout_fit_r2": output_preferred.get("rout_fit_r2"),
            "preferred_direction": output_preferred.get("direction"),
            "preferred_bias_v": output_preferred.get("bias_value_v"),
            "selection_reason": output_preferred.get("selection_reason"),
            "source_workbook": str(device_dir / "analysis.xlsx"),
            "quality_status": (ometrics.get("quality") or {}).get("status", "unknown"),
            "warning_count": len(ometrics.get("parameter_warnings", [])) + len((ometrics.get("quality") or {}).get("flags", [])),
            "gds_method_used": output_settings.get("gds_method_used"),
            "resistance_method_used": output_settings.get("resistance_method_used"),
            "resistance_fit_vd_max_v": output_settings.get("resistance_fit_vd_max_v"),
        }
        rows.append(row_out)
        by_sample.setdefault(row_out["sample_id"], []).append(row_out)

    # Join group-derived contact metrics without presenting them as independent
    # per-device fits. Statistics later deduplicate these values by tlm_group_id.
    tlm_summary = output_root / "TLM" / "tlm_master_summary.csv"
    if tlm_summary.exists():
        for tlm_row in _selected_tlm_summary_rows(tlm_summary):
                sample = str(tlm_row.get("sample_id", ""))
                group_id = f"{sample}:{tlm_row.get('tlm_id', '')}"
                for row in rows:
                    if str(row.get("sample_id", "")) != sample:
                        continue
                    _add_tlm_summary_metrics(row, tlm_row)
                    row["tlm_group_id"] = group_id

    for row in rows:
        value = row.get("rcw_ohm_um")
        row["rcw_kohm_um"] = float(value) / 1000.0 if isinstance(value, (int, float)) and np.isfinite(value) else None
    return rows, by_sample


def _physical_sweep_vds(sweep: dict[str, Any]) -> float | None:
    """Return physical Vds for one transfer sweep, preserving measured provenance."""
    def number(value: Any) -> float | None:
        if isinstance(value, (int, float)) and not isinstance(value, bool) and np.isfinite(value):
            return float(value)
        return None

    measured_vd = number(sweep.get("measured_vd_v"))
    measured_vs = number(sweep.get("measured_vs_v"))
    if measured_vd is not None and measured_vs is not None:
        return float(f"{measured_vd - measured_vs:.12g}")
    measured_vds = number(sweep.get("measured_vds_v"))
    fallback = measured_vds if measured_vds is not None else number(sweep.get("bias_value_v"))
    return float(f"{fallback:.12g}") if fallback is not None else None


def _extended_transfer_row(
    base_row: dict[str, Any], sweep: dict[str, Any], sweep_id: int,
) -> dict[str, Any]:
    """Overlay one canonical transfer sweep onto an enriched device-level row."""
    row = dict(base_row)
    for key, _, _ in BATCH_COLUMNS:
        if key in sweep:
            row[key] = sweep.get(key)

    vds = _physical_sweep_vds(sweep)
    row.update({
        "sweep_id": sweep_id,
        "vds_v": vds,
        "preferred_direction": sweep.get("direction"),
        "preferred_bias_v": vds,
        "measured_vd_v": sweep.get("measured_vd_v"),
        "measured_vs_v": sweep.get("measured_vs_v"),
        "vds_reference": sweep.get("vds_reference"),
        "selection_reason": "canonical per-sweep transfer result; not a preferred-device selection",
        "vth_v_mean": sweep.get("vth_v"),
        "gm_max": sweep.get("gm_max_s"),
        "mobility_cm2_vs_max": sweep.get("mobility_cm2_vs"),
        "ioff_method": sweep.get("ioff_method_used"),
        "ion_method_used": sweep.get("ion_method_used"),
        "_canonical_sweep_id": sweep.get("sweep_id"),
        "_master_selection_reason": base_row.get("selection_reason"),
        "_master_warning_count": base_row.get("warning_count"),
        "_master_warning_codes": base_row.get("warning_codes"),
        "_master_warning_summary": base_row.get("warning_summary"),
        "_master_hysteresis_abs_v": base_row.get("hysteresis_abs_v"),
    })

    # Extraction eligibility is authoritative. An ineligible value must not leak
    # into an extended row even when reading an older, partially enriched payload.
    if sweep.get("gm_metric_eligible") is False or sweep.get("gm_bias_homogeneous") is False:
        row["gm_max_s"] = None
        row["gm_max"] = None
    if sweep.get("mobility_metric_eligible") is False or sweep.get("mobility_bias_homogeneous") is False:
        row["mobility_cm2_vs"] = None
        row["mobility_cm2_vs_max"] = None

    sweep_warnings = [
        item.strip() for item in str(sweep.get("warnings") or "").split(";") if item.strip()
    ]
    if sweep_warnings:
        row["warning_count"] = int(row.get("warning_count") or 0) + len(sweep_warnings)
        existing_codes = str(row.get("warning_codes") or "").strip()
        row["warning_codes"] = "; ".join(filter(None, (existing_codes, "sweep")))
        existing_summary = str(row.get("warning_summary") or "").strip()
        row["warning_summary"] = "; ".join(filter(None, (
            existing_summary, "; ".join(sweep_warnings),
        )))
    return row


def collect_extended_metrics(
    output_root: Path,
    master_rows: list[dict[str, Any]] | None = None,
) -> list[dict[str, Any]]:
    """Collect one enriched row per canonical transfer sweep."""
    if master_rows is None:
        master_rows, _ = _collect_preferred_metrics(output_root)

    transfer_rows: dict[tuple[str, str], dict[str, Any]] = {}
    for row in master_rows:
        if row.get("type") != "transfer":
            continue
        key = (str(row.get("sample_id", "")), str(row.get("device", "")))
        if key in transfer_rows:
            raise ValueError(f"Duplicate transfer master row for sample/device {key!r}")
        transfer_rows[key] = row

    sources: list[tuple[Path, str]] = []
    flat_files = sorted(output_root.glob("*_metrics.json"))
    if flat_files:
        sources = [(path, path.name.removesuffix("_metrics.json")) for path in flat_files]
    else:
        sources = [
            (path, path.parent.parent.name)
            for path in sorted(output_root.rglob("transfer_metrics.json"))
        ]

    rows: list[dict[str, Any]] = []
    identities: set[tuple[str, int]] = set()
    for metrics_file, device_name in sources:
        metrics = json.loads(safe_path(metrics_file).read_text(encoding="utf-8"))
        measurement = str(
            metrics.get("analysis_type") or metrics.get("measurement_type")
            or (metrics.get("classification") or {}).get("type", "")
        ).lower()
        if "output" in measurement or not any(
            token in measurement for token in ("transfer", "idvg", "id-vg")
        ):
            continue

        if not flat_files:
            enriched_file = metrics_file.parent.parent / "extracted_metrics.json"
            if enriched_file.exists():
                enriched = json.loads(
                    safe_path(enriched_file).read_text(encoding="utf-8")
                )
                if enriched.get("sweep_results"):
                    metrics = enriched

        settings = metrics.get("analysis_settings", {})
        summary = metrics.get("device_summary") or {}
        segments = summary.get("segments")
        if not isinstance(segments, list) and metrics.get("sweep_results"):
            summary = build_device_summary(metrics, settings.get("summary_config", {}))
            segments = summary.get("segments")
        if not isinstance(segments, list):
            LOGGER.warning("No canonical transfer sweeps found for extended summary: %s", metrics_file)
            continue

        sample_id = str(metrics.get("sample_id", device_name))
        base_row = transfer_rows.get((sample_id, device_name))
        if base_row is None:
            LOGGER.warning("No transfer master row found for extended summary: %s", metrics_file)
            continue
        for ordinal, sweep in enumerate(segments, 1):
            if not isinstance(sweep, dict):
                continue
            identity = (device_name, ordinal)
            if identity in identities:
                raise ValueError(f"Duplicate extended summary identity {identity!r}")
            identities.add(identity)
            row = _extended_transfer_row(base_row, sweep, ordinal)
            preferred = summary.get("preferred") or {}
            row["_is_preferred"] = bool(
                sweep.get("sweep_id") is not None
                and sweep.get("sweep_id") == preferred.get("sweep_id")
            )
            if preferred.get("sweep_id") is None:
                row["_is_preferred"] = (
                    sweep.get("direction") == (
                        preferred.get("preferred_direction") or preferred.get("direction")
                    )
                    and _physical_sweep_vds(sweep) == _physical_sweep_vds(preferred)
                )
            rows.append(row)

    rows.sort(key=lambda row: (
        str(row.get("sample_id", "")), str(row.get("device", "")), int(row.get("sweep_id", 0)),
    ))
    return rows


def _derive_master_metrics(
    extended_rows: list[dict[str, Any]],
    preferred_rows: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], dict[str, list[dict[str, Any]]]]:
    """Collapse canonical sweep rows to one preferred transfer row per device."""
    grouped: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in extended_rows:
        grouped[(str(row.get("sample_id", "")), str(row.get("device", "")))].append(row)

    rows: list[dict[str, Any]] = []
    for key, candidates in grouped.items():
        selected = [row for row in candidates if row.get("_is_preferred")]
        if len(selected) != 1:
            raise ValueError(
                f"Expected exactly one preferred extended row for sample/device {key!r}; "
                f"found {len(selected)}"
            )
        row = dict(selected[0])
        row["selection_reason"] = row.get("_master_selection_reason")
        row["warning_count"] = row.get("_master_warning_count")
        row["warning_codes"] = row.get("_master_warning_codes")
        row["warning_summary"] = row.get("_master_warning_summary")
        # The ordinary master intentionally reports the device-wide maximum
        # absolute hysteresis, while the extended row reports this sweep pair.
        row["hysteresis_abs_v"] = row.get("_master_hysteresis_abs_v")
        rows.append(row)

    derived_keys = grouped.keys()
    for row in preferred_rows:
        key = (str(row.get("sample_id", "")), str(row.get("device", "")))
        if row.get("type") == "transfer" and key not in derived_keys:
            # Older saved results may predate canonical sweep_results. Preserve
            # their ordinary master row, but they cannot appear in the extended file.
            LOGGER.warning(
                "Transfer master row has no canonical sweeps and cannot be derived: %s",
                row.get("device"),
            )
            rows.append(row)

    # Output measurements are intentionally absent from the extended transfer
    # table, so retain their independently selected rows in the ordinary master.
    rows.extend(row for row in preferred_rows if row.get("type") == "output")
    rows.sort(key=lambda row: (
        str(row.get("sample_id", "")), str(row.get("device", "")), str(row.get("type", "")),
    ))
    by_sample: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        by_sample.setdefault(str(row.get("sample_id", "")), []).append(row)
    return rows, by_sample


def collect_metrics(
    output_root: Path,
) -> tuple[list[dict[str, Any]], dict[str, list[dict[str, Any]]]]:
    """Collect master rows, deriving transfer selections from per-sweep rows."""
    preferred_rows, _ = _collect_preferred_metrics(output_root)
    extended_rows = collect_extended_metrics(output_root, preferred_rows)
    return _derive_master_metrics(extended_rows, preferred_rows)


def write_extended_summary_excel(
    rows: list[dict[str, Any]], output_path: Path,
) -> None:
    """Write the canonical per-transfer-sweep workbook without aggregate statistics."""
    if not HAS_OPENPYXL:
        LOGGER.warning("openpyxl not available — writing extended CSV only")
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Extended FET Summary"
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    for column, (_, label, width) in enumerate(EXTENDED_BATCH_COLUMNS, 1):
        cell = ws.cell(row=1, column=column, value=label)
        cell.font, cell.fill, cell.alignment, cell.border = (
            header_font, header_fill, header_align, thin_border,
        )
        ws.column_dimensions[openpyxl.utils.get_column_letter(column)].width = width
    for row_index, row in enumerate(rows, 2):
        for column, (key, _, _) in enumerate(EXTENDED_BATCH_COLUMNS, 1):
            value = row.get(key)
            cell = ws.cell(row=row_index, column=column, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center")
            if isinstance(value, float):
                cell.number_format = "0.000E+00" if abs(value) < 0.01 and value != 0 else "0.00"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    if rows:
        from openpyxl.worksheet.table import Table, TableStyleInfo
        ref = f"A1:{openpyxl.utils.get_column_letter(ws.max_column)}{ws.max_row}"
        table = Table(displayName="ExtendedBatchTable", ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    wb.save(prepare_write_path(output_path))
    wb.close()
    LOGGER.info("Extended master summary written to %s", output_path)


def write_extended_summary_csv(
    rows: list[dict[str, Any]], output_path: Path,
) -> None:
    """Write the canonical per-transfer-sweep CSV, including headers when empty."""
    with prepare_write_path(output_path).open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=[key for key, _, _ in EXTENDED_BATCH_COLUMNS],
            extrasaction="ignore",
        )
        writer.writeheader()
        writer.writerows(rows)
    LOGGER.info("Extended master CSV written to %s", output_path)


def write_summary_excel(
    rows: list[dict[str, Any]],
    output_path: Path,
) -> None:
    """Write master_summary.xlsx with formatting."""
    if not HAS_OPENPYXL:
        LOGGER.warning("openpyxl not available — writing CSV only")
        return

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FET Summary"

    # Header styling
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    # Columns
    columns = [
        ("Sample ID", 24),
        ("Device", 30),
        ("Type", 10),
        ("Measurement", 16),
        ("Vth (V)", 12),
        ("SS (mV/dec)", 14),
        ("Ion/Ioff (log10)", 16),
        ("Ion at const Vov (A)", 19),
        ("Const Vov (V)", 14),
        ("Ion at const Vg (A)", 18),
        ("Const Vg (V)", 13),
        ("Ioff above Ig (A)", 18),
        ("Ioff leakage factor", 18),
        ("Ioff method", 30),
        ("μ_FE (cm²/V·s)", 16),
        ("gm,max", 14),
        ("gds,max", 14),
        ("Ion method used", 24),
        ("Configured Ion method", 24),
        ("Vth method", 22),
        ("SS method used", 34),
        ("Noise floor (A)", 16),
        ("Resistance method", 28),
        ("Resistance |Vd| max (V)", 22),
    ]

    columns = [(label, width) for _, label, width in BATCH_COLUMNS]

    # Write header
    for col_idx, (name, width) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    # Data rows
    for row_idx, row_data in enumerate(rows, 2):
        values = [
            row_data.get("sample_id", ""),
            row_data.get("device", ""),
            row_data.get("type", ""),
            row_data.get("measurement_type", ""),
            row_data.get("vth_v_mean"),
            row_data.get("ss_min_mv_dec"),
            row_data.get("ion_ioff_log10"),
            row_data.get("ion_const_vov_a"),
            row_data.get("ion_const_vov_v"),
            row_data.get("ion_const_vg_a"),
            row_data.get("ion_const_vg_v"),
            row_data.get("ioff_above_ig_a"),
            row_data.get("ioff_leakage_factor"),
            row_data.get("ioff_method"),
            row_data.get("mobility_cm2_vs_max"),
            row_data.get("gm_max"),
            row_data.get("gds_max"),
            row_data.get("ion_method_used"),
            row_data.get("ion_method_configured"),
            row_data.get("vth_method"),
            row_data.get("ss_method_used"),
            row_data.get("noise_floor_a"),
            row_data.get("resistance_method_used"),
            row_data.get("resistance_fit_vd_max_v"),
        ]
        values = [row_data.get(key) for key, _, _ in BATCH_COLUMNS]
        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center")
            if isinstance(val, float) and val is not None:
                cell.number_format = "0.000E+00" if abs(val) < 0.01 and val != 0 else "0.00"

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    # Dashboard — compact run overview with links to source workbooks.
    dashboard = wb.create_sheet("Dashboard", 0)
    dashboard.append(["FET Analysis Dashboard"])
    dashboard["A1"].font = Font(size=18, bold=True, color="2F5496")
    dashboard.append(["Measurements", len(rows)])
    dashboard.append(["Samples", len({r.get("sample_id") for r in rows})])
    dashboard.append(["Warnings", sum(int(r.get("warning_count", 0)) for r in rows)])
    dashboard.append([])
    dashboard.append(["Sample ID", "Measurements", "Transfer", "Output", "Warnings"])
    for cell in dashboard[6]:
        cell.font, cell.fill = header_font, header_fill
    for sample in sorted({str(r.get("sample_id", "")) for r in rows}):
        subset = [r for r in rows if str(r.get("sample_id", "")) == sample]
        dashboard.append([sample, len(subset), sum(r.get("type") == "transfer" for r in subset),
                          sum(r.get("type") == "output" for r in subset),
                          sum(int(r.get("warning_count", 0)) for r in subset)])
    dashboard.column_dimensions["A"].width = 28
    for col in "BCDE": dashboard.column_dimensions[col].width = 15
    dashboard.freeze_panes = "A6"

    # Statistically defensible sample summaries: do not mix metric names.
    stat_sheet = wb.create_sheet("Sample Statistics")
    stat_headers = ["Sample ID", "Metric", "N files", "N valid", "Missing", "Mean", "Median", "Std dev", "Min", "Max", "SEM", "95% CI half-width"]
    stat_sheet.append(stat_headers)
    for cell in stat_sheet[1]: cell.font, cell.fill, cell.alignment = header_font, header_fill, header_align
    metric_keys = ["vth_v_mean", "ss_min_mv_dec", "ion_ioff_log10", "ion_const_vov_a",
                   "ion_const_vg_a", "ioff_above_ig_a", "mobility_cm2_vs_max", "gm_max", "gds_max"]
    metric_keys = STATISTIC_METRICS
    for sample in sorted({str(r.get("sample_id", "")) for r in rows}):
        subset = [r for r in rows if str(r.get("sample_id", "")) == sample]
        for metric in metric_keys:
            source_rows = subset
            if metric in {"rc_ohm", "rcw_ohm_um", "tlm_fit_r2"}:
                deduped = {}
                for item in subset:
                    deduped.setdefault(item.get("tlm_group_id") or item.get("device"), item)
                source_rows = list(deduped.values())
            values = np.asarray([r[metric] for r in source_rows if isinstance(r.get(metric), (int, float)) and np.isfinite(r[metric])], dtype=float)
            if not len(values): continue
            std = float(np.std(values, ddof=1)) if len(values) > 1 else None
            sem = std / np.sqrt(len(values)) if std is not None else None
            stat_sheet.append([sample, metric, len(source_rows), len(values), len(source_rows) - len(values),
                               float(np.mean(values)), float(np.median(values)), std,
                               float(np.min(values)), float(np.max(values)), sem,
                               1.96 * sem if sem is not None else None])
    stat_sheet.freeze_panes = "A2"; stat_sheet.auto_filter.ref = stat_sheet.dimensions
    for col in range(1, 13): stat_sheet.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 20

    warning_sheet = wb.create_sheet("Warnings")
    warning_sheet.append(["Sample ID", "Device", "Status", "Warning Count", "Warning Codes", "Brief Warning", "Source Workbook"])
    for cell in warning_sheet[1]: cell.font, cell.fill, cell.alignment = header_font, header_fill, header_align
    for item in rows:
        if int(item.get("warning_count", 0)):
            warning_sheet.append([
                item.get("sample_id"), item.get("device"), item.get("quality_status"),
                item.get("warning_count"), item.get("warning_codes"),
                item.get("warning_summary"), item.get("source_workbook"),
            ])
    warning_sheet.freeze_panes = "A2"; warning_sheet.auto_filter.ref = warning_sheet.dimensions
    for col, width in enumerate([24, 32, 18, 16, 32, 80, 45], 1):
        warning_sheet.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width

    best_sheet = wb.create_sheet("Best per File")
    best_headers = ["Sample ID", "Device", "Metric", "Value", "Unit", "Selection", "Measured Vds (V)", "Sweep type", "Direction", "Sweep ID", "Remarks"]
    best_sheet.append(best_headers)
    for cell in best_sheet[1]: cell.font, cell.fill, cell.alignment = header_font, header_fill, header_align
    units = {item["key"]: item["unit"] for item in METRIC_DEFINITIONS}
    for row in rows:
        for metric, item in (row.get("_best_results") or {}).items():
            if not isinstance(item, dict):
                continue
            best_sheet.append([
                row.get("sample_id"), row.get("device"), metric,
                item.get("range") if metric == "vth_v_range" else item.get("value"),
                "V" if metric == "vth_v_range" else units.get(metric, ""), item.get("selection"),
                item.get("measured_vds_v"), item.get("sweep_type"), item.get("direction"),
                item.get("sweep_id"), item.get("remarks"),
            ])
    best_sheet.freeze_panes = "A2"; best_sheet.auto_filter.ref = best_sheet.dimensions
    for column, width in enumerate([24, 32, 28, 16, 14, 18, 18, 18, 22, 45, 110], 1):
        best_sheet.column_dimensions[openpyxl.utils.get_column_letter(column)].width = width

    from openpyxl.worksheet.table import Table, TableStyleInfo

    def _dedupe_header_row(sheet) -> None:
        seen: dict[str, int] = {}
        for column in range(1, sheet.max_column + 1):
            cell = sheet.cell(row=1, column=column)
            base = str(cell.value).strip() if cell.value is not None else ""
            if not base:
                base = f"Column{column}"
            count = seen.get(base, 0) + 1
            seen[base] = count
            cell.value = base if count == 1 else f"{base}_{count}"

    for index, sheet in enumerate((ws, stat_sheet, warning_sheet, best_sheet), 1):
        if sheet.max_row > 1 and sheet.max_column > 0:
            _dedupe_header_row(sheet)
            ref = f"A1:{openpyxl.utils.get_column_letter(sheet.max_column)}{sheet.max_row}"
            table = Table(displayName=f"BatchTable{index}", ref=ref)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            sheet.add_table(table)
    wb.save(prepare_write_path(output_path))
    LOGGER.info("Master summary written to %s", output_path)


def write_summary_csv(
    rows: list[dict[str, Any]],
    by_sample: dict[str, list[dict[str, Any]]],
    output_path: Path,
) -> None:
    """Write master_summary.csv and samples_summary.csv."""
    if not rows:
        return
    fieldnames = [
        "sample_id", "device", "type", "measurement_type",
        "vth_v_mean", "ss_min_mv_dec", "ion_ioff_log10",
        "ion_const_vov_a", "ion_const_vov_v", "ion_const_vg_a", "ion_const_vg_v",
        "ioff_above_ig_a", "ioff_leakage_factor", "ioff_method",
        "mobility_cm2_vs_max", "gm_max", "gds_max",
        "ion_method_used", "ion_method_configured", "vth_method",
        "ss_method_used", "noise_floor_a", "gds_method_used",
        "resistance_method_used", "resistance_fit_vd_max_v",
    ]
    fieldnames = [key for key, _, _ in BATCH_COLUMNS]
    with open(prepare_write_path(output_path), "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    LOGGER.info("Master CSV written to %s", output_path)

    best_path = output_path.parent / "best_per_file.csv"
    best_fields = ["sample_id", "device", "metric", "value", "unit", "selection", "measured_vds_v", "sweep_type", "direction", "sweep_id", "remarks"]
    units = {item["key"]: item["unit"] for item in METRIC_DEFINITIONS}
    with open(prepare_write_path(best_path), "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=best_fields)
        writer.writeheader()
        for row in rows:
            for metric, item in (row.get("_best_results") or {}).items():
                if not isinstance(item, dict):
                    continue
                writer.writerow({
                    "sample_id": row.get("sample_id"), "device": row.get("device"), "metric": metric,
                    "value": item.get("range") if metric == "vth_v_range" else item.get("value"),
                    "unit": "V" if metric == "vth_v_range" else units.get(metric, ""),
                    "selection": item.get("selection"), "measured_vds_v": item.get("measured_vds_v"),
                    "sweep_type": item.get("sweep_type"), "direction": item.get("direction"),
                    "sweep_id": item.get("sweep_id"), "remarks": item.get("remarks"),
                })
    LOGGER.info("Best-per-file CSV written to %s", best_path)

    # Per-sample grouped summary
    sample_path = output_path.parent / "samples_summary.csv"
    with open(prepare_write_path(sample_path), "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["sample_id", "n_devices", "n_transfer", "n_output",
                          "ss_min_mv_dec", "ion_ioff_log10_min", "ion_ioff_log10_max"])
        for sample_id, srows in sorted(by_sample.items()):
            devices = set(r["device"] for r in srows)
            n_transfer = sum(1 for r in srows if r.get("type") == "transfer")
            n_output = sum(1 for r in srows if r.get("type") == "output")
            ss_vals = [r["ss_min_mv_dec"] for r in srows
                       if r.get("ss_min_mv_dec") is not None and r.get("type") == "transfer"]
            io_vals = [r["ion_ioff_log10"] for r in srows
                       if r.get("ion_ioff_log10") is not None and r.get("type") == "transfer"]
            writer.writerow([
                sample_id, len(devices), n_transfer, n_output,
                f"{min(ss_vals):.0f}" if ss_vals else "",
                f"{min(io_vals):.1f}" if io_vals else "",
                f"{max(io_vals):.1f}" if io_vals else "",
            ])
    LOGGER.info("Samples CSV written to %s", sample_path)

    statistics_path = output_path.parent / "metric_statistics.csv"
    with open(prepare_write_path(statistics_path), "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["sample_id", "metric", "n_files", "n_valid", "missing",
                         "mean", "median", "std_dev", "min", "max", "sem", "ci95_half_width"])
        for sample_id, srows in sorted(by_sample.items()):
            for metric in STATISTIC_METRICS:
                source_rows = srows
                if metric in {"rc_ohm", "rcw_ohm_um", "tlm_fit_r2"}:
                    deduped = {}
                    for item in srows:
                        deduped.setdefault(item.get("tlm_group_id") or item.get("device"), item)
                    source_rows = list(deduped.values())
                values = np.asarray([
                    row[metric] for row in source_rows
                    if isinstance(row.get(metric), (int, float)) and np.isfinite(row[metric])
                ], dtype=float)
                if not len(values):
                    continue
                std = float(np.std(values, ddof=1)) if len(values) > 1 else None
                sem = std / np.sqrt(len(values)) if std is not None else None
                writer.writerow([
                    sample_id, metric, len(source_rows), len(values), len(source_rows) - len(values),
                    float(np.mean(values)), float(np.median(values)), std,
                    float(np.min(values)), float(np.max(values)), sem,
                    1.96 * sem if sem is not None else None,
                ])
    LOGGER.info("Metric statistics CSV written to %s", statistics_path)


def plot_batch_statistics(
    rows: list[dict[str, Any]],
    by_sample: dict[str, list[dict[str, Any]]],
    output_dir: Path,
    config: dict[str, Any] | None = None,
) -> None:
    """Generate batch statistical overview plots, grouped by sample ID."""
    if not HAS_MPL or not rows:
        return

    plot_cfg = (config or {}).get("plots", {})
    colors = (config or {}).get("plots", {}).get("palette", [
        "#0072B2", "#E69F00", "#009E73",
    ])

    plt.rcParams.update({
        "font.size": plot_cfg.get("font_size", 11),
        "figure.dpi": plot_cfg.get("dpi", 150),
        "savefig.dpi": plot_cfg.get("dpi", 150),
        "savefig.bbox": "tight",
    })

    # Aggregate by sample_id (transfer rows only)
    ss_by_sample: dict[str, list[float]] = {}
    io_by_sample: dict[str, list[float]] = {}
    for r in rows:
        if r.get("type") != "transfer":
            continue
        sid = r["sample_id"]
        ss = r.get("ss_min_mv_dec")
        io = r.get("ion_ioff_log10")
        if ss is not None:
            ss_by_sample.setdefault(sid, []).append(ss)
        if io is not None:
            io_by_sample.setdefault(sid, []).append(io)

    # Best per sample: min SS (lower is better), max Ion/Ioff (higher is better)
    ss_agg: dict[str, float] = {sid: min(vals) for sid, vals in ss_by_sample.items()}
    ion_ioff_agg: dict[str, float] = {sid: max(vals) for sid, vals in io_by_sample.items()}

    # ── SS bar chart (grouped by sample_id, best[min] SS per sample) ──
    if ss_agg:
        ss_labels = list(ss_agg.keys())
        ss_values = list(ss_agg.values())
        fig1, ax1 = plt.subplots(figsize=(max(6, len(ss_labels) * 1.5), 5))
        bars = ax1.bar(range(len(ss_labels)), ss_values, color=colors[0])
        ax1.set_xticks(range(len(ss_labels)))
        ax1.set_xticklabels(ss_labels, rotation=30, ha="right", fontsize=9)
        ax1.set_ylabel("SS (mV/dec)")
        ax1.set_title("Subthreshold Swing — Best per Sample (Min SS)")
        for bar, val in zip(bars, ss_values):
            ax1.text(bar.get_x() + bar.get_width()/2, bar.get_height() + 5,
                     f"{val:.0f}", ha="center", fontsize=8)
        ax1.grid(True, alpha=0.3, axis="y")
        fig1.tight_layout()
        fig1.savefig(prepare_write_path(output_dir / "batch_ss.png"))
        plt.close(fig1)

    # ── Ion/Ioff bar chart (grouped by sample_id, best[max] per sample) ─
    if ion_ioff_agg:
        io_labels = list(ion_ioff_agg.keys())
        io_values = list(ion_ioff_agg.values())
        fig2, ax2 = plt.subplots(figsize=(max(6, len(io_labels) * 1.5), 5))
        bars2 = ax2.bar(range(len(io_labels)), io_values, color=colors[1])
        ax2.set_xticks(range(len(io_labels)))
        ax2.set_xticklabels(io_labels, rotation=30, ha="right", fontsize=9)
        ax2.set_ylabel("log10(Ion/Ioff)")
        ax2.set_title("Ion/Ioff Ratio — Best per Sample (Max Ion/Ioff)")
        for bar, val in zip(bars2, io_values):
            if val is not None:
                ax2.text(bar.get_x() + bar.get_width()/2, bar.get_height() + 0.1,
                         f"{val:.1f}", ha="center", fontsize=8)
        ax2.grid(True, alpha=0.3, axis="y")
        fig2.tight_layout()
        fig2.savefig(prepare_write_path(output_dir / "batch_ion_ioff.png"))
        plt.close(fig2)

    # Per-sample metric distributions. For N<5, a dot plot avoids a misleading histogram.
    metric_specs = {
        "vth_v_mean": "Vth (V)", "ss_min_mv_dec": "SS (mV/dec)",
        "ion_ioff_log10": "log10(Ion/Ioff)", "mobility_cm2_vs_max": "μFE (cm²/V·s)",
        "gm_max": "gm,max (S)", "gds_max": "gds,max (S)",
    }

    # Per-metric box charts grouped by sample.  Boxes show the distribution,
    # whiskers span min→max, orange markers show the mean, and the in-plot label
    # reports mean/stddev/N so it remains visible inside the axes.
    bar_dir = output_dir / "metric_bars"
    safe_mkdir(bar_dir, parents=True, exist_ok=True)
    for key, label in metric_specs.items():
        sample_labels: list[str] = []
        distributions: list[np.ndarray] = []
        for sample_id, sample_rows in sorted(by_sample.items()):
            values = np.asarray([
                r[key] for r in sample_rows
                if isinstance(r.get(key), (int, float)) and np.isfinite(r[key])
            ], dtype=float)
            if not len(values):
                continue
            sample_labels.append(str(sample_id))
            distributions.append(values)
        if not distributions:
            continue
        fig, ax = plt.subplots(figsize=(max(6, len(sample_labels) * 1.5), 5))
        positions = np.arange(1, len(distributions) + 1)
        bp = ax.boxplot(
            distributions,
            positions=positions,
            widths=0.55,
            whis=(0, 100),
            patch_artist=True,
            showmeans=True,
            meanprops={"marker": "D", "markerfacecolor": colors[1], "markeredgecolor": "black", "markersize": 5},
            medianprops={"color": "black", "linewidth": 1.2},
        )
        for box in bp["boxes"]:
            box.set(facecolor=colors[0], alpha=0.45)
        ax.set_xticks(positions)
        ax.set_xticklabels(sample_labels, rotation=30, ha="right", fontsize=9)
        ax.set_ylabel(label)
        ax.set_title(f"{label} — Distribution by Sample (whiskers=min–max)")
        for x_pos, values in zip(positions, distributions):
            mean = float(np.mean(values))
            std = float(np.std(values, ddof=1)) if len(values) > 1 else 0.0
            ax.text(
                x_pos,
                mean,
                f"μ={mean:.3g}\nσ={std:.2g}\nN={len(values)}",
                ha="center",
                va="center",
                fontsize=7,
                bbox={"boxstyle": "round,pad=0.2", "facecolor": "white", "alpha": 0.75, "edgecolor": "none"},
            )
        ax.grid(True, alpha=0.3, axis="y")
        fig.tight_layout()
        fig.savefig(prepare_write_path(bar_dir / f"{key}_by_sample.png"))
        plt.close(fig)

    # Dedicated SS comparison: minimum and average SS side-by-side per sample.
    ss_labels: list[str] = []
    ss_min_values: list[float] = []
    ss_mean_values: list[float] = []
    ss_counts: list[int] = []
    for sample_id, sample_rows in sorted(by_sample.items()):
        values = np.asarray([
            r["ss_min_mv_dec"] for r in sample_rows
            if isinstance(r.get("ss_min_mv_dec"), (int, float)) and np.isfinite(r["ss_min_mv_dec"])
        ], dtype=float)
        if not len(values):
            continue
        ss_labels.append(str(sample_id))
        ss_min_values.append(float(np.min(values)))
        ss_mean_values.append(float(np.mean(values)))
        ss_counts.append(int(len(values)))
    if ss_labels:
        fig, ax = plt.subplots(figsize=(max(6, len(ss_labels) * 1.7), 5))
        x = np.arange(len(ss_labels))
        width = 0.38
        min_bars = ax.bar(x - width / 2, ss_min_values, width, label="Min SS", color=colors[0])
        mean_bars = ax.bar(x + width / 2, ss_mean_values, width, label="Avg SS", color=colors[1])
        ax.set_xticks(x)
        ax.set_xticklabels(ss_labels, rotation=30, ha="right", fontsize=9)
        ax.set_ylabel("SS (mV/dec)")
        ax.set_title("Subthreshold Swing — Min vs Average per Sample")
        for bars, values in ((min_bars, ss_min_values), (mean_bars, ss_mean_values)):
            for bar, value in zip(bars, values):
                ax.text(
                    bar.get_x() + bar.get_width() / 2,
                    bar.get_height() / 2,
                    f"{value:.0f}",
                    ha="center",
                    va="center",
                    fontsize=8,
                    color="white",
                    fontweight="bold",
                )
        for index, count in enumerate(ss_counts):
            ax.text(index, max(ss_min_values[index], ss_mean_values[index]), f"N={count}", ha="center", va="bottom", fontsize=8)
        ax.legend()
        ax.grid(True, alpha=0.3, axis="y")
        fig.tight_layout()
        fig.savefig(prepare_write_path(output_dir / "batch_ss_min_avg.png"))
        fig.savefig(prepare_write_path(bar_dir / "ss_min_avg_by_sample.png"))
        plt.close(fig)

    histogram_dir = output_dir / "histograms"
    safe_mkdir(histogram_dir, parents=True, exist_ok=True)
    for sample_id, sample_rows in sorted(by_sample.items()):
        safe_sample = "".join(c if c.isalnum() or c in "-_" else "_" for c in str(sample_id))
        for key, label in metric_specs.items():
            values = np.asarray([r[key] for r in sample_rows
                                 if isinstance(r.get(key), (int, float)) and np.isfinite(r[key])], dtype=float)
            if not len(values):
                continue
            fig, ax = plt.subplots(figsize=(7, 5))
            if len(values) >= 5:
                ax.hist(values, bins="auto", color=colors[0], edgecolor="black", alpha=0.8)
                ax.set_ylabel("Count")
            else:
                ax.scatter(values, np.zeros_like(values), color=colors[0], s=50)
                ax.set_yticks([]); ax.set_ylabel("Individual measurements")
            mean = float(np.mean(values)); median = float(np.median(values))
            std = float(np.std(values, ddof=1)) if len(values) > 1 else None
            ax.axvline(mean, color=colors[1], linestyle="--", label=f"Mean={mean:.3g}")
            ax.axvline(median, color=colors[2], linestyle=":", label=f"Median={median:.3g}")
            ax.set_xlabel(label); ax.set_title(f"{sample_id}: {label} (N={len(values)}, SD={std:.3g})" if std is not None else f"{sample_id}: {label} (N=1)")
            ax.grid(True, alpha=0.25); ax.legend(fontsize=8); fig.tight_layout()
            fig.savefig(prepare_write_path(
                histogram_dir / f"{safe_sample}_{key}.png"
            )); plt.close(fig)

    LOGGER.info("Batch statistics plots written to %s", output_dir)


def generate_batch_summary(
    output_root: Path,
    config: dict[str, Any] | None = None,
) -> Path:
    """Full batch summary: collect, write Excel+CSV, plot statistics.

    Returns path to the summary directory.
    """
    summary_dir = output_root / "batch_summary"
    safe_mkdir(summary_dir, parents=True, exist_ok=True)

    preferred_rows, _ = _collect_preferred_metrics(output_root)
    extended_rows = collect_extended_metrics(output_root, preferred_rows)
    rows, by_sample = _derive_master_metrics(extended_rows, preferred_rows)

    if not rows:
        LOGGER.warning("No device metrics found in %s", output_root)
        return summary_dir

    write_summary_excel(rows, summary_dir / "master_summary.xlsx")
    write_summary_csv(rows, by_sample, summary_dir / "master_summary.csv")
    write_extended_summary_excel(
        extended_rows, summary_dir / "extended_master_summary.xlsx",
    )
    write_extended_summary_csv(
        extended_rows, summary_dir / "extended_master_summary.csv",
    )
    if (config or {}).get("execution", {}).get("generate_batch_plots", False):
        plot_batch_statistics(rows, by_sample, summary_dir, config)
    else:
        for png in safe_path(summary_dir).glob("*.png"):
            safe_unlink(png)
        for name in ("metric_bars", "histograms"):
            if safe_path(summary_dir / name).exists():
                safe_rmtree(summary_dir / name)

    LOGGER.info("Batch summary complete: %d devices across %d samples",
                len(rows), len(by_sample))
    return summary_dir
