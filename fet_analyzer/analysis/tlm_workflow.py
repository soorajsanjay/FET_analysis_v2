"""Dedicated cross-length TLM analysis grouped by sample label and TLM ID."""
from __future__ import annotations
import csv
import re
from datetime import datetime, timezone
from collections import defaultdict
from pathlib import Path
from typing import Any
import numpy as np
from fet_analyzer.analysis.ion_bias import field_to_voltage, resolve_ion_bias
from fet_analyzer.analysis.numerics import nearest_abs_current_at_voltage
from fet_analyzer.analysis.tlm import (
    extract_tlm, _interpolate_current, _interpolate_prepared, _prepare_interpolation,
)
from fet_analyzer.utils.logging import LOGGER
from fet_analyzer.reports.html_report import write_html_report
from fet_analyzer.temp_workspace import temporary_directory
from fet_analyzer.path_utils import (
    prepare_write_path, safe_mkdir, safe_path, safe_rmtree, safe_unlink,
)

def _identity(record: dict[str, Any]) -> tuple[str, str, float] | None:
    name = record["source_file"].stem
    match = re.search(r"(TLM\d*)_(\d+(?:\.\d+)?)\s*(?:um|µm|μm)", name, re.I)
    if not match: return None
    sample = record["classification"].get("filename_info", {}).get("sample_label", "unknown")
    return str(sample), match.group(1).upper(), float(match.group(2))

def _common_film_thickness(records: list[tuple[float, dict[str, Any]]]) -> float | None:
    values = {float(record["device_params"]["film_thickness_nm"])
              for _, record in records if record["device_params"].get("film_thickness_nm") is not None}
    if len(values) > 1:
        LOGGER.warning("TLM film resistivity unavailable: devices use different film thicknesses")
        return None
    return next(iter(values)) if values else None

def _fit(points: list[tuple[float, float]], width: float, r2_warning: float,
         tlm_cfg: dict[str, Any] | None = None,
         film_thickness_nm: float | None = None) -> dict[str, Any]:
    cfg = tlm_cfg or {}
    return extract_tlm(
        points, width_um=width, r2_warning=r2_warning,
        min_length_span_ratio=float(cfg.get("acceptance_min_length_span_ratio", 2.0)),
        max_relative_uncertainty=float(cfg.get("acceptance_max_relative_uncertainty", 1.0)),
        max_leverage=float(cfg.get("acceptance_max_leverage", 0.85)),
        max_residual_curvature=float(cfg.get("acceptance_max_residual_curvature", 0.8)),
        max_abs_lt_to_lmin_ratio=float(cfg.get("acceptance_max_lt_to_lmin_ratio", 10.0)),
        film_thickness_nm=film_thickness_nm,
    )

def _output_resistance(record: dict[str, Any], vd_max: float) -> dict[float | str, float]:
    result = {}
    c = record["classification"]; sweep = c.get("sweep_variable", "Vd"); drain = c.get("drain_current_raw_column", "Id")
    from fet_analyzer.analysis.segmentation import group_by_bias
    for segments in group_by_bias(record["segments"]).values():
        forward = [segment for segment in segments if segment.direction == "forward"]
        candidates = forward or (segments[:1] if len(segments) == 1 else [])
        if not candidates:
            continue
        segment = candidates[0]
        vd = np.asarray(segment.data.get(sweep, []), float); current = np.asarray(segment.data.get(drain, []), float)
        valid = np.isfinite(vd) & np.isfinite(current) & (np.abs(vd) <= vd_max)
        if valid.sum() < 3: continue
        slope = float(np.polyfit(vd[valid], current[valid], 1)[0])
        if abs(slope) > 1e-15:
            if c.get("is_ungated_tlm"):
                key = "ungated"
            else:
                key = round(segment.bias_level, 6) if segment.bias_level is not None else "floating_gate"
            result[key] = 1.0 / abs(slope)
    return result

def _transfer_rows(records: list[tuple[float, dict[str, Any]]], config: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []; factor = float(config.get("transfer", {}).get("ioff_leakage_factor", 3.0))
    r2_warn = float(config.get("advanced", {}).get("tlm_r2_warning", 0.9)); width = float(records[0][1]["device_params"].get("channel_width_um", 100))
    film_thickness_nm = _common_film_thickness(records)
    by_vds: dict[float, list[tuple[float, Any, dict[str, Any], Any, Any]]] = defaultdict(list)
    for length, record in records:
        for segment in record["segments"]:
            if segment.direction == "forward" and segment.bias_level is not None:
                c = record["classification"]
                drain = c.get("drain_current_raw_column", "Id")
                gate = c.get("gate_current_column") or c.get("gate_leakage_column")
                sweep = np.asarray(segment.sweep_values, float)
                drain_prepared = _prepare_interpolation(
                    sweep, np.asarray(segment.data.get(drain, []), float)
                )
                gate_prepared = (
                    _prepare_interpolation(sweep, np.asarray(segment.data[gate], float))
                    if gate and segment.data.get(gate) else None
                )
                by_vds[round(float(segment.bias_level), 6)].append(
                    (length, segment, record, drain_prepared, gate_prepared)
                )
    for vds, members in sorted(by_vds.items()):
        if len({x[0] for x in members}) < 3: continue
        grids = [{round(float(v), 9) for v in seg.sweep_values if np.isfinite(v)}
                 for _, seg, _, _, _ in members]
        common_vg = sorted(set.intersection(*grids)) if grids else []
        for vg in common_vg:
            points = []; excluded = 0
            for length, segment, record, drain_prepared, gate_prepared in members:
                id_at = _interpolate_prepared(drain_prepared, vg)
                ig_at = _interpolate_prepared(gate_prepared, vg) if gate_prepared else None
                if id_at is None or abs(id_at) <= 1e-15 or (ig_at is not None and abs(id_at) <= factor * abs(ig_at)):
                    excluded += 1; continue
                points.append((length, abs(vds) / abs(id_at)))
            if len(points) < 3: continue
            fit = _fit(points, width, r2_warn, config.get("tlm", {}), film_thickness_nm)
            rows.append({"vds_v": vds, "vg_v": vg, "n_lengths": len({point[0] for point in points}), "excluded": excluded,
                         "expected_lengths": len({item[0] for item in members}),
                         "points": points, **fit})
    return rows


def _tlm_role(record: dict[str, Any]) -> str:
    classification = record.get("classification", {})
    role = classification.get("tlm_role")
    if role:
        return str(role)
    sweep = str(classification.get("sweep_variable", "")).lower()
    if classification.get("is_ungated_tlm"):
        return "ungated_ltlm"
    return "gated_transfer" if sweep in {"vg", "vbg"} else "output_sanity"


def _output_fit_rows(
    records: list[tuple[float, dict[str, Any]]], config: dict[str, Any],
) -> list[dict[str, Any]]:
    if not records:
        return []
    width = float(records[0][1]["device_params"].get("channel_width_um", 100))
    vd_max = float(config.get("tlm", {}).get("tlm_output_vd_max", 0.1))
    r2_warning = float(config.get("advanced", {}).get("tlm_r2_warning", 0.9))
    thickness = _common_film_thickness(records)
    by_bias: dict[float | str, list[tuple[float, float]]] = defaultdict(list)
    for length, record in records:
        for bias, resistance in _output_resistance(record, vd_max).items():
            by_bias[bias].append((length, resistance))
    return [
        {
            "gate_bias": bias,
            "n_lengths": len({point[0] for point in points}),
            "expected_lengths": len({length for length, _ in records}),
            "points": points,
            **_fit(points, width, r2_warning, config.get("tlm", {}), thickness),
        }
        for bias, points in sorted(by_bias.items(), key=lambda item: str(item[0]))
        if len({point[0] for point in points}) >= 3
    ]

def _length_metrics(records: list[tuple[float, dict[str, Any]]]) -> list[dict[str, Any]]:
    rows = []
    for length, record in records:
        m = record.get("metrics", {}); summary = m.get("summary", {})
        base = {"source_file": record["source_file"].name, "lch_um": length}
        preferred = (m.get("device_summary") or {}).get("preferred", {})
        values = {"vth_v": preferred.get("vth_v", summary.get("vth_v_mean")),
                  "ss_mv_dec": preferred.get("ss_min_mv_dec", summary.get("ss_mv_dec_min")),
                  "mobility_cm2_vs": preferred.get("mobility_cm2_vs", summary.get("mobility_cm2_vs_max")),
                  "hysteresis_v": preferred.get("hysteresis_abs_v", summary.get("hysteresis_v_max")),
                  "ion_ioff_log10": preferred.get("ion_ioff_log10"),
                  "ion_a": preferred.get("ion_configured_a"),
                  "ioff_a": preferred.get("ioff_configured_a"),
                  "gm_max_s": preferred.get("gm_max_s", m.get("gm_max_s"))}
        rows.append({**base, **values})
    return rows

def _ion_overdrive_values(records: list[tuple[float, dict[str, Any]]], config: dict[str, Any]) -> dict[str, dict[str, Any]]:
    """Return maximum, configured-overdrive, and maximum-common-overdrive Ion."""
    prepared = []
    for length, record in records:
        forward = [s for s in record["segments"] if s.direction == "forward" and s.bias_level is not None]
        if not forward: continue
        segment = max(forward, key=lambda s: abs(float(s.bias_level)))
        vth_candidates = list(record.get("metrics", {}).get("per_bias_vth", {}).values())
        vth_result = min(vth_candidates, key=lambda item: abs(abs(float(item.get("vds_v", 0))) - abs(float(segment.bias_level)))) if vth_candidates else {}
        vth = vth_result.get("vth_v")
        if vth is None: continue
        c = record["classification"]; drain = c.get("drain_current_raw_column", "Id")
        vg = np.asarray(segment.sweep_values, float); current = np.asarray(segment.data.get(drain, []), float)
        finite = np.isfinite(vg) & np.isfinite(current); vg, current = vg[finite], current[finite]
        if len(vg) < 2: continue
        prepared.append((record, float(vth), vg, current, float(np.min(vg)-vth), float(np.max(vg)-vth)))
    if not prepared: return {}
    common_low = max(item[4] for item in prepared); common_high = min(item[5] for item in prepared)
    polarity = str(prepared[0][0]["device_params"].get("polarity", "p")).lower()
    common_vov = common_low if polarity == "p" else common_high
    result = {}
    for record, vth, vg, current, _, _ in prepared:
        ion_bias = resolve_ion_bias(config, record.get("device_params"))
        configured = ion_bias["overdrive_v"]
        values = {
            "ion_max_a": float(np.max(np.abs(current))),
            "configured_overdrive_v": configured,
            "configured_overdrive_input_v": ion_bias["overdrive_input_v"],
            "configured_overdrive_field_mv_cm": ion_bias["overdrive_field_mv_cm"],
            "configured_oxide_thickness_nm": ion_bias["oxide_thickness_nm"],
            "configured_overdrive_source": ion_bias["overdrive_source"],
            "ion_bias_warnings": "; ".join(ion_bias["warnings"]),
            "common_overdrive_v": common_vov,
        }
        if configured is not None:
            read = nearest_abs_current_at_voltage(
                vg.tolist(), current.tolist(), vth + float(configured)
            )
            actual_vov = (
                float(read["actual_v"]) - vth if read["actual_v"] is not None else None
            )
            values.update({
                "ion_fixed_overdrive_a": read["current_a"],
                "configured_overdrive_actual_v": actual_vov,
                "configured_overdrive_target_vg_v": read["requested_v"],
                "configured_overdrive_actual_vg_v": read["actual_v"],
                "configured_overdrive_delta_vg_v": read["delta_v"],
            })
        else: values["ion_fixed_overdrive_a"] = None
        read = _interpolate_current(vg, current, vth + common_vov); values["ion_common_overdrive_a"] = abs(read) if read is not None else None
        result[record["source_file"].name] = values
    return result

def _write_group(
    group_dir: Path, sample: str, tlm_id: str,
    records: list[tuple[float, dict[str, Any]]], config: dict[str, Any],
) -> tuple[Path, dict[str, Any]]:
    import openpyxl
    from openpyxl.worksheet.table import Table, TableStyleInfo
    import matplotlib; matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    safe_mkdir(group_dir, parents=True, exist_ok=True)
    keep_plots = bool(config.get("execution", {}).get("save_plots", False))
    temporary_plots: Path | None = None
    if keep_plots:
        plots = group_dir / "plots"; safe_mkdir(plots, parents=True, exist_ok=True)
        for pattern in ("output_tlm_*.png", "ungated_ltlm_*.png", "transfer_tlm_*.png", "*_vs_vg.png", "*_vs_lch.png"):
            for old in safe_path(plots).glob(pattern): safe_unlink(old)
    else:
        if safe_path(group_dir / "plots").exists(): safe_rmtree(group_dir / "plots")
        temporary_plots = temporary_directory(config, "tlm_", group_dir); plots = temporary_plots
    width = float(records[0][1]["device_params"].get("channel_width_um", 100)); vd_max = float(config.get("tlm", {}).get("tlm_output_vd_max", 0.1)); r2w = float(config.get("advanced", {}).get("tlm_r2_warning", .9))
    film_thickness_nm = _common_film_thickness(records)
    roles: dict[str, list[tuple[float, dict[str, Any]]]] = defaultdict(list)
    for item in records:
        roles[_tlm_role(item[1])].append(item)
    transfer_records = roles["gated_transfer"]
    ungated_records = roles["ungated_ltlm"]
    output_sanity_records = roles["output_sanity"]
    ungated_fits = _output_fit_rows(ungated_records, config)
    output_sanity_fits = _output_fit_rows(output_sanity_records, config)
    transfer_fits = _transfer_rows(transfer_records, config) if transfer_records else []
    length_rows = _length_metrics(transfer_records)
    ion_definitions = _ion_overdrive_values(transfer_records, config)
    for row in length_rows:
        row.update(ion_definitions.get(row["source_file"], {}))
        ioff = row.get("ioff_a")
        for ion_key, ratio_key in [("ion_max_a","ion_ioff_max_log10"),
                                   ("ion_fixed_overdrive_a","ion_ioff_fixed_overdrive_log10"),
                                   ("ion_common_overdrive_a","ion_ioff_common_overdrive_log10")]:
            ion = row.get(ion_key)
            row[ratio_key] = float(np.log10(ion / ioff)) if ion is not None and ioff is not None and ion > 0 and ioff > 0 else None

    wb = openpyxl.Workbook(); wb.active.title = "README"; wb.active.append([f"TLM analysis: {sample} / {tlm_id}"]); wb.active.append(["Regression", "Ordinary least squares with a free intercept"]); wb.active.append(["Ungated LTLM", "Low-Vd resistance is fitted without requiring a gate-bias column"]); wb.active.append(["Width (um)", width]); wb.active.append(["Low-Vd fit limit (V)", vd_max])
    def table_sheet(name, headers, rows):
        seen = {}
        unique_headers = []
        for index, header in enumerate(headers, 1):
            base = str(header).strip() if header is not None else ""
            if not base:
                base = f"Column{index}"
            count = seen.get(base, 0) + 1
            seen[base] = count
            unique_headers.append(base if count == 1 else f"{base}_{count}")
        ws = wb.create_sheet(name); ws.append(unique_headers)
        for row in rows: ws.append(row)
        ref = f"A1:{openpyxl.utils.get_column_letter(len(unique_headers))}{max(1, len(rows) + 1)}"
        ws.freeze_panes="A2"; ws.auto_filter.ref=ref
        if rows:
            t=Table(displayName="Tbl"+re.sub(r"\W", "", name), ref=ref); t.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True); ws.add_table(t)
        return ws
    table_sheet("Source Files", ["sample_label","tlm_id","lch_um","filename","tlm_role","width_um","film_thickness_nm"], [[sample,tlm_id,l,r["source_file"].name,_tlm_role(r),r["device_params"].get("channel_width_um",100),r["device_params"].get("film_thickness_nm")] for l,r in records])
    tr_rows=[]
    diagnostic_rows=[]
    for f in transfer_fits:
        slope_std=f.get("slope_std"); intercept_std=f.get("intercept_std")
        tr_rows.append([f["vds_v"],f["vg_v"],f["n_lengths"],f["excluded"],f.get("rsh_ohm_sq"),f.get("rcw_ohm_um"),f.get("rc_ohm"),f.get("lt_um"),f.get("rhoc_ohm_cm2"),f.get("film_thickness_nm"),f.get("rho_film_ohm_cm"),f.get("r2"),f.get("slope_ohm_per_um"),slope_std,1.96*slope_std if slope_std is not None else None,f.get("intercept_ohm"),intercept_std,1.96*intercept_std if intercept_std is not None else None,f.get("acceptance_status"),"; ".join(f.get("acceptance_reasons",[])),f.get("length_span_ratio"),f.get("slope_relative_uncertainty"),f.get("intercept_relative_uncertainty"),f.get("max_observation_leverage"),f.get("residual_curvature"),f.get("lt_to_lmin_ratio"),"; ".join(f.get("warnings",[]))])
        for length, resistance in f.get("points", []):
            measured=resistance*width; predicted=f.get("slope_ohm_per_um",0)*length+f.get("intercept_ohm",0)
            diagnostic_rows.append([f["vds_v"],f["vg_v"],length,measured,predicted,measured-predicted])
    table_sheet("TLM Parameters vs Vg", ["vds_v","vg_v","n_lengths","excluded_leakage","rsh_ohm_sq","rcw_ohm_um","rc_ohm","lt_um","rhoc_ohm_cm2","film_thickness_nm","rho_film_ohm_cm","r2","slope_ohm_per_um","slope_std","slope_95ci","intercept_ohm_um","intercept_std","intercept_95ci","acceptance_status","acceptance_reasons","length_span_ratio","slope_relative_uncertainty","intercept_relative_uncertainty","max_observation_leverage","residual_curvature","lt_to_lmin_ratio","warnings"], tr_rows)
    table_sheet("Transfer Fit Diagnostics", ["vds_v","vg_v","lch_um","measured_rtotal_w_ohm_um","predicted_rtotal_w_ohm_um","residual_ohm_um"], diagnostic_rows)
    headers=["source_file","lch_um","vth_v","ss_mv_dec","mobility_cm2_vs","hysteresis_v","ion_ioff_log10","ion_a","ioff_a","gm_max_s",
             "ion_max_a","ion_fixed_overdrive_a","configured_overdrive_v","configured_overdrive_input_v",
             "configured_overdrive_field_mv_cm","configured_oxide_thickness_nm","configured_overdrive_source",
             "configured_overdrive_actual_v","configured_overdrive_target_vg_v",
             "configured_overdrive_actual_vg_v","configured_overdrive_delta_vg_v",
             "ion_bias_warnings","ion_common_overdrive_a","common_overdrive_v"]
    headers.extend(["ion_ioff_max_log10","ion_ioff_fixed_overdrive_log10","ion_ioff_common_overdrive_log10"])
    table_sheet("Metrics vs Lch", headers, [[r.get(h) for h in headers] for r in length_rows])

    output_headers = ["gate_bias_v","lch_um","rtotal_ohm","rtotal_w_ohm_um","rsh_ohm_sq","rcw_ohm_um","rc_ohm","lt_um","rhoc_ohm_cm2","film_thickness_nm","rho_film_ohm_cm","r2","acceptance_status","acceptance_reasons","length_span_ratio","slope_relative_uncertainty","intercept_relative_uncertainty","max_observation_leverage","residual_curvature","lt_to_lmin_ratio","warnings"]
    def output_rows(fits):
        return [[f["gate_bias"],l,rt,rt*width,f.get("rsh_ohm_sq"),f.get("rcw_ohm_um"),f.get("rc_ohm"),f.get("lt_um"),f.get("rhoc_ohm_cm2"),f.get("film_thickness_nm"),f.get("rho_film_ohm_cm"),f.get("r2"),f.get("acceptance_status"),"; ".join(f.get("acceptance_reasons",[])),f.get("length_span_ratio"),f.get("slope_relative_uncertainty"),f.get("intercept_relative_uncertainty"),f.get("max_observation_leverage"),f.get("residual_curvature"),f.get("lt_to_lmin_ratio"),"; ".join(f.get("warnings",[]))] for f in fits for l,rt in f["points"]]
    table_sheet("Ungated LTLM Fits", output_headers, output_rows(ungated_fits))
    table_sheet("Output Sanity Check", output_headers, output_rows(output_sanity_fits))

    html_plots: list[Path] = []
    for fit_index, f in enumerate(ungated_fits):
        x=np.array([p[0] for p in f["points"]]); y=np.array([p[1]*width for p in f["points"]]); fig,ax=plt.subplots(); ax.scatter(x,y,label="Measured")
        if f.get("slope_ohm_per_um") is not None: ax.plot(x,f["slope_ohm_per_um"]*x+f["intercept_ohm"],label=f"Fit R²={f.get('r2',float('nan')):.3g}")
        ax.set(xlabel="Lch (µm)",ylabel="Rtotal×W (Ω·µm)",title=f"{sample} {tlm_id}: Ungated LTLM"); ax.legend(); ax.grid(alpha=.3); fig.tight_layout()
        plot_path = plots/f"ungated_ltlm_{fit_index}.png"; fig.savefig(prepare_write_path(plot_path)); plt.close(fig); html_plots.append(plot_path)
    if transfer_fits:
        # Representative raw TLM fits spanning the shared Vg range.
        selected_indices = sorted(set(np.linspace(0, len(transfer_fits) - 1, min(6, len(transfer_fits)), dtype=int)))
        for plot_index, fit_index in enumerate(selected_indices):
            f = transfer_fits[fit_index]
            x = np.asarray([point[0] for point in f["points"]], dtype=float)
            y = np.asarray([point[1] * width for point in f["points"]], dtype=float)
            fig, ax = plt.subplots()
            ax.scatter(x, y, label="Measured")
            if f.get("slope_ohm_per_um") is not None:
                order = np.argsort(x)
                prediction = f["slope_ohm_per_um"] * x[order] + f["intercept_ohm"]
                ax.plot(x[order], prediction,
                        label=f"Fit R²={f.get('r2', float('nan')):.3g}")
                if f.get("slope_std") is not None and f.get("intercept_std") is not None:
                    uncertainty = 1.96 * np.sqrt((x[order] * f["slope_std"]) ** 2 + f["intercept_std"] ** 2)
                    ax.fill_between(x[order], prediction - uncertainty, prediction + uncertainty,
                                    alpha=.18, label="Approx. 95% CI")
            ax.set(xlabel="Lch (µm)", ylabel="Rtotal×W (Ω·µm)",
                   title=f"{sample} {tlm_id}: transfer TLM, Vg={f['vg_v']:g} V, Vds={f['vds_v']:g} V")
            ax.grid(alpha=.3); ax.legend(); fig.tight_layout()
            plot_path = plots / f"transfer_tlm_fit_{plot_index}_vg_{f['vg_v']:g}.png"; fig.savefig(prepare_write_path(plot_path)); html_plots.append(plot_path)
            plt.close(fig)
            predicted = f["slope_ohm_per_um"] * x + f["intercept_ohm"]
            residual = y - predicted
            fig, ax = plt.subplots()
            ax.axhline(0, color="0.35", linewidth=1)
            ax.scatter(x, residual)
            ax.set(xlabel="Lch (µm)", ylabel="Residual Rtotal×W (Ω·µm)",
                   title=f"{sample} {tlm_id}: TLM residuals, Vg={f['vg_v']:g} V")
            ax.grid(alpha=.3); fig.tight_layout()
            plot_path = plots / f"transfer_tlm_residuals_{plot_index}_vg_{f['vg_v']:g}.png"; fig.savefig(prepare_write_path(plot_path)); html_plots.append(plot_path)
            plt.close(fig)
        for key,label in [("rc_ohm","Rc (Ω)"),("rcw_ohm_um","RcW (Ω·µm)"),("rsh_ohm_sq","Rsh (Ω/sq)"),("rho_film_ohm_cm","Film resistivity (Ω·cm)"),("lt_um","LT (µm)"),("rhoc_ohm_cm2","ρc (Ω·cm²)"),("r2","Fit R²")]:
            fig,ax=plt.subplots()
            for vds in sorted({f["vds_v"] for f in transfer_fits}):
                subset=[f for f in transfer_fits if f["vds_v"]==vds and f.get(key) is not None]; ax.plot([f["vg_v"] for f in subset],[f[key] for f in subset],marker=".",label=f"Vds={vds:g} V")
            ax.set(xlabel="Vg (V)",ylabel=label,title=f"{sample} {tlm_id}: {label} vs Vg"); ax.grid(alpha=.3); ax.legend(); fig.tight_layout(); plot_path=plots/f"{key}_vs_vg.png"; fig.savefig(prepare_write_path(plot_path)); plt.close(fig); html_plots.append(plot_path)
    for key,label,logy in [("ion_ioff_log10","log10(Ion/Ioff)",False),("ion_ioff_max_log10","log10(Ion/Ioff), maximum current",False),("ion_ioff_fixed_overdrive_log10","log10(Ion/Ioff), configured overdrive",False),("ion_ioff_common_overdrive_log10","log10(Ion/Ioff), common overdrive",False),("ss_mv_dec","SS (mV/dec)",False),("mobility_cm2_vs","Mobility (cm²/V·s)",False),("vth_v","Vth (V)",False),("ion_a","Ion max/min method (A)",True),("ion_max_a","Ion maximum measured (A)",True),("ion_fixed_overdrive_a","Ion configured overdrive (A)",True),("ion_common_overdrive_a","Ion maximum common overdrive (A)",True),("ioff_a","Ioff (A)",True),("hysteresis_v","Hysteresis (V)",False),("gm_max_s","gm,max (S)",True)]:
        pts=[r for r in length_rows if r.get(key) is not None]
        if not pts: continue
        fig,ax=plt.subplots(); ax.scatter([r["lch_um"] for r in pts],[r[key] for r in pts],marker="o");
        if logy: ax.set_yscale("log")
        ax.set(xlabel="Lch (µm)",ylabel=label,title=f"{sample} {tlm_id}: {label} vs Lch"); ax.grid(alpha=.3); fig.tight_layout(); plot_path=plots/f"{key}_vs_lch.png"; fig.savefig(prepare_write_path(plot_path)); plt.close(fig); html_plots.append(plot_path)
    path=group_dir/f"TLM_{sample}_{tlm_id}.xlsx"; wb.save(prepare_write_path(path)); wb.close()
    report=group_dir/f"TLM_{sample}_{tlm_id}_report.html"
    tlm_cfg = config.get("tlm", {})
    representative_ion_bias = resolve_ion_bias(config, records[0][1].get("device_params"))
    settings = [
        ("Transfer read mode", tlm_cfg.get("transfer_read_mode")),
        ("Transfer read Vg (V)", tlm_cfg.get("transfer_read_vg_v")),
        ("Transfer gate field (MV/cm)", tlm_cfg.get("transfer_gate_field_mv_cm")),
        ("Transfer overdrive, resolved (V)", representative_ion_bias["overdrive_v"]),
        ("Transfer overdrive input (V)", representative_ion_bias["overdrive_input_v"]),
        ("Transfer overdrive field (MV/cm)", representative_ion_bias["overdrive_field_mv_cm"]),
        ("TLM transfer overdrive field (MV/cm)", tlm_cfg.get("transfer_overdrive_field_mv_cm")),
        ("Overdrive source", representative_ion_bias["overdrive_source"]),
        ("Oxide thickness for conversion (nm)", representative_ion_bias["oxide_thickness_nm"]),
        ("Semiconductor film thickness (nm)", film_thickness_nm),
        ("Transfer read Vds (V)", tlm_cfg.get("transfer_read_vds_v")),
        ("Output resistance fit |Vd| max (V)", tlm_cfg.get("tlm_output_vd_max", 0.1)),
        ("TLM R² warning threshold", config.get("advanced", {}).get("tlm_r2_warning", 0.9)),
    ]
    settings_md = "\n".join(
        f"| {label} | {'not configured' if value is None else value} |"
        for label, value in settings
    )
    result_rows = []
    for label, fit in (("Gated transfer", _best_fit(transfer_fits)), ("Ungated LTLM", _best_fit(ungated_fits))):
        if fit:
            result_rows.append(
                f"| {label} | {fit.get('rsh_ohm_sq')} | {fit.get('rho_film_ohm_cm')} | "
                f"{fit.get('rhoc_ohm_cm2')} | {fit.get('r2')} | {fit.get('acceptance_status')} |"
            )
    result_table = (
        "## Representative TLM Results\n\n"
        "| Mode | Rsh (ohm/sq) | Film resistivity (ohm cm) | Contact resistivity (ohm cm2) | R2 | Status |\n"
        "|---|---:|---:|---:|---:|---|\n" + "\n".join(result_rows) + "\n\n"
        if result_rows else ""
    )
    report_lines = (
        f"# TLM {sample} / {tlm_id}\n\n**Files:** {len(records)}\n**Width:** {width:g} µm\n\n"
        f"{result_table}## Analysis Configuration Used\n\n| Setting | Effective value |\n|---|---|\n{settings_md}\n\n## Plots\n"
    ).splitlines()
    write_html_report(report, f"TLM {sample} / {tlm_id}", report_lines, html_plots)
    if temporary_plots is not None:
        safe_rmtree(temporary_plots, ignore_errors=True)
    return path, _summary_row(
        sample, tlm_id, records, config,
        ungated_fits=ungated_fits, output_fits=output_sanity_fits, transfer_fits=transfer_fits,
        length_rows=length_rows,
    )

def _best_fit(fits: list[dict[str, Any]]) -> dict[str, Any]:
    valid = [fit for fit in fits if isinstance(fit.get("r2"), (int, float)) and np.isfinite(fit["r2"])]
    if valid:
        rank = {"accepted": 2, "review": 1, "rejected_nonphysical": 0}
        return max(valid, key=lambda fit: (rank.get(fit.get("acceptance_status"), -1), fit["r2"]))
    return fits[0] if fits else {}


def _accepted_best_fit(fits: list[dict[str, Any]]) -> dict[str, Any]:
    accepted = [fit for fit in fits if fit.get("acceptance_status") == "accepted"]
    return _best_fit(accepted)


def _accepted_near_zero_transfer_fit(
    fits: list[dict[str, Any]], config: dict[str, Any],
) -> dict[str, Any]:
    accepted = [fit for fit in fits if fit.get("acceptance_status") == "accepted"]
    return _near_zero_transfer_fit(accepted, config)


def _index_fields(
    ungated_fits: list[dict[str, Any]], transfer_fits: list[dict[str, Any]],
    config: dict[str, Any], source_files: dict[str, str] | None = None,
) -> dict[str, Any]:
    fit = _accepted_best_fit(ungated_fits)
    source = "ungated_ltlm" if fit else ""
    condition = "accepted explicit LTLM low-Vd fit" if fit else ""
    actual_vg = None
    vds = None
    if not fit:
        fit = _accepted_near_zero_transfer_fit(transfer_fits, config)
        source = "gated_transfer_vg0" if fit else "unavailable"
        condition = "accepted common measured Vg nearest 0 V" if fit else "no accepted ungated or gated-near-zero TLM fit"
        actual_vg = fit.get("vg_v") if fit else None
        vds = fit.get("vds_v") if fit else None
    files = (source_files or {}).get(source, "")
    return {
        "index_source": source, "index_condition": condition,
        "index_actual_vg_v": actual_vg, "index_vds_v": vds,
        "index_rsh_ohm_sq": fit.get("rsh_ohm_sq"),
        "index_rc_ohm": fit.get("rc_ohm"), "index_rcw_ohm_um": fit.get("rcw_ohm_um"),
        "index_lt_um": fit.get("lt_um"), "index_rho_film_ohm_cm": fit.get("rho_film_ohm_cm"),
        "index_rhoc_ohm_cm2": fit.get("rhoc_ohm_cm2"), "index_r2": fit.get("r2"),
        "index_acceptance_status": fit.get("acceptance_status"),
        "index_acceptance_reasons": "; ".join(fit.get("acceptance_reasons", [])),
        "index_source_files": files,
    }


def _mode_summary_fields(prefix: str, fits: list[dict[str, Any]]) -> dict[str, Any]:
    best = _best_fit(fits)
    return {
        f"{prefix}_fit_count": len(fits),
        f"{prefix}_best_gate_bias_v": best.get("gate_bias"),
        f"{prefix}_best_r2": best.get("r2"),
        f"{prefix}_acceptance_status": best.get("acceptance_status"),
        f"{prefix}_acceptance_reasons": "; ".join(best.get("acceptance_reasons", [])),
        f"{prefix}_warnings": "; ".join(best.get("warnings", [])),
        f"{prefix}_rsh_ohm_sq": best.get("rsh_ohm_sq"),
        f"{prefix}_rcw_ohm_um": best.get("rcw_ohm_um"), f"{prefix}_rc_ohm": best.get("rc_ohm"),
        f"{prefix}_lt_um": best.get("lt_um"), f"{prefix}_rhoc_ohm_cm2": best.get("rhoc_ohm_cm2"),
        f"{prefix}_rho_film_ohm_cm": best.get("rho_film_ohm_cm"),
    }

def _nearest_transfer_fit(
    fits: list[dict[str, Any]], target_vg: float, config: dict[str, Any],
) -> dict[str, Any]:
    """Select the fitted common-Vg point nearest a requested gate voltage."""
    valid = [
        fit for fit in fits
        if isinstance(fit.get("vg_v"), (int, float)) and np.isfinite(fit["vg_v"])
    ]
    if not valid:
        return {}
    preferred_vds = config.get("summary", {}).get("preferred_vd_v")
    rank = {"accepted": 2, "review": 1, "rejected_nonphysical": 0}
    def key(fit: dict[str, Any]) -> tuple[float, float, int, float, float]:
        vds = fit.get("vds_v")
        vds_distance = (
            abs(abs(float(vds)) - abs(float(preferred_vds)))
            if isinstance(vds, (int, float)) and preferred_vds is not None
            else 0.0
        )
        r2 = float(fit["r2"]) if isinstance(fit.get("r2"), (int, float)) and np.isfinite(fit["r2"]) else -1.0
        return (
            abs(float(fit["vg_v"]) - float(target_vg)), vds_distance,
            -rank.get(str(fit.get("acceptance_status", "")), -1), -r2,
            float(vds) if isinstance(vds, (int, float)) and np.isfinite(vds) else float("inf"),
        )
    return min(valid, key=key)


def _near_zero_transfer_fit(
    fits: list[dict[str, Any]], config: dict[str, Any],
) -> dict[str, Any]:
    return _nearest_transfer_fit(fits, 0.0, config)


def _maximum_fittable_transfer_fit(
    fits: list[dict[str, Any]], config: dict[str, Any],
) -> dict[str, Any]:
    """Select the valid fit at the largest measured/common |Vg|."""
    valid = [
        fit for fit in fits
        if isinstance(fit.get("vg_v"), (int, float)) and np.isfinite(fit["vg_v"])
    ]
    if not valid:
        return {}
    max_abs_vg = max(abs(float(fit["vg_v"])) for fit in valid)
    candidates = [fit for fit in valid if np.isclose(abs(float(fit["vg_v"])), max_abs_vg)]
    # Reuse the standard Vds/status/R2 tie breakers while locking the target to
    # the selected extreme.  The signed target keeps +Vg/-Vg ties deterministic.
    signed_target = max(float(fit["vg_v"]) for fit in candidates)
    return _nearest_transfer_fit(candidates, signed_target, config)


def _fixed_transfer_target(
    records: list[tuple[float, dict[str, Any]]], config: dict[str, Any],
) -> dict[str, Any]:
    """Resolve fixed TLM Vg, giving configured electric field precedence."""
    tlm_cfg = config.get("tlm", {})
    requested_field = tlm_cfg.get("transfer_gate_field_mv_cm")
    requested_vg = tlm_cfg.get("transfer_read_vg_v")
    if requested_field is None and requested_vg is None:
        return {"requested_vg_v": None, "requested_field_mv_cm": None,
                "oxide_thickness_nm": None, "source": None,
                "warning": "fixed TLM Vg/electric field is not configured"}
    if requested_field is None:
        try:
            value = float(requested_vg)
        except (TypeError, ValueError):
            value = None
        return {"requested_vg_v": value, "requested_field_mv_cm": None,
                "oxide_thickness_nm": None, "source": "fixed_vg", "warning": ""}
    oxide_values = {
        float(record["device_params"]["oxide_thickness_nm"])
        for _, record in records
        if isinstance(record.get("device_params", {}).get("oxide_thickness_nm"), (int, float))
        and float(record["device_params"]["oxide_thickness_nm"]) > 0
    }
    if len(oxide_values) != 1:
        detail = "missing" if not oxide_values else "inconsistent across TLM devices"
        return {"requested_vg_v": None, "requested_field_mv_cm": requested_field,
                "oxide_thickness_nm": None, "source": "electric_field",
                "warning": f"fixed-field TLM unavailable: oxide thickness is {detail}"}
    oxide_nm = next(iter(oxide_values))
    return {
        "requested_vg_v": field_to_voltage(requested_field, oxide_nm),
        "requested_field_mv_cm": float(requested_field),
        "oxide_thickness_nm": oxide_nm, "source": "electric_field", "warning": "",
    }


def _operating_point_fields(
    prefix: str, fit: dict[str, Any], *, selection_reason: str,
    requested_vg_v: float | None = None, requested_field_mv_cm: float | None = None,
    oxide_thickness_nm: float | None = None, source: str | None = None,
) -> dict[str, Any]:
    actual_vg = fit.get("vg_v")
    delta_vg = (
        float(actual_vg) - float(requested_vg_v)
        if isinstance(actual_vg, (int, float)) and requested_vg_v is not None else None
    )
    actual_field = (
        float(actual_vg) / (float(oxide_thickness_nm) * 0.1)
        if isinstance(actual_vg, (int, float))
        and isinstance(oxide_thickness_nm, (int, float)) and oxide_thickness_nm > 0
        else None
    )
    return {
        f"{prefix}_requested_vg_v": requested_vg_v,
        f"{prefix}_requested_field_mv_cm": requested_field_mv_cm,
        f"{prefix}_actual_field_mv_cm": actual_field,
        f"{prefix}_field_oxide_thickness_nm": oxide_thickness_nm,
        f"{prefix}_source": source,
        f"{prefix}_vg_v": actual_vg, f"{prefix}_vg_delta_v": delta_vg,
        f"{prefix}_vds_v": fit.get("vds_v"),
        f"{prefix}_rsh_ohm_sq": fit.get("rsh_ohm_sq"),
        f"{prefix}_rcw_ohm_um": fit.get("rcw_ohm_um"),
        f"{prefix}_rc_ohm": fit.get("rc_ohm"), f"{prefix}_lt_um": fit.get("lt_um"),
        f"{prefix}_rhoc_ohm_cm2": fit.get("rhoc_ohm_cm2"),
        f"{prefix}_rho_film_ohm_cm": fit.get("rho_film_ohm_cm"),
        f"{prefix}_r2": fit.get("r2"),
        f"{prefix}_acceptance_status": fit.get("acceptance_status"),
        f"{prefix}_acceptance_reasons": "; ".join(fit.get("acceptance_reasons", [])),
        f"{prefix}_selection_reason": selection_reason,
    }

def _summary_row(
    sample: str, tlm_id: str, records: list[tuple[float, dict[str, Any]]],
    config: dict[str, Any], *,
    ungated_fits: list[dict[str, Any]] | None = None,
    output_fits: list[dict[str, Any]] | None = None,
    transfer_fits: list[dict[str, Any]] | None = None,
    length_rows: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    width = float(records[0][1]["device_params"].get("channel_width_um", 100))
    film_thickness_nm = _common_film_thickness(records)
    vd_max = float(config.get("tlm", {}).get("tlm_output_vd_max", 0.1))
    r2w = float(config.get("advanced", {}).get("tlm_r2_warning", .9))
    role_records: dict[str, list[tuple[float, dict[str, Any]]]] = defaultdict(list)
    for item in records: role_records[_tlm_role(item[1])].append(item)
    transfer_records = role_records["gated_transfer"]
    if ungated_fits is None: ungated_fits = _output_fit_rows(role_records["ungated_ltlm"], config)
    if output_fits is None: output_fits = _output_fit_rows(role_records["output_sanity"], config)
    if transfer_fits is None:
        transfer_fits = _transfer_rows(transfer_records, config) if transfer_records else []
    if length_rows is None:
        length_rows = _length_metrics(transfer_records)

    def vals(key: str) -> list[float]:
        return [
            float(row[key]) for row in length_rows
            if isinstance(row.get(key), (int, float)) and np.isfinite(row[key])
        ]

    best_transfer = _best_fit(transfer_fits)
    zero_transfer = _near_zero_transfer_fit(transfer_fits, config)
    fixed_target = _fixed_transfer_target(records, config)
    fixed_transfer = (
        _nearest_transfer_fit(transfer_fits, fixed_target["requested_vg_v"], config)
        if fixed_target["requested_vg_v"] is not None else {}
    )
    maximum_transfer = _maximum_fittable_transfer_fit(transfer_fits, config)
    lengths = sorted({length for length, _ in records})
    ss = vals("ss_mv_dec")
    mobility = vals("mobility_cm2_vs")
    gm = vals("gm_max_s")
    ion = vals("ion_ioff_log10")
    vth = vals("vth_v")
    row = {
        "sample_id": sample,
        "tlm_id": tlm_id,
        "n_files": len(records),
        "n_lengths": len(lengths),
        "lch_min_um": min(lengths) if lengths else None,
        "lch_max_um": max(lengths) if lengths else None,
        "width_um": width,
        "film_thickness_nm": film_thickness_nm,
        "analysis_level": "individual",
        "n_tlm_structures": 1,
        "transfer_fit_count": len(transfer_fits),
        "transfer_best_vg_v": best_transfer.get("vg_v"),
        "transfer_best_vds_v": best_transfer.get("vds_v"),
        "transfer_best_r2": best_transfer.get("r2"),
        "transfer_acceptance_status": best_transfer.get("acceptance_status"),
        "transfer_acceptance_reasons": "; ".join(best_transfer.get("acceptance_reasons", [])),
        "transfer_rsh_ohm_sq": best_transfer.get("rsh_ohm_sq"),
        "transfer_rcw_ohm_um": best_transfer.get("rcw_ohm_um"),
        "transfer_rc_ohm": best_transfer.get("rc_ohm"),
        "transfer_lt_um": best_transfer.get("lt_um"),
        "transfer_rhoc_ohm_cm2": best_transfer.get("rhoc_ohm_cm2"),
        "transfer_rho_film_ohm_cm": best_transfer.get("rho_film_ohm_cm"),
        "ss_min_mv_dec": min(ss) if ss else None,
        "ss_mean_mv_dec": float(np.mean(ss)) if ss else None,
        "mobility_max_cm2_vs": max(mobility) if mobility else None,
        "gm_max_s": max(gm) if gm else None,
        "ion_ioff_max_log10": max(ion) if ion else None,
        "vth_mean_v": float(np.mean(vth)) if vth else None,
    }
    row.update(_operating_point_fields(
        "transfer_zero", zero_transfer,
        selection_reason=("measured/common Vg nearest 0 V; ties use preferred |Vds|, acceptance, then R2" if zero_transfer else "no transfer-TLM fit available"),
    ))
    fixed_reason = fixed_target["warning"] or (
        "configured electric field converted using common oxide thickness; nearest measured/common Vg used"
        if fixed_target["source"] == "electric_field" else
        "configured fixed Vg; nearest measured/common Vg used"
    )
    row.update(_operating_point_fields(
        "transfer_fixed", fixed_transfer, selection_reason=fixed_reason,
        requested_vg_v=fixed_target["requested_vg_v"],
        requested_field_mv_cm=fixed_target["requested_field_mv_cm"],
        oxide_thickness_nm=fixed_target["oxide_thickness_nm"], source=fixed_target["source"],
    ))
    row.update(_operating_point_fields(
        "transfer_max", maximum_transfer,
        selection_reason=("largest |measured/common Vg| with a TLM fit; ties use preferred |Vds|, acceptance, then R2" if maximum_transfer else "no transfer-TLM fit available"),
    ))
    row.update(_mode_summary_fields("ungated", ungated_fits))
    row.update(_mode_summary_fields("output", output_fits))
    source_files = {
        "ungated_ltlm": "; ".join(Path(r.get("source_file", "")).name for _, r in role_records["ungated_ltlm"] if r.get("source_file")),
        "gated_transfer_vg0": "; ".join(Path(r.get("source_file", "")).name for _, r in transfer_records if r.get("source_file")),
    }
    row.update(_index_fields(ungated_fits, transfer_fits, config, source_files))
    return row


def _mean_std(values: list[float]) -> tuple[float, float]:
    return float(np.mean(values)), float(np.std(values, ddof=1)) if len(values) > 1 else 0.0


def _aggregate_transfer_rows(
    records: list[tuple[str, float, dict[str, Any]]], config: dict[str, Any],
    width: float, film_thickness_nm: float | None,
) -> list[dict[str, Any]]:
    """Pool TLM structures by Lch, fit the per-length means, and retain spread."""
    factor = float(config.get("transfer", {}).get("ioff_leakage_factor", 3.0))
    r2_warn = float(config.get("advanced", {}).get("tlm_r2_warning", 0.9))
    by_vds: dict[float, list[tuple[str, float, Any, dict[str, Any], Any, Any]]] = defaultdict(list)
    for tlm_id, length, record in records:
        if str(record["classification"].get("sweep_variable", "")).lower() not in {"vg", "vbg"}:
            continue
        for segment in record["segments"]:
            if segment.direction == "forward" and segment.bias_level is not None:
                classification = record["classification"]
                drain = classification.get("drain_current_raw_column", "Id")
                gate = classification.get("gate_current_column") or classification.get("gate_leakage_column")
                sweep = np.asarray(segment.sweep_values, float)
                drain_prepared = _prepare_interpolation(
                    sweep, np.asarray(segment.data.get(drain, []), float)
                )
                gate_prepared = (
                    _prepare_interpolation(sweep, np.asarray(segment.data[gate], float))
                    if gate and segment.data.get(gate) else None
                )
                by_vds[round(float(segment.bias_level), 6)].append(
                    (tlm_id, length, segment, record, drain_prepared, gate_prepared)
                )
    rows: list[dict[str, Any]] = []
    for vds, members in sorted(by_vds.items()):
        grids = [{round(float(v), 9) for v in segment.sweep_values if np.isfinite(v)}
                 for _, _, segment, _, _, _ in members]
        common_vg = sorted(set.intersection(*grids)) if grids else []
        for vg in common_vg:
            by_length: dict[float, list[dict[str, Any]]] = defaultdict(list)
            excluded = 0
            for tlm_id, length, segment, record, drain_prepared, gate_prepared in members:
                id_at = _interpolate_prepared(drain_prepared, vg)
                ig_at = _interpolate_prepared(gate_prepared, vg) if gate_prepared else None
                if id_at is None or abs(id_at) <= 1e-15 or (ig_at is not None and abs(id_at) <= factor * abs(ig_at)):
                    excluded += 1
                    continue
                by_length[length].append({
                    "tlm_id": tlm_id, "source_file": record["source_file"].name,
                    "id_abs_a": abs(float(id_at)), "rtotal_ohm": abs(vds) / abs(float(id_at)),
                })
            length_stats = []
            for length, observations in sorted(by_length.items()):
                id_mean, id_std = _mean_std([item["id_abs_a"] for item in observations])
                r_mean, r_std = _mean_std([item["rtotal_ohm"] for item in observations])
                length_stats.append({
                    "lch_um": length, "id_mean_a": id_mean, "id_std_a": id_std,
                    "rtotal_mean_ohm": r_mean, "rtotal_std_ohm": r_std,
                    "n_observations": len(observations),
                    "n_tlm_structures": len({item["tlm_id"] for item in observations}),
                    "tlm_ids": "; ".join(sorted({item["tlm_id"] for item in observations})),
                    "source_files": "; ".join(item["source_file"] for item in observations),
                })
            if len(length_stats) < 3:
                continue
            points = [(item["lch_um"], item["rtotal_mean_ohm"]) for item in length_stats]
            fit = _fit(points, width, r2_warn, config.get("tlm", {}), film_thickness_nm)
            rows.append({
                "vds_v": vds, "vg_v": vg, "n_lengths": len(length_stats),
                "n_observations": sum(item["n_observations"] for item in length_stats),
                "n_tlm_structures": len({tlm_id for tlm_id, _, _, _, _, _ in members}),
                "excluded": excluded, "points": points, "length_stats": length_stats, **fit,
            })
    return rows


def _aggregate_output_rows(
    records: list[tuple[str, float, dict[str, Any]]], config: dict[str, Any],
    width: float, film_thickness_nm: float | None, role: str,
) -> list[dict[str, Any]]:
    vd_max = float(config.get("tlm", {}).get("tlm_output_vd_max", 0.1))
    r2_warn = float(config.get("advanced", {}).get("tlm_r2_warning", 0.9))
    by_bias_length: dict[float | str, dict[float, list[dict[str, Any]]]] = defaultdict(lambda: defaultdict(list))
    for tlm_id, length, record in records:
        if _tlm_role(record) != role:
            continue
        for bias, resistance in _output_resistance(record, vd_max).items():
            by_bias_length[bias][length].append({
                "tlm_id": tlm_id, "source_file": record["source_file"].name,
                "rtotal_ohm": float(resistance),
            })
    rows: list[dict[str, Any]] = []
    for bias, by_length in by_bias_length.items():
        length_stats = []
        for length, observations in sorted(by_length.items()):
            r_mean, r_std = _mean_std([item["rtotal_ohm"] for item in observations])
            length_stats.append({
                "lch_um": length, "rtotal_mean_ohm": r_mean, "rtotal_std_ohm": r_std,
                "n_observations": len(observations),
                "n_tlm_structures": len({item["tlm_id"] for item in observations}),
                "tlm_ids": "; ".join(sorted({item["tlm_id"] for item in observations})),
                "source_files": "; ".join(item["source_file"] for item in observations),
            })
        if len(length_stats) < 3:
            continue
        points = [(item["lch_um"], item["rtotal_mean_ohm"]) for item in length_stats]
        fit = _fit(points, width, r2_warn, config.get("tlm", {}), film_thickness_nm)
        rows.append({
            "gate_bias": bias, "n_lengths": len(length_stats),
            "n_observations": sum(item["n_observations"] for item in length_stats),
            "n_tlm_structures": len({item["tlm_id"] for stats in length_stats for item in by_length[stats["lch_um"]]}),
            "points": points, "length_stats": length_stats, **fit,
        })
    return rows


def _write_sample_master(
    output_root: Path, sample: str,
    structures: dict[str, list[tuple[float, dict[str, Any]]]], config: dict[str, Any],
) -> tuple[Path | None, dict[str, Any] | None]:
    """Write the sample-level statistical TLM result across independent structures."""
    eligible = {key: value for key, value in structures.items() if len({item[0] for item in value}) >= 3}
    if len(eligible) < 2:
        return None, None
    records = [(tlm_id, length, record) for tlm_id, members in eligible.items() for length, record in members]
    widths = {float(record["device_params"].get("channel_width_um", 100)) for _, _, record in records}
    if len(widths) != 1:
        LOGGER.warning("Sample-level TLM %s skipped: structures use different channel widths", sample)
        return None, None
    width = next(iter(widths))
    thicknesses = {float(record["device_params"]["film_thickness_nm"])
                   for _, _, record in records if record["device_params"].get("film_thickness_nm") is not None}
    film_thickness_nm = next(iter(thicknesses)) if len(thicknesses) == 1 else None
    if len(thicknesses) > 1:
        LOGGER.warning("Sample-level TLM %s: mixed film thicknesses; bulk resistivity is not reported", sample)
    transfer_fits = _aggregate_transfer_rows(records, config, width, film_thickness_nm)
    ungated_fits = _aggregate_output_rows(records, config, width, film_thickness_nm, "ungated_ltlm")
    output_fits = _aggregate_output_rows(records, config, width, film_thickness_nm, "output_sanity")
    if not transfer_fits and not ungated_fits and not output_fits:
        return None, None

    import openpyxl
    from openpyxl.worksheet.table import Table, TableStyleInfo
    import matplotlib; matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    group_dir = output_root / sample / "MASTER"
    safe_mkdir(group_dir, parents=True, exist_ok=True)
    keep_plots = bool(config.get("execution", {}).get("save_plots", False))
    temporary_plots: Path | None = None
    if keep_plots:
        plots = group_dir / "plots"; safe_mkdir(plots, parents=True, exist_ok=True)
        for pattern in ("master_transfer_*.png", "master_ungated_*.png", "master_output_*.png"):
            for old in safe_path(plots).glob(pattern): safe_unlink(old)
    else:
        if safe_path(group_dir / "plots").exists(): safe_rmtree(group_dir / "plots")
        temporary_plots = temporary_directory(config, "tlm_master_", group_dir); plots = temporary_plots

    wb = openpyxl.Workbook(); readme = wb.active; readme.title = "README"
    readme.append([f"Master statistical TLM: {sample}"])
    readme.append(["Method", "Independent TLM structures are pooled at each Lch; mean and sample standard deviation are reported; the TLM fit uses per-length mean Rtotal."])
    readme.append(["TLM structures", len(eligible)]); readme.append(["Width (um)", width]); readme.append(["Film thickness (nm)", film_thickness_nm])

    def table_sheet(name: str, headers: list[str], rows: list[list[Any]]) -> None:
        ws = wb.create_sheet(name); ws.append(headers)
        for row in rows: ws.append(row)
        ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(headers))}{max(1, len(rows)+1)}"
        if rows:
            ref = ws.auto_filter.ref
            table = Table(displayName="Tbl" + re.sub(r"\W", "", name), ref=ref)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
        for index, header in enumerate(headers, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(index)].width = min(max(len(header) + 2, 12), 34)

    table_sheet("Source Structures", ["sample_label","tlm_id","lch_um","filename","tlm_role","width_um","film_thickness_nm"],
                [[sample, tlm_id, length, record["source_file"].name, _tlm_role(record), width, record["device_params"].get("film_thickness_nm")]
                 for tlm_id, length, record in records])
    fit_headers = ["mode","read_bias_v","vds_v","n_tlm_structures","n_lengths","n_observations","rsh_ohm_sq","rho_film_ohm_cm","rcw_ohm_um","rc_ohm","lt_um","rhoc_ohm_cm2","r2","acceptance_status","acceptance_reasons"]
    fit_rows = []
    for fit in transfer_fits:
        fit_rows.append(["transfer",fit.get("vg_v"),fit.get("vds_v"),fit.get("n_tlm_structures"),fit.get("n_lengths"),fit.get("n_observations"),fit.get("rsh_ohm_sq"),fit.get("rho_film_ohm_cm"),fit.get("rcw_ohm_um"),fit.get("rc_ohm"),fit.get("lt_um"),fit.get("rhoc_ohm_cm2"),fit.get("r2"),fit.get("acceptance_status"),"; ".join(fit.get("acceptance_reasons",[]))])
    for fit in ungated_fits:
        fit_rows.append(["ungated",fit.get("gate_bias"),None,fit.get("n_tlm_structures"),fit.get("n_lengths"),fit.get("n_observations"),fit.get("rsh_ohm_sq"),fit.get("rho_film_ohm_cm"),fit.get("rcw_ohm_um"),fit.get("rc_ohm"),fit.get("lt_um"),fit.get("rhoc_ohm_cm2"),fit.get("r2"),fit.get("acceptance_status"),"; ".join(fit.get("acceptance_reasons",[]))])
    for fit in output_fits:
        fit_rows.append(["output_sanity",fit.get("gate_bias"),None,fit.get("n_tlm_structures"),fit.get("n_lengths"),fit.get("n_observations"),fit.get("rsh_ohm_sq"),fit.get("rho_film_ohm_cm"),fit.get("rcw_ohm_um"),fit.get("rc_ohm"),fit.get("lt_um"),fit.get("rhoc_ohm_cm2"),fit.get("r2"),fit.get("acceptance_status"),"; ".join(fit.get("acceptance_reasons",[]))])
    table_sheet("Aggregate Fits", fit_headers, fit_rows)

    stat_headers = ["mode","read_bias_v","vds_v","lch_um","id_mean_a","id_std_a","rtotal_mean_ohm","rtotal_std_ohm","n_observations","n_tlm_structures","tlm_ids","source_files"]
    stat_rows = []
    for fit in transfer_fits:
        for stat in fit["length_stats"]:
            stat_rows.append(["transfer",fit["vg_v"],fit["vds_v"],stat["lch_um"],stat.get("id_mean_a"),stat.get("id_std_a"),stat["rtotal_mean_ohm"],stat["rtotal_std_ohm"],stat["n_observations"],stat["n_tlm_structures"],stat["tlm_ids"],stat["source_files"]])
    for fit in ungated_fits:
        for stat in fit["length_stats"]:
            stat_rows.append(["ungated",fit["gate_bias"],None,stat["lch_um"],None,None,stat["rtotal_mean_ohm"],stat["rtotal_std_ohm"],stat["n_observations"],stat["n_tlm_structures"],stat["tlm_ids"],stat["source_files"]])
    for fit in output_fits:
        for stat in fit["length_stats"]:
            stat_rows.append(["output_sanity",fit["gate_bias"],None,stat["lch_um"],None,None,stat["rtotal_mean_ohm"],stat["rtotal_std_ohm"],stat["n_observations"],stat["n_tlm_structures"],stat["tlm_ids"],stat["source_files"]])
    table_sheet("Length Statistics", stat_headers, stat_rows)

    diagnostics = []
    for mode, fits in (("transfer", transfer_fits), ("ungated", ungated_fits), ("output_sanity", output_fits)):
        for fit in fits:
            bias = fit.get("vg_v") if mode == "transfer" else fit.get("gate_bias")
            for stat in fit["length_stats"]:
                measured = stat["rtotal_mean_ohm"] * width
                predicted = fit.get("slope_ohm_per_um", 0) * stat["lch_um"] + fit.get("intercept_ohm", 0)
                diagnostics.append([mode,bias,fit.get("vds_v"),stat["lch_um"],measured,predicted,measured-predicted,stat["rtotal_std_ohm"]*width])
    table_sheet("Fit Diagnostics", ["mode","read_bias_v","vds_v","lch_um","measured_mean_rtotal_w_ohm_um","predicted_rtotal_w_ohm_um","residual_ohm_um","rtotal_w_std_ohm_um"], diagnostics)

    html_plots: list[Path] = []
    if transfer_fits:
        selected = sorted(set(np.linspace(0, len(transfer_fits)-1, min(6, len(transfer_fits)), dtype=int)))
        for plot_index, fit_index in enumerate(selected):
            fit = transfer_fits[fit_index]; stats = fit["length_stats"]
            x = np.asarray([item["lch_um"] for item in stats], float)
            r = np.asarray([item["rtotal_mean_ohm"]*width for item in stats], float)
            rerr = np.asarray([item["rtotal_std_ohm"]*width for item in stats], float)
            current = np.asarray([item["id_mean_a"] for item in stats], float)
            ierr = np.asarray([item["id_std_a"] for item in stats], float)
            order = np.argsort(x); prediction = fit["slope_ohm_per_um"]*x[order] + fit["intercept_ohm"]
            fig, axes = plt.subplots(1, 2, figsize=(11, 4.4))
            axes[0].errorbar(x, r, yerr=rerr, fmt="o", capsize=3, label="mean ± SD")
            axes[0].plot(x[order], prediction, label=f"fit R²={fit.get('r2', float('nan')):.3g}")
            axes[0].set(xlabel="Lch (µm)", ylabel="Rtotal×W (Ω·µm)", title="Aggregate TLM fit"); axes[0].grid(alpha=.3); axes[0].legend()
            axes[1].errorbar(x, current, yerr=ierr, fmt="o", capsize=3)
            axes[1].set(xlabel="Lch (µm)", ylabel="|Id| (A)", title="Current spread by length"); axes[1].set_yscale("log"); axes[1].grid(alpha=.3)
            fig.suptitle(f"{sample} master TLM: Vg={fit['vg_v']:g} V, Vds={fit['vds_v']:g} V")
            fig.tight_layout(); plot_path=plots / f"master_transfer_{plot_index}_vg_{fit['vg_v']:g}.png"; fig.savefig(prepare_write_path(plot_path)); plt.close(fig); html_plots.append(plot_path)
    if ungated_fits:
        selected = sorted(set(np.linspace(0, len(ungated_fits)-1, min(6, len(ungated_fits)), dtype=int)))
        for plot_index, fit_index in enumerate(selected):
            fit = ungated_fits[fit_index]; stats = fit["length_stats"]
            x = np.asarray([item["lch_um"] for item in stats], float)
            r = np.asarray([item["rtotal_mean_ohm"]*width for item in stats], float)
            rerr = np.asarray([item["rtotal_std_ohm"]*width for item in stats], float)
            order = np.argsort(x); prediction = fit["slope_ohm_per_um"]*x[order] + fit["intercept_ohm"]
            fig, ax = plt.subplots(figsize=(6.4, 4.4))
            ax.errorbar(x, r, yerr=rerr, fmt="o", capsize=3, label="mean ± SD")
            ax.plot(x[order], prediction, label=f"fit R²={fit.get('r2', float('nan')):.3g}")
            ax.set(xlabel="Lch (µm)", ylabel="Rtotal×W (Ω·µm)", title=f"{sample} master Ungated LTLM")
            ax.grid(alpha=.3); ax.legend(); fig.tight_layout()
            plot_path=plots / f"master_ungated_{plot_index}.png"; fig.savefig(prepare_write_path(plot_path)); plt.close(fig); html_plots.append(plot_path)

    path = group_dir / f"TLM_{sample}_MASTER.xlsx"; wb.save(prepare_write_path(path)); wb.close()
    best_transfer = _best_fit(transfer_fits); best_ungated = _best_fit(ungated_fits)
    zero_transfer = _near_zero_transfer_fit(transfer_fits, config)
    summary_records = [(length, record) for _, length, record in records]
    fixed_target = _fixed_transfer_target(summary_records, config)
    fixed_transfer = (
        _nearest_transfer_fit(transfer_fits, fixed_target["requested_vg_v"], config)
        if fixed_target["requested_vg_v"] is not None else {}
    )
    maximum_transfer = _maximum_fittable_transfer_fit(transfer_fits, config)
    result_lines = [
        f"# Master Statistical TLM — {sample}", "",
        f"**Independent TLM structures:** {len(eligible)}", f"**Source files:** {len(records)}", "",
        "Measurements are first grouped by channel length. Mean, sample standard deviation, count, and source TLM IDs are retained. The TLM line is fitted through the mean Rtotal at each length; raw individual TLM reports remain separate.", "",
        "## Representative Results", "",
        "| Mode | Rsh (ohm/sq) | Film resistivity (ohm cm) | Contact resistivity (ohm cm2) | R2 | Status |",
        "|---|---:|---:|---:|---:|---|",
    ]
    for label, fit in (("Transfer", best_transfer), ("Ungated LTLM", best_ungated)):
        if fit:
            result_lines.append(f"| {label} | {fit.get('rsh_ohm_sq')} | {fit.get('rho_film_ohm_cm')} | {fit.get('rhoc_ohm_cm2')} | {fit.get('r2')} | {fit.get('acceptance_status')} |")
    result_lines.extend(["", "## Plots", ""])
    report = group_dir / f"TLM_{sample}_MASTER_report.html"
    write_html_report(report, f"Master TLM {sample}", result_lines, html_plots)
    if temporary_plots is not None: safe_rmtree(temporary_plots, ignore_errors=True)

    summary = {
        "sample_id": sample, "tlm_id": "MASTER", "analysis_level": "master_aggregate",
        "n_tlm_structures": len(eligible), "n_files": len(records),
        "n_lengths": len({length for _, length, _ in records}),
        "lch_min_um": min(length for _, length, _ in records), "lch_max_um": max(length for _, length, _ in records),
        "width_um": width, "film_thickness_nm": film_thickness_nm,
        "transfer_fit_count": len(transfer_fits),
        "transfer_best_vg_v": best_transfer.get("vg_v") if best_transfer else None,
        "transfer_best_vds_v": best_transfer.get("vds_v") if best_transfer else None,
        "transfer_best_r2": best_transfer.get("r2") if best_transfer else None,
        "transfer_acceptance_status": best_transfer.get("acceptance_status") if best_transfer else None,
        "transfer_acceptance_reasons": "; ".join(best_transfer.get("acceptance_reasons", [])) if best_transfer else "",
        "transfer_rsh_ohm_sq": best_transfer.get("rsh_ohm_sq") if best_transfer else None,
        "transfer_rho_film_ohm_cm": best_transfer.get("rho_film_ohm_cm") if best_transfer else None,
        "transfer_rcw_ohm_um": best_transfer.get("rcw_ohm_um") if best_transfer else None,
        "transfer_rc_ohm": best_transfer.get("rc_ohm") if best_transfer else None,
        "transfer_lt_um": best_transfer.get("lt_um") if best_transfer else None,
        "transfer_rhoc_ohm_cm2": best_transfer.get("rhoc_ohm_cm2") if best_transfer else None,
        "transfer_zero_vg_v": zero_transfer.get("vg_v") if zero_transfer else None,
        "transfer_zero_vds_v": zero_transfer.get("vds_v") if zero_transfer else None,
        "transfer_zero_rsh_ohm_sq": zero_transfer.get("rsh_ohm_sq") if zero_transfer else None,
        "transfer_zero_rcw_ohm_um": zero_transfer.get("rcw_ohm_um") if zero_transfer else None,
        "transfer_zero_rc_ohm": zero_transfer.get("rc_ohm") if zero_transfer else None,
        "transfer_zero_lt_um": zero_transfer.get("lt_um") if zero_transfer else None,
        "transfer_zero_rhoc_ohm_cm2": zero_transfer.get("rhoc_ohm_cm2") if zero_transfer else None,
        "transfer_zero_rho_film_ohm_cm": zero_transfer.get("rho_film_ohm_cm") if zero_transfer else None,
        "transfer_zero_r2": zero_transfer.get("r2") if zero_transfer else None,
        "transfer_zero_acceptance_status": zero_transfer.get("acceptance_status") if zero_transfer else None,
        "transfer_zero_acceptance_reasons": "; ".join(zero_transfer.get("acceptance_reasons", [])) if zero_transfer else "",
        "transfer_zero_selection_reason": (
            "measured/common Vg nearest 0 V; ties use preferred |Vds|, acceptance, then R2"
            if zero_transfer else ""
        ),
    }
    summary.update(_operating_point_fields(
        "transfer_zero", zero_transfer,
        selection_reason=("measured/common Vg nearest 0 V; ties use preferred |Vds|, acceptance, then R2" if zero_transfer else "no transfer-TLM fit available"),
    ))
    fixed_reason = fixed_target["warning"] or (
        "configured electric field converted using common oxide thickness; nearest measured/common Vg used"
        if fixed_target["source"] == "electric_field" else
        "configured fixed Vg; nearest measured/common Vg used"
    )
    summary.update(_operating_point_fields(
        "transfer_fixed", fixed_transfer, selection_reason=fixed_reason,
        requested_vg_v=fixed_target["requested_vg_v"],
        requested_field_mv_cm=fixed_target["requested_field_mv_cm"],
        oxide_thickness_nm=fixed_target["oxide_thickness_nm"], source=fixed_target["source"],
    ))
    summary.update(_operating_point_fields(
        "transfer_max", maximum_transfer,
        selection_reason=("largest |measured/common Vg| with a TLM fit; ties use preferred |Vds|, acceptance, then R2" if maximum_transfer else "no transfer-TLM fit available"),
    ))
    summary.update(_mode_summary_fields("ungated", ungated_fits))
    summary.update(_mode_summary_fields("output", output_fits))
    summary.update(_index_fields(ungated_fits, transfer_fits, config, {
        "ungated_ltlm": "; ".join(record["source_file"].name for _, _, record in records if _tlm_role(record) == "ungated_ltlm"),
        "gated_transfer_vg0": "; ".join(record["source_file"].name for _, _, record in records if _tlm_role(record) == "gated_transfer"),
    }))
    return path, summary

def _write_master_summary(output_root: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        return
    headers = list(rows[0].keys())
    csv_path = output_root / "tlm_master_summary.csv"
    with open(prepare_write_path(csv_path), "w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)
    try:
        import openpyxl
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError:
        LOGGER.warning("openpyxl not available - wrote CSV only for TLM master summary")
        return
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TLM Summary"
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header) for header in headers])
    ws.freeze_panes = "A2"
    ref = f"A1:{openpyxl.utils.get_column_letter(len(headers))}{len(rows) + 1}"
    ws.auto_filter.ref = ref
    table = Table(displayName="TlmMasterSummary", ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    for index, header in enumerate(headers, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(index)].width = min(max(len(header) + 2, 12), 28)
    wb.save(prepare_write_path(output_root / "tlm_master_summary.xlsx"))
    wb.close()


def _finite_number(value: Any) -> float | None:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if np.isfinite(number) else None


def _device_sweep_value(record: dict[str, Any], vds_v: float, section: str, key: str) -> float | None:
    candidates: list[tuple[float, float]] = []
    for sweep in record.get("metrics", {}).get("sweep_results", []):
        identity = sweep.get("identity", {})
        if identity.get("direction") != "forward":
            continue
        measured = _finite_number(identity.get("measured_vds_v"))
        bias = _finite_number(identity.get("bias_value_v"))
        actual_vds = measured if measured is not None else bias
        value = _finite_number((sweep.get(section) or {}).get(key))
        if actual_vds is not None and value is not None:
            candidates.append((abs(actual_vds - float(vds_v)), value))
    return min(candidates, key=lambda item: item[0])[1] if candidates else None


def _polarity(records: list[tuple[float, dict[str, Any]]]) -> str:
    values = [str(record.get("device_params", {}).get("polarity", "")).lower() for _, record in records]
    return "n" if values and all(value.startswith("n") for value in values) else "p"


def _onward_envelope(values: list[float], polarity: str) -> float | None:
    if not values:
        return None
    return max(values) if polarity == "n" else min(values)


def _gated_operating_points(
    records: list[tuple[float, dict[str, Any]]], fits: list[dict[str, Any]], config: dict[str, Any],
) -> list[dict[str, Any]]:
    polarity = _polarity(records)
    points: list[dict[str, Any]] = []

    zero = _near_zero_transfer_fit(fits, config)
    points.append({
        "condition": "near_zero", "fit": zero, "requested_vg_v": 0.0,
        "target_vg_min_v": 0.0, "target_vg_max_v": 0.0,
        "selection_reason": "nearest common measured Vg to 0 V; no acceptance-driven movement",
    })

    fixed_targets = []
    fixed_sources = []
    fixed_fields = []
    oxide_values = []
    for _, record in records:
        resolved = resolve_ion_bias(config, record.get("device_params"))
        target = _finite_number(resolved.get("fixed_vg_v"))
        if target is not None:
            fixed_targets.append(target)
        if resolved.get("fixed_vg_source"):
            fixed_sources.append(str(resolved["fixed_vg_source"]))
        if resolved.get("gate_field_mv_cm") is not None:
            fixed_fields.append(float(resolved["gate_field_mv_cm"]))
        if resolved.get("oxide_thickness_nm") is not None:
            oxide_values.append(float(resolved["oxide_thickness_nm"]))
    fixed_complete = len(fixed_targets) == len(records)
    fixed_target = _onward_envelope(fixed_targets, polarity) if fixed_complete else None
    fixed = _nearest_transfer_fit(fits, fixed_target, config) if fixed_target is not None else {}
    points.append({
        "condition": "ion_fixed_gate_or_field", "fit": fixed,
        "requested_vg_v": fixed_target,
        "target_vg_min_v": min(fixed_targets) if fixed_targets else None,
        "target_vg_max_v": max(fixed_targets) if fixed_targets else None,
        "requested_field_mv_cm": fixed_fields[0] if len(set(fixed_fields)) == 1 and fixed_fields else None,
        "oxide_thickness_min_nm": min(oxide_values) if oxide_values else None,
        "oxide_thickness_max_nm": max(oxide_values) if oxide_values else None,
        "source": fixed_sources[0] if len(set(fixed_sources)) == 1 and fixed_sources else None,
        "selection_reason": (
            "unified Ion fixed-gate/field setting; common Vg nearest the on-ward envelope of per-device targets"
            if fixed_target is not None else
            "fixed-gate/field Ion condition is unavailable for one or more TLM lengths"
        ),
    })

    preferred_vds = config.get("summary", {}).get("preferred_vd_v")
    vds_values = sorted({float(fit["vds_v"]) for fit in fits if _finite_number(fit.get("vds_v")) is not None},
                        key=lambda value: abs(abs(value) - abs(float(preferred_vds or 0))))
    overdrive_candidates: list[tuple[dict[str, Any], list[float], list[float], list[str]]] = []
    for vds in vds_values:
        targets = []
        vths = []
        sources = []
        for _, record in records:
            resolved = resolve_ion_bias(config, record.get("device_params"))
            overdrive = _finite_number(resolved.get("overdrive_v"))
            vth = _device_sweep_value(record, vds, "vth", "vth_v")
            if overdrive is not None and vth is not None:
                targets.append(vth + overdrive)
                vths.append(vth)
                if resolved.get("overdrive_source"):
                    sources.append(str(resolved["overdrive_source"]))
        target = _onward_envelope(targets, polarity)
        subset = [fit for fit in fits if np.isclose(float(fit.get("vds_v")), vds)]
        if target is not None and len(targets) == len(records) and subset:
            overdrive_candidates.append((_nearest_transfer_fit(subset, target, config), targets, vths, sources))
    overdrive_fit, overdrive_targets, vths, overdrive_sources = overdrive_candidates[0] if overdrive_candidates else ({}, [], [], [])
    points.append({
        "condition": "ion_overdrive_or_field", "fit": overdrive_fit,
        "requested_vg_v": _onward_envelope(overdrive_targets, polarity),
        "target_vg_min_v": min(overdrive_targets) if overdrive_targets else None,
        "target_vg_max_v": max(overdrive_targets) if overdrive_targets else None,
        "vth_min_v": min(vths) if vths else None, "vth_max_v": max(vths) if vths else None,
        "source": overdrive_sources[0] if len(set(overdrive_sources)) == 1 and overdrive_sources else None,
        "selection_reason": (
            "unified Ion overdrive/field setting applied to each device Vth; common Vg nearest the on-ward target envelope"
            if overdrive_targets else "overdrive Ion condition or per-device Vth is unavailable"
        ),
    })

    ioff_fit: dict[str, Any] = {}
    ioff_targets: list[float] = []
    attempted: list[str] = []
    expected_lengths = len({length for length, _ in records})
    for vds in vds_values:
        targets = [
            value for _, record in records
            if (value := _device_sweep_value(record, vds, "subthreshold_swing", "ss_ioff_vg")) is not None
        ]
        if len(targets) != expected_lengths:
            attempted.append(f"Vds={vds:g}: missing per-device Ioff Vg")
            continue
        envelope = _onward_envelope(targets, polarity)
        candidates = [fit for fit in fits if np.isclose(float(fit.get("vds_v")), vds)]
        if polarity == "n":
            candidates = sorted((fit for fit in candidates if float(fit["vg_v"]) >= float(envelope)), key=lambda fit: float(fit["vg_v"]))
        else:
            candidates = sorted((fit for fit in candidates if float(fit["vg_v"]) <= float(envelope)), key=lambda fit: -float(fit["vg_v"]))
        for fit in candidates:
            attempted.append(f"Vds={vds:g}, Vg={float(fit['vg_v']):g}: {fit.get('acceptance_status')}")
            if (fit.get("acceptance_status") == "accepted" and int(fit.get("excluded", 0)) == 0
                    and int(fit.get("n_lengths", 0)) == expected_lengths
                    and _finite_number(fit.get("slope_ohm_per_um")) is not None
                    and _finite_number(fit.get("intercept_ohm")) is not None):
                ioff_fit = fit
                ioff_targets = targets
                break
        if ioff_fit:
            break
    points.append({
        "condition": "near_ioff_first_accepted", "fit": ioff_fit,
        "requested_vg_v": _onward_envelope(ioff_targets, polarity),
        "target_vg_min_v": min(ioff_targets) if ioff_targets else None,
        "target_vg_max_v": max(ioff_targets) if ioff_targets else None,
        "selection_reason": (
            "first accepted all-length leakage-qualified common-Vg fit from the Ioff envelope toward On"
            if ioff_fit else "no accepted all-length fit found from the per-device Ioff envelope toward On"
        ),
        "attempted_candidates": "; ".join(attempted),
    })
    return points


MASTER_TLM_HEADERS = [
    "sample_id", "tlm_id", "analysis_mode", "condition", "polarity", "requested_vg_v",
    "actual_vg_v", "target_vg_min_v", "target_vg_max_v", "requested_field_mv_cm",
    "oxide_thickness_min_nm", "oxide_thickness_max_nm", "vth_min_v", "vth_max_v", "vds_v",
    "gate_bias_v", "n_files", "n_lengths", "expected_lengths", "excluded_points", "width_um",
    "film_thickness_nm", "rho_film_ohm_cm", "rsh_ohm_sq", "rc_ohm", "rcw_ohm_um", "lt_um",
    "rhoc_ohm_cm2", "r2", "slope_std", "intercept_std", "acceptance_status",
    "acceptance_reasons", "warnings", "selection_source", "selection_reason", "attempted_candidates",
    "source_files", "plot_file",
]


def _master_tlm_row(
    sample: str, tlm_id: str, mode: str, condition: str,
    records: list[tuple[float, dict[str, Any]]], fit: dict[str, Any], **extra: Any,
) -> dict[str, Any]:
    width = float(records[0][1].get("device_params", {}).get("channel_width_um", 100)) if records else None
    reasons = fit.get("acceptance_reasons", [])
    warnings = fit.get("warnings", [])
    row = {
        "sample_id": sample, "tlm_id": tlm_id, "analysis_mode": mode, "condition": condition,
        "polarity": _polarity(records) if records else None,
        "actual_vg_v": fit.get("vg_v"), "vds_v": fit.get("vds_v"),
        "gate_bias_v": fit.get("gate_bias"), "n_files": len(records),
        "n_lengths": fit.get("n_lengths", len({point[0] for point in fit.get("points", [])})),
        "expected_lengths": len({length for length, _ in records}), "excluded_points": fit.get("excluded"),
        "width_um": width, "film_thickness_nm": fit.get("film_thickness_nm", _common_film_thickness(records) if records else None),
        "rho_film_ohm_cm": fit.get("rho_film_ohm_cm"), "rsh_ohm_sq": fit.get("rsh_ohm_sq"),
        "rc_ohm": fit.get("rc_ohm"), "rcw_ohm_um": fit.get("rcw_ohm_um"), "lt_um": fit.get("lt_um"),
        "rhoc_ohm_cm2": fit.get("rhoc_ohm_cm2"), "r2": fit.get("r2"),
        "slope_std": fit.get("slope_std"), "intercept_std": fit.get("intercept_std"),
        "acceptance_status": fit.get("acceptance_status", "unavailable"),
        "acceptance_reasons": "; ".join(reasons) if isinstance(reasons, list) else reasons,
        "warnings": "; ".join(warnings) if isinstance(warnings, list) else warnings,
        "source_files": "; ".join(record["source_file"].name for _, record in records),
    }
    row.update(extra)
    return {header: row.get(header) for header in MASTER_TLM_HEADERS}


def _write_master_fit_plot(path: Path, row: dict[str, Any], fit: dict[str, Any]) -> None:
    import matplotlib; matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    points = fit.get("points", [])
    if not points or _finite_number(fit.get("slope_ohm_per_um")) is None:
        return
    width = float(row["width_um"])
    x = np.asarray([point[0] for point in points], float)
    measured = np.asarray([point[1] * width for point in points], float)
    order = np.argsort(x)
    predicted = float(fit["slope_ohm_per_um"]) * x + float(fit["intercept_ohm"])
    fig, axes = plt.subplots(1, 2, figsize=(11, 4.4))
    axes[0].scatter(x, measured, label="Measured")
    axes[0].plot(x[order], predicted[order], label="Free-intercept OLS")
    axes[0].set(xlabel="Channel length (µm)", ylabel="Rtotal × W (Ω·µm)", title="TLM fit")
    axes[0].grid(alpha=.3); axes[0].legend()
    axes[1].axhline(0, color="black", linewidth=.8)
    axes[1].scatter(x, measured - predicted)
    axes[1].set(xlabel="Channel length (µm)", ylabel="Residual (Ω·µm)", title="Fit residuals")
    axes[1].grid(alpha=.3)
    def compact(value: Any) -> str:
        number = _finite_number(value)
        return f"{number:.4g}" if number is not None else "unavailable"
    fig.suptitle(
        f"{row['sample_id']} {row['tlm_id']} · {row['analysis_mode']} · {row['condition']}\n"
        f"R²={compact(row.get('r2'))} · Rsh={compact(row.get('rsh_ohm_sq'))} Ω/sq · "
        f"RcW={compact(row.get('rcw_ohm_um'))} Ω·µm · {row.get('acceptance_status')}"
    )
    fig.tight_layout(); fig.savefig(prepare_write_path(path)); plt.close(fig)


def _write_gated_context_plot(
    path: Path, sample: str, tlm_id: str, fits: list[dict[str, Any]], selected: list[dict[str, Any]],
) -> None:
    import matplotlib; matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    usable = [fit for fit in fits if _finite_number(fit.get("vg_v")) is not None]
    if not usable:
        return
    fig, axes = plt.subplots(2, 2, figsize=(11, 8))
    for ax, key, label in (
        (axes[0, 0], "rsh_ohm_sq", "Sheet resistance (Ω/sq)"),
        (axes[0, 1], "rcw_ohm_um", "Contact resistance × width (Ω·µm)"),
        (axes[1, 0], "lt_um", "Transfer length (µm)"),
        (axes[1, 1], "r2", "Fit R²"),
    ):
        for vds in sorted({float(fit["vds_v"]) for fit in usable}):
            subset = sorted(
                (fit for fit in usable if np.isclose(float(fit["vds_v"]), vds) and _finite_number(fit.get(key)) is not None),
                key=lambda fit: float(fit["vg_v"]),
            )
            if subset:
                ax.plot([fit["vg_v"] for fit in subset], [fit[key] for fit in subset], marker=".", label=f"Vds={vds:g} V")
        for row in selected:
            if _finite_number(row.get("actual_vg_v")) is not None and _finite_number(row.get(key)) is not None:
                ax.scatter([row["actual_vg_v"]], [row[key]], s=55, marker="x", linewidths=2)
                ax.annotate(row["condition"], (row["actual_vg_v"], row[key]), fontsize=7, xytext=(4, 4), textcoords="offset points")
        ax.set(xlabel="Gate voltage (V)", ylabel=label); ax.grid(alpha=.3)
    axes[0, 0].legend(fontsize=8)
    fig.suptitle(f"{sample} {tlm_id}: gated TLM parameter context")
    fig.tight_layout(); fig.savefig(prepare_write_path(path)); plt.close(fig)


def _write_extended_master_tlm(
    output_root: Path,
    groups: dict[tuple[str, str], list[tuple[float, dict[str, Any]]]],
    config: dict[str, Any],
) -> Path | None:
    rows_by_sheet: dict[str, list[dict[str, Any]]] = {
        "Gated TLM": [], "Ungated LTLM": [], "Output Sanity Check": [], "Fit Audit": [],
    }
    plots = output_root / "master_tlm_plots"
    safe_mkdir(plots, parents=True, exist_ok=True)
    for old in plots.glob("*.png"):
        safe_unlink(old)
    for (sample, tlm_id), members in sorted(groups.items()):
        roles: dict[str, list[tuple[float, dict[str, Any]]]] = defaultdict(list)
        for item in members:
            roles[_tlm_role(item[1])].append(item)
        gated = roles["gated_transfer"]
        if len({length for length, _ in gated}) >= 3:
            fits = _transfer_rows(gated, config)
            for fit in fits:
                rows_by_sheet["Fit Audit"].append(_master_tlm_row(
                    sample, tlm_id, "gated_transfer", "candidate_common_vg", gated, fit,
                    selection_reason="candidate fit retained for operating-point audit",
                ))
            for selected in _gated_operating_points(gated, fits, config):
                fit = selected.pop("fit")
                rows_by_sheet["Gated TLM"].append(_master_tlm_row(
                    sample, tlm_id, "gated_transfer", selected.pop("condition"), gated, fit, **selected,
                ))
        for role, sheet, mode in (
            ("ungated_ltlm", "Ungated LTLM", "ungated_ltlm"),
            ("output_sanity", "Output Sanity Check", "output_sanity"),
        ):
            role_records = roles[role]
            if len({length for length, _ in role_records}) < 3:
                continue
            for fit in _output_fit_rows(role_records, config):
                condition = "ungated_low_vd" if role == "ungated_ltlm" else "output_gate_bias"
                row = _master_tlm_row(
                    sample, tlm_id, mode, condition, role_records, fit,
                    selection_source="explicit LTLM filename" if role == "ungated_ltlm" else "gated output characteristic",
                    selection_reason=(
                        "low-Vd resistance from explicit ungated LTLM IV files"
                        if role == "ungated_ltlm" else "gated output-characteristic fit retained only as a sanity check"
                    ),
                )
                rows_by_sheet[sheet].append(row)
                rows_by_sheet["Fit Audit"].append(dict(row))

    reported = [row for sheet in ("Gated TLM", "Ungated LTLM", "Output Sanity Check") for row in rows_by_sheet[sheet]]
    if not reported:
        return None
    fit_lookup: dict[tuple[str, str, str, str, Any, Any], dict[str, Any]] = {}
    for (sample, tlm_id), members in sorted(groups.items()):
        gated = [item for item in members if _tlm_role(item[1]) == "gated_transfer"]
        for fit in _transfer_rows(gated, config) if len({item[0] for item in gated}) >= 3 else []:
            fit_lookup[(sample, tlm_id, "gated_transfer", "", fit.get("vg_v"), fit.get("vds_v"))] = fit
        for role in ("ungated_ltlm", "output_sanity"):
            subset = [item for item in members if _tlm_role(item[1]) == role]
            for fit in _output_fit_rows(subset, config) if len({item[0] for item in subset}) >= 3 else []:
                fit_lookup[(sample, tlm_id, role, "", fit.get("gate_bias"), None)] = fit
    for index, row in enumerate(reported, 1):
        if row["analysis_mode"] == "output_sanity":
            continue
        key_value = row.get("actual_vg_v") if row["analysis_mode"] == "gated_transfer" else row.get("gate_bias_v")
        fit = fit_lookup.get((row["sample_id"], row["tlm_id"], row["analysis_mode"], "", key_value, row.get("vds_v") if row["analysis_mode"] == "gated_transfer" else None), {})
        safe = re.sub(r"[^A-Za-z0-9._-]+", "_", f"{index}_{row['sample_id']}_{row['tlm_id']}_{row['analysis_mode']}_{row['condition']}")
        plot_path = plots / f"{safe}.png"
        _write_master_fit_plot(plot_path, row, fit)
        if plot_path.exists():
            row["plot_file"] = str(plot_path.relative_to(output_root))
    for (sample, tlm_id), members in sorted(groups.items()):
        gated = [item for item in members if _tlm_role(item[1]) == "gated_transfer"]
        if len({item[0] for item in gated}) < 3:
            continue
        fits = _transfer_rows(gated, config)
        selected = [row for row in rows_by_sheet["Gated TLM"] if row["sample_id"] == sample and row["tlm_id"] == tlm_id]
        safe = re.sub(r"[^A-Za-z0-9._-]+", "_", f"{sample}_{tlm_id}_gated_parameter_context")
        _write_gated_context_plot(plots / f"{safe}.png", sample, tlm_id, fits, selected)

    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    workbook = openpyxl.Workbook()
    readme = workbook.active; readme.title = "README"
    readme_rows = [
        ["Master TLM Summary"],
        ["Generated UTC", datetime.now(timezone.utc).isoformat()],
        ["Purpose", "Additive comparison of gated transfer TLM, explicit ungated LTLM, and gated output-characteristic sanity checks."],
        ["Regression", "Ordinary least squares with a free intercept; Rtotal×W is fitted against channel length."],
        ["Gated conditions", "Near 0 V; unified Ion fixed gate/field; unified Ion overdrive/field; first accepted all-length fit from the measured Ioff envelope toward On."],
        ["Important", "No generic best gated fit is selected. Output-characteristic TLM is an Excel-only sanity-check table: it is not plotted and never substitutes for gated transfer or explicit ungated LTLM."],
        ["Resistivity", "rho_film_ohm_cm requires semiconductor film thickness. rhoc_ohm_cm2 is contact resistivity and is reported separately."],
    ]
    for values in readme_rows: readme.append(values)
    readme.column_dimensions["A"].width = 24; readme.column_dimensions["B"].width = 110
    readme["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    readme["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    readme.merge_cells("A1:B1")
    readme.freeze_panes = "A2"
    for sheet_name, rows in rows_by_sheet.items():
        ws = workbook.create_sheet(sheet_name); ws.append(MASTER_TLM_HEADERS)
        for row in rows:
            values = [row.get(header) for header in MASTER_TLM_HEADERS]
            ws.append(values)
            plot_col = MASTER_TLM_HEADERS.index("plot_file") + 1
            if row.get("plot_file"):
                cell = ws.cell(ws.max_row, plot_col)
                cell.hyperlink = row["plot_file"]
                cell.style = "Hyperlink"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(MASTER_TLM_HEADERS))}{max(1, ws.max_row)}"
        if rows:
            table = Table(displayName="Tbl" + re.sub(r"\W", "", sheet_name), ref=ws.auto_filter.ref)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
        for col, header in enumerate(MASTER_TLM_HEADERS, 1):
            ws.cell(1, col).font = Font(bold=True, color="FFFFFF")
            ws.cell(1, col).fill = PatternFill("solid", fgColor="1F4E78")
            ws.cell(1, col).alignment = Alignment(wrap_text=True)
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = min(max(len(header) + 2, 12), 28)
    path = output_root / "master_tlm_summary.xlsx"
    workbook.save(prepare_write_path(path)); workbook.close()
    return path

def generate_tlm_workflow(records: list[dict[str, Any]], output_root: Path, config: dict[str, Any]) -> list[Path]:
    # Device plots may run in child processes, so establish the same explicit
    # DPI/font/grid settings in the parent before rendering TLM figures.
    from fet_analyzer.plotting.transfer_plots import _setup_style
    _setup_style(config)
    groups: dict[tuple[str,str], list[tuple[float,dict[str,Any]]]] = defaultdict(list)
    for record in records:
        if "tlm" in set((record.get("parameter_preflight") or {}).get("suppressed_metrics", [])):
            LOGGER.warning("TLM excluded %s: strict parameter preflight", record.get("source_file"))
            continue
        ident=_identity(record)
        if ident: groups[(ident[0],ident[1])].append((ident[2],record))
    generated=[]; summary_rows=[]
    for (sample,tlm_id),members in groups.items():
        if len({m[0] for m in members}) < 3: continue
        group_path, group_summary = _write_group(
            output_root/sample/tlm_id, sample, tlm_id, members, config
        )
        generated.append(group_path)
        summary_rows.append(group_summary)
    by_sample: dict[str, dict[str, list[tuple[float, dict[str, Any]]]]] = defaultdict(dict)
    for (sample, tlm_id), members in groups.items():
        by_sample[sample][tlm_id] = members
    for sample, structures in sorted(by_sample.items()):
        master_path, master_summary = _write_sample_master(output_root, sample, structures, config)
        if master_path is not None:
            generated.append(master_path)
        if master_summary is not None:
            summary_rows.append(master_summary)
    _write_master_summary(output_root, summary_rows)
    extended = _write_extended_master_tlm(output_root, groups, config)
    if extended is not None:
        generated.append(extended)
    return generated
