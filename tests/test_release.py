from __future__ import annotations
import contextlib
import http.client
import io
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch
import openpyxl
from fet_analyzer.cli import main
from fet_analyzer.parsers.csv_parser import parse_csv_data
from fet_analyzer.dashboard.server import create_server


class ReleaseRegressionTests(unittest.TestCase):
    def test_excel_reads_numeric_measurements_and_closes_workbook(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "input.xlsx"
            wb = openpyxl.Workbook()
            for row in [("Vg", "Vd", "Id"), (0, .1, 1e-8), (1, .1, 2e-8)]:
                wb.active.append(row)
            wb.save(path)
            wb.close()
            parsed = parse_csv_data(path)
            self.assertEqual(parsed["data"]["Vg"], [0, 1])
            self.assertEqual(parsed["num_rows"], 2)
            path.unlink()

    def test_excel_rejects_ambiguous_sheets_formulas_and_legacy_binary(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "input.xlsx"
            wb = openpyxl.Workbook()
            wb.active.append(["Vg", "Id"])
            wb.active.append([1, "=2+2"])
            wb.save(path)
            with self.assertRaisesRegex(ValueError, "formulas"):
                parse_csv_data(path)
            wb.create_sheet("Another measurement")
            wb.save(path)
            wb.close()
            with self.assertRaisesRegex(ValueError, "exactly one"):
                parse_csv_data(path)
            with self.assertRaisesRegex(ValueError, "Legacy"):
                parse_csv_data(Path(folder) / "input.xls")

    def test_summary_only_does_not_require_raw_input_or_ion_read(self):
        with tempfile.TemporaryDirectory() as folder:
            with patch("fet_analyzer.reports.batch_summary.generate_batch_summary") as summary, patch("fet_analyzer.reports.index_page.generate_html_index"):
                result = main(["--input", str(Path(folder) / "missing"), "--output", folder, "--summary-only"])
            self.assertEqual(result, 0)
            summary.assert_called_once()
            self.assertFalse((Path(folder) / "device_parameters.txt").exists())

    def test_input_override_controls_default_output(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            config = root / "config.yaml"
            config.write_text("transfer:\n  ion_fixed_vg_v: -1\n")
            with patch("fet_analyzer.cli.discover_files", return_value=([], [])) as discover:
                self.assertEqual(main(["--input", folder, "--config", str(config), "--dry-run"]), 0)
            self.assertEqual(discover.call_args.kwargs["excluded_roots"][0], root / "output")

    def test_tlm_dry_run_never_calls_continuation(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            with patch("fet_analyzer.cli.discover_files", return_value=([root / "LTLM.csv"], [])), patch("fet_analyzer.tlm_continue.tlm_preflight", return_value={"status": "pass", "eligible_groups": ["demo"]}) as preflight, patch("fet_analyzer.tlm_continue.run_tlm_continuation") as run:
                with contextlib.redirect_stdout(io.StringIO()):
                    self.assertEqual(main(["--input", folder, "--output", folder, "--only", "tlm", "--tlm-continue", "--dry-run"]), 0)
                preflight.assert_called_once()
                run.assert_not_called()
            self.assertEqual(list(root.iterdir()), [])

    def test_invalid_mode_combinations_and_negative_workers_fail(self):
        for args in (["--summary-only", "--dry-run"], ["--summary-only", "--only", "tlm"], ["--workers", "-1"]):
            with contextlib.redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):
                main(args)

    def test_browser_accepts_local_requests_rejects_foreign_origin_and_host(self):
        with tempfile.TemporaryDirectory() as folder:
            with patch("fet_analyzer.dashboard.server.DashboardState._remember_project"):
                server, url = create_server(Path(folder), port=0)
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            try:
                port = server.server_address[1]
                for headers, expected in [({}, 200), ({"Origin": url}, 200), ({"Origin": "https://external.invalid"}, 403), ({"Host": "external.invalid"}, 403), ({"Sec-Fetch-Site": "cross-site"}, 403)]:
                    conn = http.client.HTTPConnection("127.0.0.1", port, timeout=5)
                    try:
                        conn.request("GET", "/", headers=headers)
                        response = conn.getresponse()
                        self.assertEqual(response.status, expected)
                        response.read()
                    finally:
                        conn.close()
            finally:
                server.shutdown()
                server.server_close()
                thread.join(timeout=5)
