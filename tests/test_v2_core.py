from __future__ import annotations

import tempfile
import unittest
import csv
import json
from unittest.mock import patch
from pathlib import Path

import numpy as np

from fet_analyzer.analysis.classifier import MeasurementType, classify_measurement
from fet_analyzer.analysis.geometry import infer_geometry
from fet_analyzer.analysis.ion_bias import resolve_ion_bias
from fet_analyzer.analysis.common_overdrive import apply_max_common_overdrive
from fet_analyzer.analysis.numerics import (
    compute_gm, compute_gm_with_diagnostics, compute_ss,
    nearest_abs_current_at_voltage,
)
from fet_analyzer.analysis.output_sweep import (
    analyze_output_sweep,
    summarize_resistance_by_gate_bias,
)
from fet_analyzer.analysis.sample_database import (
    build_observations,
    discover_comparisons,
)
from fet_analyzer.analysis.segmentation import SweepSegment, segment_sweeps
from fet_analyzer.analysis.transfer_sweep import analyze_transfer_sweep, analyze_transfer_sweeps
from fet_analyzer.analysis.tlm import extract_tlm, extract_rtotal
from fet_analyzer.analysis.tlm_workflow import (
    _aggregate_transfer_rows, _maximum_fittable_transfer_fit,
    _gated_operating_points, _index_fields, _near_zero_transfer_fit, _output_resistance,
    _summary_row, _write_extended_master_tlm,
)
from fet_analyzer.analysis.metric_summary import (
    build_best_results, choose_ion, select_preferred_sweep, transfer_sweep_metrics,
)
from fet_analyzer.analysis.quality import assess_quality
from fet_analyzer.device_parameters import load_master_table, resolve_device_parameters
from fet_analyzer.run_manifest import input_manifest_entry
from fet_analyzer.temp_workspace import RunWorkspace, temporary_directory
from fet_analyzer.reports.batch_summary import (
    BATCH_COLUMNS, EXTENDED_BATCH_COLUMNS, _add_tlm_summary_metrics,
    collect_extended_metrics, collect_metrics, generate_batch_summary,
    write_summary_csv, write_summary_excel,
)
from fet_analyzer.output_structure import create_output_structure
from fet_analyzer.file_discovery import discover_files
from fet_analyzer.reports.html_report import write_html_report
from fet_analyzer.reports.device_report import generate_device_report
from fet_analyzer.reports.index_page import generate_html_index
from fet_analyzer.config import DEFAULT_CONFIG, deep_merge
from fet_analyzer.preflight import apply_preflight_suppression, assess_parameter_preflight
from fet_analyzer.schema import SCHEMA_VERSION, attach_canonical_result, validate_canonical_result
from fet_analyzer.validation import validate_config
from fet_analyzer.tlm_continue import run_tlm_continuation
from fet_analyzer import analyze_transfer


class ReportingConsistencyTests(unittest.TestCase):
    def test_single_constant_vds_is_classified_and_reported_directly(self):
        parsed = {
            "columns": ["Vg", "Vds", "Id"],
            "data": {"Vg": [-1, 0, 1], "Vds": [-2.0, -2.0, -2.0], "Id": [-1e-9, -1e-7, -1e-5]},
        }
        classification = classify_measurement(
            {}, parsed["columns"], parsed["data"],
            "IdVg_K11a_SnO_SampleB_TLM1_5um_1-2.csv",
        )
        self.assertEqual(classification["bias_variable"], "Vds")
        segments = segment_sweeps(parsed, classification)
        self.assertTrue(segments)
        self.assertEqual(segments[0].measured_vds_v, -2.0)
        self.assertEqual(segments[0].vds_reference, "Vds measured directly")

    def test_lone_terminal_voltage_does_not_fabricate_vds(self):
        parsed = {
            "columns": ["Vg", "Vd", "Id"],
            "data": {"Vg": [-1, 0, 1], "Vd": [-2.0, -2.0, -2.0], "Id": [-1e-9, -1e-7, -1e-5]},
        }
        classification = {"sweep_variable": "Vg", "bias_variable": "Vd"}
        segment = segment_sweeps(parsed, classification)[0]
        self.assertEqual(segment.measured_vd_v, -2.0)
        self.assertIsNone(segment.measured_vds_v)
        self.assertIsNone(segment.bias_level)
        self.assertIn("other terminal voltage was not recorded", segment.vds_reference)

    def test_single_measured_vds_runs_full_transfer_analysis_without_config_substitution(self):
        vg = np.linspace(-5.0, 5.0, 101)
        drain = 1e-12 + 2e-5 / (1.0 + np.exp(-(vg - 0.5) * 2.0))
        result = analyze_transfer(
            {
                "Vg": vg,
                "Vds": np.full_like(vg, -1.3),
                "Id": drain,
                "Ig": np.full_like(vg, 1e-14),
            },
            {
                "channel_width_um": 100,
                "channel_length_um": 10,
                "oxide_thickness_nm": 90,
                "dielectric_constant": 3.9,
                "polarity": "n",
            },
            {"transfer": {"preferred_vd_v": -9.0}},
            source_file="IdVg_single_vds.csv",
        )
        sweep = result["sweep_results"][0]
        self.assertEqual(sweep["identity"]["measured_vds_v"], -1.3)
        self.assertEqual(sweep["identity"]["bias_value_v"], -1.3)
        self.assertEqual(sweep["identity"]["vds_reference"], "Vds measured directly")
        self.assertIsNotNone(sweep["vth"]["vth_v"])
        self.assertIsNotNone(sweep["mobility"]["mobility_cm2_vs"])
        self.assertIsNotNone(sweep["subthreshold_swing"]["ss_mv_dec"])
        self.assertIsNotNone(sweep["on_off"]["ion_ioff"])
        self.assertIsNotNone(sweep["gm"]["gm_max_s"])

    def test_gate_sweep_bias_detection_prefers_stepped_vd_over_grounded_vs(self):
        parsed = {
            "columns": ["Vg", "Vd", "Vs", "Id"],
            "data": {
                "Vg": [-1, 0, 1, -1, 0, 1],
                "Vd": [-0.1, -0.1, -0.1, -1.0, -1.0, -1.0],
                "Vs": [0.0] * 6,
                "Id": [-1e-9, -1e-7, -1e-5] * 2,
            },
        }
        classification = classify_measurement(
            {}, parsed["columns"], parsed["data"], "IdVg_multi_bias.csv"
        )
        self.assertEqual(classification["bias_variable"], "Vd")

    def test_measured_vds_uses_vd_minus_vs(self):
        parsed = {
            "columns": ["Vg", "Vd", "Vs", "Id"],
            "data": {
                "Vg": [-1, 0, 1], "Vd": [0.2, 0.2, 0.2],
                "Vs": [1.2, 1.2, 1.2], "Id": [-1e-9, -1e-7, -1e-5],
            },
        }
        classification = {"sweep_variable": "Vg", "bias_variable": "Vd"}
        segment = segment_sweeps(parsed, classification)[0]
        self.assertAlmostEqual(segment.measured_vds_v, -1.0)
        self.assertEqual(segment.vds_reference, "Vd and Vs measured; Vds = Vd - Vs")

    def test_trailing_singleton_bias_is_quarantined_before_direction_and_gm(self):
        first = np.linspace(5.0, -5.0, 101).tolist()
        second = np.linspace(-5.0, 5.0, 101).tolist()
        vg = first + second + [5.0]
        vd = [-0.1] * 202 + [-1.0]
        drain = [-(1.0e-6 + (5.0 - value) * 1.5e-7) for value in first]
        drain += [-(2.5e-6 - (value + 5.0) * 1.0e-7) for value in second]
        drain += [-1.0e-5]
        parsed = {
            "columns": ["Vg", "Vd", "Vs", "Id"],
            "data": {"Vg": vg, "Vd": vd, "Vs": [0.0] * len(vg), "Id": drain},
        }
        classification = {
            "sweep_variable": "Vg", "bias_variable": "Vd",
            "drain_current_raw_column": "Id", "device_polarity": "p",
            "source_filename": "IdVg_trailing_singleton.ztr",
        }
        segments = segment_sweeps(parsed, classification)
        self.assertEqual(len(segments), 2)
        self.assertTrue(all(segment.measured_vds_v == -0.1 for segment in segments))
        self.assertTrue(all(-1.0 not in segment.data["Vd"] for segment in segments))
        excluded = classification["segmentation_diagnostics"]["excluded_bias_runs"]
        self.assertEqual(len(excluded), 1)
        self.assertEqual(excluded[0]["start_index"], 202)
        self.assertEqual(excluded[0]["n_points"], 1)
        self.assertEqual(excluded[0]["measured_vds_v"], -1.0)

    def test_mixed_vds_segment_is_rejected_again_at_extraction_boundary(self):
        segment = SweepSegment(
            direction="forward", bias_level=-0.1, bias_variable="Vd",
            sweep_variable="Vg", sweep_values=[-2, -1, 0, 1, 2],
            data={
                "Vg": [-2, -1, 0, 1, 2],
                "Vd": [-0.1, -0.1, -0.1, -0.1, -1.0],
                "Vs": [0.0] * 5,
                "Id": [-5e-6, -4e-6, -3e-6, -2e-6, -1e-5],
            },
            measured_vds_v=-0.1,
        )
        result = analyze_transfer_sweep(
            segment,
            {"drain_current_raw_column": "Id", "source_filename": "mixed.csv"},
            {"transfer": {"smooth_method": "none", "ss_require_above_gate_leakage": False}},
            {
                "channel_width_um": 100, "channel_length_um": 5,
                "oxide_thickness_nm": 90, "dielectric_constant": 3.9,
                "polarity": "p",
            },
        )
        self.assertFalse(result.gm["bias_homogeneous"])
        self.assertFalse(result.gm["metric_eligible"])
        self.assertIsNone(result.gm["gm_max_s"])
        self.assertFalse(result.mobility["metric_eligible"])
        self.assertIsNone(result.mobility["mobility_cm2_vs"])
        self.assertIn("mixed measured drain bias", " ".join(result.warnings))

    def test_supported_near_boundary_gm_peak_remains_eligible(self):
        vg = [-5.0, -4.9, -4.8, -4.7, -4.6, -4.5]
        result = analyze_transfer(
            {
                "Vg": vg, "Vds": [-0.1] * len(vg),
                "Id": [0.0, 1.0e-6, 1.1e-6, 1.2e-6, 1.3e-6, 1.4e-6],
            },
            {
                "channel_width_um": 100, "channel_length_um": 5,
                "oxide_thickness_nm": 90, "dielectric_constant": 3.9,
                "polarity": "p",
            },
            {"transfer": {"smooth_method": "none", "ss_require_above_gate_leakage": False}},
            source_file="IdVg_valid_boundary.csv",
        )
        sweep = result["sweep_results"][0]
        self.assertEqual(sweep["gm"]["gm_max_vg"], -4.9)
        self.assertEqual(sweep["gm"]["gm_peak_position"], "near_lower_boundary")
        self.assertEqual(sweep["gm"]["gm_peak_support_vg_v"], [-5.0, -4.9, -4.8])
        self.assertEqual(sweep["gm"]["gm_peak_support_vds_v"], [-0.1, -0.1, -0.1])
        self.assertTrue(sweep["gm"]["metric_eligible"])
        self.assertTrue(sweep["mobility"]["metric_eligible"])
        self.assertIsNotNone(sweep["mobility"]["mobility_cm2_vs"])

    def test_small_vds_noise_within_configured_tolerance_remains_one_sweep(self):
        vg = [-2, -1, 0, 1, 2]
        result = analyze_transfer(
            {
                "Vg": vg,
                "Vds": [-0.10001, -0.1, -0.09999, -0.1, -0.10001],
                "Id": [1e-9, 1e-8, 1e-7, 1e-6, 1e-5],
            },
            {
                "channel_width_um": 100, "channel_length_um": 5,
                "oxide_thickness_nm": 90, "dielectric_constant": 3.9,
                "polarity": "n",
            },
            {"transfer": {
                "vds_segmentation_tolerance_v": 5e-5,
                "smooth_method": "none", "ss_require_above_gate_leakage": False,
            }},
            source_file="IdVg_vds_noise.csv",
        )
        self.assertEqual(len(result["sweep_results"]), 1)
        sweep = result["sweep_results"][0]
        self.assertTrue(sweep["gm"]["bias_homogeneous"])
        self.assertAlmostEqual(sweep["identity"]["measured_vds_spread_v"], 2e-5)

    def test_best_mobility_excludes_explicitly_ineligible_candidate(self):
        sweeps = [
            {
                "sweep_id": "valid", "direction": "forward",
                "measured_vds_v": -0.1, "mobility_cm2_vs": 7.5,
                "mobility_metric_eligible": True,
                "mobility_bias_homogeneous": True,
            },
            {
                "sweep_id": "contaminated", "direction": "reverse",
                "measured_vds_v": -0.1, "mobility_cm2_vs": 300.0,
                "mobility_metric_eligible": False,
                "mobility_bias_homogeneous": False,
            },
        ]
        best = build_best_results(sweeps, {})
        self.assertEqual(best["mobility_cm2_vs"]["value"], 7.5)
        excluded = best["mobility_cm2_vs_excluded_candidates"]
        self.assertEqual(excluded[0]["sweep_id"], "contaminated")

    def test_gate_leakage_warning_uses_point_coverage(self):
        drain = [2.0] * 9 + [0.5]
        gate = [1.0] * 10
        result = assess_quality(
            [{"cleaned": {"Id": drain, "Ig": gate}, "indices_kept": list(range(10))}],
            {"sweep_results": [], "analysis_settings": {}},
            {"advanced": {"gate_leakage_id_dominance_fraction": 0.90}},
        )
        self.assertAlmostEqual(result["gate_leakage"]["passing_fraction"], 0.9)
        self.assertNotIn("gate_leakage_dominance", {item["code"] for item in result["flags"]})
        result = assess_quality(
            [{"cleaned": {"Id": [2.0] * 8 + [0.5, 0.5], "Ig": gate}, "indices_kept": list(range(10))}],
            {"sweep_results": [], "analysis_settings": {}},
            {"advanced": {"gate_leakage_id_dominance_fraction": 0.90}},
        )
        self.assertIn("gate_leakage_dominance", {item["code"] for item in result["flags"]})

    def test_open_and_short_checks_are_warning_only_and_auditable(self):
        transfer_metrics = {
            "sweep_results": [{"identity": {"measurement_type": "transfer"}}],
            "analysis_settings": {"noise_floor_a": 1e-13},
            "device_geometry": {"channel_width_um": 10.0},
        }
        open_result = assess_quality(
            [{"cleaned": {"Id": [2e-12] * 12}, "indices_kept": list(range(12))}],
            transfer_metrics, {"advanced": {}},
        )
        self.assertIn("possible_open", {item["code"] for item in open_result["flags"]})
        short_result = assess_quality(
            [{"cleaned": {"Id": [2e-3] * 12}, "indices_kept": list(range(12))}],
            transfer_metrics, {"advanced": {}},
        )
        self.assertIn("possible_short", {item["code"] for item in short_result["flags"]})
        self.assertEqual(short_result["status"], "warning")

    def test_best_result_remarks_retain_bias_sweep_and_read_condition(self):
        best = build_best_results([{
            "sweep_id": "transfer:0:forward:-2V", "direction": "forward",
            "measured_vds_v": -2.0, "ion_const_vg_a": 5e-5,
            "ion_const_vg_requested_v": -4.0, "ion_const_vg_v": -3.9,
            "ion_gate_field_requested_mv_cm": -0.4,
            "ion_gate_field_actual_mv_cm": -0.39, "warnings": "nearest Vg used",
        }], {})
        item = best["ion_const_vg_a"]
        self.assertEqual(item["sweep_type"], "transfer")
        self.assertEqual(item["measured_vds_v"], -2.0)
        self.assertIn("forward", item["remarks"].lower())
        self.assertIn("requested Vg=-4.0", item["remarks"])
        self.assertIn("nearest Vg used", item["remarks"])


class AdaptiveGmSmoothingTests(unittest.TestCase):
    def test_review_audit_flags_gated_semantics_without_changing_them(self):
        from fet_analyzer.analysis.gm_review import review_gm_semantics
        audit = review_gm_semantics()
        self.assertTrue(audit["fixed_gap_isolation"]["current_and_isolated_differ"])
        self.assertEqual(audit["tied_peak"]["currently_selected_vg"], 2.0)
        self.assertEqual(audit["tied_peak"]["equivalent_peak_vg"], [2.0, 3.0, 4.0])
        self.assertTrue(all(not section["production_behavior_changed"] for section in audit.values()))
    @staticmethod
    def adaptive(vg, current, **overrides):
        config = {"smooth_method": "adaptive_savgol", "savgol_order": 1}
        config.update(overrides)
        return compute_gm_with_diagnostics(list(vg), list(current), config)

    def test_clean_curve_selects_smallest_legal_stable_window(self):
        vg = np.linspace(-4, 4, 163)
        current = 1e-6 * np.tanh(vg / 2)
        result = self.adaptive(vg, current)
        region = result["smoothing"]["regions"][0]
        self.assertEqual(region["candidate_windows"], [9, 11, 13, 15])
        self.assertEqual(region["selected_window_points"], 9)
        self.assertEqual(region["selection_status"], "accepted")
        raw_vg, raw_gm = compute_gm(vg.tolist(), current.tolist(), 0, 1)
        selected_peak = np.nanmax(np.abs(result["gm_for_metrics"]))
        self.assertLessEqual(abs(selected_peak - np.nanmax(np.abs(raw_gm))) / selected_peak, 0.05)
        self.assertLessEqual(abs(result["vg"][int(np.nanargmax(np.abs(result["gm_for_metrics"])))]), np.median(np.diff(raw_vg)))

    def test_short_one_candidate_and_two_candidate_fallbacks(self):
        outcomes = {}
        for count in (43, 63, 83):
            vg = np.linspace(-3, 3, count)
            result = self.adaptive(
                vg, 1e-6 * np.tanh(vg / 2),
                **({"adaptive_gm_snr_plateau_tolerance": 0.5} if count == 83 else {}),
            )
            outcomes[count] = result
        unavailable = outcomes[43]
        self.assertFalse(unavailable["smoothing"]["metric_eligible"])
        self.assertTrue(np.isfinite(unavailable["gm"]).any())
        self.assertFalse(np.isfinite(unavailable["gm_for_metrics"]).any())
        one = outcomes[63]["smoothing"]["regions"][0]
        self.assertEqual(one["candidate_windows"], [5])
        self.assertEqual(one["selection_status"], "review")
        two = outcomes[83]["smoothing"]["regions"][0]
        self.assertEqual(two["candidate_windows"], [5, 7])
        self.assertEqual(two["selected_window_points"], 5)
        self.assertEqual(two["selection_status"], "accepted")

    def test_duplicates_irregular_spacing_and_cleaned_gap_are_region_safe(self):
        left = np.linspace(-6, -2, 83)
        right = np.linspace(2, 6, 83) ** 1.01
        vg = np.concatenate([left, [left[20]], right])
        current = 1e-6 * np.tanh(vg / 2)
        result = self.adaptive(vg, current)
        self.assertEqual(len(result["vg"]), len(np.unique(vg)))
        self.assertEqual(len(result["smoothing"]["regions"]), 2)
        self.assertTrue(all(
            region["selected_window_points"] < 0.10 * region["n_usable_points"]
            for region in result["smoothing"]["regions"]
            if region["selected_window_points"] is not None
        ))

    def test_no_plateau_uses_review_fallback_with_auditable_candidates(self):
        rng = np.random.default_rng(7)
        vg = np.linspace(-4, 4, 163)
        current = 1e-6 * np.tanh(vg) + rng.normal(0, 7e-8, len(vg))
        result = self.adaptive(
            vg, current, adaptive_gm_peak_tolerance=0.0,
            adaptive_gm_snr_plateau_tolerance=0.0,
        )
        region = result["smoothing"]["regions"][0]
        self.assertEqual(region["selection_status"], "review")
        self.assertEqual(region["selection_reason"], "best_snr_no_stable_plateau")
        self.assertEqual(sum(bool(item["selected"]) for item in region["candidates"]), 1)

    def test_noise_can_increase_window_and_narrow_peak_guard_is_enforced(self):
        vg = np.linspace(-4, 4, 163)
        rng = np.random.default_rng(1)
        noisy = 1e-6 * np.tanh(vg / 2) + rng.normal(0, 2e-8, len(vg))
        noisy_region = self.adaptive(vg, noisy)["smoothing"]["regions"][0]
        self.assertGreater(noisy_region["selected_window_points"], noisy_region["candidate_windows"][0])

        narrow = 1e-6 * np.tanh(vg / 0.20)
        narrow_region = self.adaptive(vg, narrow)["smoothing"]["regions"][0]
        base_peak = narrow_region["candidates"][0]["gm_peak_s"]
        selected = next(item for item in narrow_region["candidates"] if item["selected"])
        self.assertLessEqual(abs(selected["gm_peak_s"] - base_peak) / base_peak, 0.05)

    def test_n_and_p_polarities_have_matching_adaptive_provenance(self):
        vg = np.linspace(-4, 4, 163)
        n_result = self.adaptive(vg, 1e-6 * np.tanh(vg / 2))
        p_result = self.adaptive(vg, -1e-6 * np.tanh(vg / 2))
        self.assertEqual(
            n_result["smoothing"]["peak_window_points"],
            p_result["smoothing"]["peak_window_points"],
        )
        self.assertAlmostEqual(
            float(np.nanmax(np.abs(n_result["gm_for_metrics"]))),
            float(np.nanmax(np.abs(p_result["gm_for_metrics"]))),
        )

    def test_fixed_and_unsmoothed_wrapper_remain_compatible(self):
        vg = np.linspace(-2, 2, 31).tolist()
        current = (1e-6 * np.tanh(np.asarray(vg))).tolist()
        for window in (0, 5):
            old_vg, old_gm = compute_gm(vg, current, window, 1)
            result = compute_gm_with_diagnostics(vg, current, {
                "smooth_method": "savgol" if window else "none",
                "savgol_window": window, "savgol_order": 1,
            })
            np.testing.assert_allclose(old_vg, result["vg"], equal_nan=True)
            np.testing.assert_allclose(old_gm, result["gm"], equal_nan=True)


class IonBiasConfigurationTests(unittest.TestCase):
    def test_electric_field_overrides_voltage_and_uses_device_oxide(self):
        resolved = resolve_ion_bias(
            {"transfer": {
                "ion_overdrive_v": -3.0,
                "ion_overdrive_field_mv_cm": -2.0,
                "ion_fixed_vg_v": -5.0,
            }},
            {"oxide_thickness_nm": 20.0},
        )
        self.assertEqual(resolved["fixed_vg_v"], -5.0)
        self.assertEqual(resolved["overdrive_input_v"], -3.0)
        self.assertAlmostEqual(resolved["overdrive_v"], -4.0)
        self.assertEqual(resolved["overdrive_source"], "electric_field")
        self.assertEqual(resolved["warnings"], [])

    def test_electric_field_without_oxide_is_flagged_and_not_guessed(self):
        resolved = resolve_ion_bias(
            {"transfer": {"ion_overdrive_field_mv_cm": 1.5}},
            {"oxide_thickness_nm": None},
        )
        self.assertIsNone(resolved["overdrive_v"])
        self.assertEqual(resolved["overdrive_source"], "electric_field")
        self.assertTrue(any("oxide_thickness_nm" in warning for warning in resolved["warnings"]))

    def test_gate_field_resolves_device_specific_fixed_vg(self):
        resolved = resolve_ion_bias(
            {"transfer": {"ion_method": "fixed_gate_field", "ion_gate_field_mv_cm": -2.0}},
            {"oxide_thickness_nm": 20.0},
        )
        self.assertAlmostEqual(resolved["fixed_vg_v"], -4.0)
        self.assertEqual(resolved["fixed_vg_source"], "electric_field")

    def test_optional_gate_field_takes_precedence_over_fixed_vg(self):
        resolved = resolve_ion_bias(
            {"transfer": {
                "ion_method": "maximum_measured", "report_ion_at_fixed_vg": True,
                "ion_fixed_vg_v": 1.0, "ion_gate_field_mv_cm": 2.0,
            }},
            {"oxide_thickness_nm": 20.0},
        )
        self.assertEqual(resolved["fixed_vg_v"], 4.0)
        self.assertEqual(resolved["fixed_vg_source"], "electric_field")

    def test_nearest_voltage_read_records_requested_and_actual(self):
        read = nearest_abs_current_at_voltage([0.0, 2.0], [1e-9, 3e-9], 1.6)
        self.assertEqual(read["requested_v"], 1.6)
        self.assertEqual(read["actual_v"], 2.0)
        self.assertAlmostEqual(read["delta_v"], 0.4)
        self.assertEqual(read["current_a"], 3e-9)

    def test_equidistant_voltage_read_uses_conservative_measured_point(self):
        read = nearest_abs_current_at_voltage(
            [-29.0, -28.0], [3.8e-6, 1.9e-6], -28.5,
        )
        self.assertEqual(read["requested_v"], -28.5)
        self.assertEqual(read["actual_v"], -28.0)
        self.assertEqual(read["current_a"], 1.9e-6)
        self.assertEqual(read["delta_v"], 0.5)
        self.assertFalse(read["exact"])
        self.assertTrue(read["equidistant_tie"])
        self.assertEqual(read["candidate_voltages_v"], [-29.0, -28.0])
        self.assertIn("lowest median |Id|", read["selection_reason"])

    def test_equidistant_voltage_read_handles_duplicates_deterministically(self):
        read = nearest_abs_current_at_voltage(
            [-29.0, -29.0, -28.0, -28.0],
            [4.0e-6, 6.0e-6, 2.0e-6, 4.0e-6],
            -28.5,
        )
        self.assertEqual(read["actual_v"], -28.0)
        self.assertEqual(read["current_a"], 3.0e-6)

    def test_fixed_gate_field_half_step_runs_without_maximum_current_fallback(self):
        vg = [-30.0, -29.0, -28.0, -27.0, -26.0]
        result = analyze_transfer(
            {
                "Vg": vg,
                "Vds": [-0.1] * len(vg),
                "Id": [-5.0e-6, -4.0e-6, -2.0e-6, -1.0e-6, -0.5e-6],
                "Ig": [1.0e-12] * len(vg),
            },
            {
                "channel_width_um": 100,
                "channel_length_um": 5,
                "oxide_thickness_nm": 95,
                "dielectric_constant": 4.06,
                "polarity": "p",
            },
            {
                "transfer": {
                    "ion_method": "fixed_gate_field",
                    "ion_gate_field_mv_cm": -3.0,
                    "report_ion_at_fixed_vg": True,
                    "smooth_method": "none",
                }
            },
            source_file="IdVg_half_step.ztr",
        )
        sweep = result["sweep_results"][0]
        current = sweep["current_summary"]
        self.assertEqual(current["ion_const_vg_requested_v"], -28.5)
        self.assertEqual(current["ion_const_vg_v"], -28.0)
        self.assertEqual(current["ion_const_vg_a"], 2.0e-6)
        self.assertTrue(current["ion_const_vg_equidistant_tie"])
        self.assertEqual(current["ion_const_vg_candidate_voltages_v"], [-29.0, -28.0])
        self.assertIn("lower |Id|", " ".join(sweep["warnings"]))
        preferred = result["device_summary"]["preferred"]
        self.assertEqual(preferred["ion_method_used"], "constant_vg")
        self.assertEqual(preferred["ion_configured_a"], 2.0e-6)

    def test_compliance_cleaning_is_opt_in_by_default(self):
        self.assertFalse(DEFAULT_CONFIG["advanced"]["compliance_cleaning_enabled"])

    def test_config_editing_warns_and_readiness_rejects_missing_fixed_voltage(self):
        config = deep_merge(DEFAULT_CONFIG, {"transfer": {"ion_method": "fixed_vg", "ion_fixed_vg_v": None}})
        report = validate_config(config)
        self.assertTrue(any(issue.code == "missing_fixed_vg" for issue in report.warnings))
        from fet_analyzer.validation import validate_analysis_readiness
        readiness = validate_analysis_readiness(config)
        self.assertTrue(any(issue.code == "missing_fixed_vg" for issue in readiness.errors))

    def test_config_validation_rejects_nonpositive_film_thickness(self):
        config = deep_merge(DEFAULT_CONFIG, {"device_defaults": {"film_thickness_nm": 0}})
        report = validate_config(config)
        self.assertTrue(any(issue.code == "invalid_film_thickness" for issue in report.errors))

    def test_config_validation_rejects_advertised_but_unimplemented_methods(self):
        cases = [
            {"transfer": {"smooth_method": "gaussian"}},
            {"transfer": {"vth_method": "second_derivative"}},
            {"transfer": {"vth_method": "max_gm_point"}},
            {"transfer": {"ss_method": "linear_fit"}},
            {"tlm": {"tlm_ss_method": "average"}},
        ]
        for override in cases:
            with self.subTest(override=override):
                report = validate_config(deep_merge(DEFAULT_CONFIG, override))
                self.assertTrue(any(issue.code == "method_not_implemented" for issue in report.errors))

    def test_strict_preflight_suppresses_only_dependent_metrics(self):
        config = deep_merge(DEFAULT_CONFIG, {"advanced": {"parameter_preflight_mode": "strict"}})
        preflight = assess_parameter_preflight(
            {"channel_width_um": 100, "channel_length_um": 10, "oxide_thickness_nm": 20, "dielectric_constant": 3.9},
            {"channel_width_um": "confirmed", "channel_length_um": "confirmed", "oxide_thickness_nm": "template_default", "dielectric_constant": "confirmed"},
            config,
        )
        metrics = {"summary": {"mobility_cm2_vs_max": 5.0}, "sweep_results": [{"mobility": {"mobility_cm2_vs": 5.0}, "current_summary": {"ion_const_vov_a": 1e-6}}]}
        apply_preflight_suppression(metrics, preflight)
        self.assertIsNone(metrics["summary"]["mobility_cm2_vs_max"])
        self.assertIsNone(metrics["sweep_results"][0]["current_summary"]["ion_const_vov_a"])

    def test_canonical_schema_and_in_memory_api(self):
        result = analyze_transfer(
            {"Vg": [-2, -1, 0, 1, 2], "Vd": [0.1] * 5, "Id": [1e-9, 1e-8, 1e-7, 1e-6, 1e-5]},
            {"channel_width_um": 100, "channel_length_um": 10, "oxide_thickness_nm": 20, "dielectric_constant": 3.9, "polarity": "n"},
            {"transfer": {"ss_require_above_gate_leakage": False}},
        )
        self.assertEqual(result["schema_version"], SCHEMA_VERSION)
        self.assertEqual(SCHEMA_VERSION, "5.0")
        self.assertEqual(validate_canonical_result(result["canonical_result"]), [])
        self.assertIn("current_ua_per_um", result["canonical_result"]["ion"])
        self.assertIn("fixed_gate_current_a", result["canonical_result"]["ion"])
        self.assertIn("fixed_overdrive_current_a", result["canonical_result"]["ion"])
        self.assertIn("gm_diagnostics", result["canonical_result"])

    def test_canonical_consistency_detects_renderer_drift(self):
        from fet_analyzer.schema import attach_canonical_result, validate_metrics_consistency
        metrics = {
            "analysis_type": "transfer",
            "device_summary": {"preferred": {"ion_configured_a": 1e-6}},
        }
        attach_canonical_result(metrics)
        self.assertEqual(validate_metrics_consistency(metrics), [])
        metrics["device_summary"]["preferred"]["ion_configured_a"] = 2e-6
        self.assertTrue(validate_metrics_consistency(metrics))

    def test_true_maximum_common_overdrive_is_group_resolved(self):
        def record(name, low):
            vg = list(np.linspace(low, 2, 11))
            segment = SweepSegment("forward", 0.1, "Vd", "Vg", vg, {"Vg": vg, "Id": [10 ** (v / 2 - 8) for v in vg]})
            sweep = {"identity": {"measurement_type": "transfer", "sweep_index": 0, "direction": "forward", "bias_value_v": 0.1}, "vth": {"vth_v": 0.0}, "current_summary": {"id_max_a": 1e-5}, "on_off": {"ion_a": 1e-5, "ioff_a": 1e-9}, "mobility": {}, "subthreshold_swing": {}, "gm": {}, "warnings": []}
            return {"source_file": Path(name), "segments": [segment], "classification": {"drain_current_raw_column": "Id", "filename_info": {"sample_label": "S"}}, "device_params": {"polarity": "p"}, "metrics": {"analysis_settings": {"ion_method_configured": "max_common_overdrive"}, "sweep_results": [sweep]}}
        records = [record("a.csv", -8), record("b.csv", -6), record("c.csv", -5)]
        audit = apply_max_common_overdrive(records, deep_merge(DEFAULT_CONFIG, {"transfer": {"ion_method": "max_common_overdrive"}}))
        self.assertEqual(audit[0]["chosen_v"], -5.0)
        self.assertTrue(all(item["metrics"]["device_summary"]["preferred"]["ion_const_vov_v"] == -5.0 for item in records))


class MasterSummaryTests(unittest.TestCase):
    def test_extended_summary_is_primary_transfer_sweep_table(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            sweeps = [
                {
                    "sweep_id": "transfer:0:forward:-0.25V", "direction": "forward",
                    "bias_value_v": -0.25, "measured_vd_v": -0.2,
                    "measured_vs_v": 0.1, "measured_vds_v": -0.3,
                    "vds_reference": "Vd and Vs measured; Vds = Vd - Vs",
                    "vth_v": -1.0, "ss_min_mv_dec": 90.0,
                    "gm_max_s": 2e-6, "mobility_cm2_vs": 4.0,
                    "gm_metric_eligible": True, "gm_bias_homogeneous": True,
                    "mobility_metric_eligible": True, "mobility_bias_homogeneous": True,
                    "ion_configured_a": 1e-5, "ioff_configured_a": 1e-9,
                    "ion_ioff": 1e4, "warnings": "",
                },
                {
                    "sweep_id": "transfer:1:reverse:-1V", "direction": "reverse",
                    "bias_value_v": -1.0, "measured_vds_v": -1.0,
                    "vth_v": -1.2, "ss_min_mv_dec": 95.0,
                    "gm_max_s": 3e-6, "mobility_cm2_vs": 5.0,
                    "gm_metric_eligible": True, "gm_bias_homogeneous": True,
                    "mobility_metric_eligible": True, "mobility_bias_homogeneous": True,
                    "ion_configured_a": 2e-5, "ioff_configured_a": 2e-9,
                    "ion_ioff": 1e4, "warnings": "review this sweep",
                },
            ]
            payload = {
                "sample_id": "SampleB", "analysis_type": "transfer",
                "measurement_type": "IdVg", "analysis_settings": {},
                "device_geometry": {
                    "channel_length_um": 12.5, "channel_width_um": 250.0,
                    "oxide_thickness_nm": 20.0, "dielectric_constant": 4.0,
                },
                "quality": {"status": "pass", "flags": []},
                "device_summary": {
                    "preferred": {
                        **sweeps[1], "preferred_direction": "reverse",
                        "preferred_bias_v": -1.0, "selection_reason": "configured preferred sweep",
                    },
                    "segments": sweeps,
                    "best": {"mobility_cm2_vs": {"value": 5.0, "remarks": "best sweep"}},
                },
            }
            (root / "DeviceA_metrics.json").write_text(json.dumps(payload), encoding="utf-8")
            (root / "OnlyOutput_metrics.json").write_text(json.dumps({
                "sample_id": "SampleB", "analysis_type": "output",
                "measurement_type": "IdVd", "device_summary": {
                    "preferred": {"direction": "forward", "bias_value_v": -2.0},
                    "segments": [], "best": {},
                },
            }), encoding="utf-8")

            extended = collect_extended_metrics(root)
            self.assertEqual(len(extended), 2)
            self.assertEqual([row["sweep_id"] for row in extended], [1, 2])
            self.assertEqual([row["preferred_direction"] for row in extended], ["forward", "reverse"])
            self.assertAlmostEqual(extended[0]["vds_v"], -0.3)
            self.assertAlmostEqual(extended[0]["preferred_bias_v"], -0.3)
            self.assertTrue(all(row["channel_length_um"] == 12.5 for row in extended))
            self.assertTrue(all(row["channel_width_um"] == 250.0 for row in extended))
            self.assertEqual(extended[1]["mobility_cm2_vs"], 5.0)
            self.assertEqual(extended[0]["best_mobility_cm2_vs"], 5.0)
            self.assertNotIn("OnlyOutput", {row["device"] for row in extended})

            rows, _ = collect_metrics(root)
            transfer = next(row for row in rows if row["type"] == "transfer")
            output = next(row for row in rows if row["type"] == "output")
            self.assertEqual(transfer["sweep_id"], 2)
            self.assertEqual(transfer["vth_v"], -1.2)
            self.assertEqual(transfer["selection_reason"], "configured preferred sweep")
            self.assertEqual(transfer["channel_length_um"], 12.5)
            self.assertEqual(transfer["channel_width_um"], 250.0)
            self.assertIsNone(output["channel_length_um"])
            self.assertIsNone(output["channel_width_um"])

            generate_batch_summary(root, {"execution": {"generate_batch_plots": False}})
            summary_dir = root / "batch_summary"
            csv_path = summary_dir / "extended_master_summary.csv"
            xlsx_path = summary_dir / "extended_master_summary.xlsx"
            self.assertTrue(csv_path.exists())
            self.assertTrue(xlsx_path.exists())
            with csv_path.open(newline="", encoding="utf-8") as handle:
                csv_rows = list(csv.DictReader(handle))
            self.assertEqual(len(csv_rows), 2)
            self.assertEqual(csv_rows[0]["sweep_id"], "1")
            self.assertEqual(csv_rows[0]["channel_length_um"], "12.5")
            self.assertEqual(csv_rows[0]["channel_width_um"], "250.0")
            self.assertEqual(
                [key for key, _, _ in EXTENDED_BATCH_COLUMNS if key not in {"sweep_id", "vds_v"}],
                [key for key, _, _ in BATCH_COLUMNS],
            )
            import openpyxl
            workbook = openpyxl.load_workbook(xlsx_path, read_only=True)
            sheet = workbook["Extended FET Summary"]
            headers = next(sheet.iter_rows(values_only=True))
            self.assertEqual(sheet.max_row - 1, len(csv_rows))
            self.assertIn("Sweep ID", headers)
            self.assertIn("Vds (V)", headers)
            self.assertIn("Lch (μm)", headers)
            self.assertIn("Wch (μm)", headers)
            workbook.close()

            master_csv_path = summary_dir / "master_summary.csv"
            with master_csv_path.open(newline="", encoding="utf-8") as handle:
                master_rows = list(csv.DictReader(handle))
            transfer_csv = next(row for row in master_rows if row["type"] == "transfer")
            output_csv = next(row for row in master_rows if row["type"] == "output")
            self.assertEqual(transfer_csv["channel_length_um"], "12.5")
            self.assertEqual(transfer_csv["channel_width_um"], "250.0")
            self.assertEqual(output_csv["channel_length_um"], "")
            self.assertEqual(output_csv["channel_width_um"], "")

            master_workbook = openpyxl.load_workbook(
                summary_dir / "master_summary.xlsx", read_only=True,
            )
            master_headers = next(master_workbook["FET Summary"].iter_rows(values_only=True))
            self.assertIn("Lch (μm)", master_headers)
            self.assertIn("Wch (μm)", master_headers)
            master_workbook.close()

    def test_batch_row_includes_all_selected_tlm_operating_points(self):
        row = {"type": "transfer"}
        tlm_row = {
            "analysis_level": "master_aggregate", "tlm_id": "MASTER",
            "transfer_rcw_ohm_um": "12000", "transfer_best_r2": "0.98",
            "transfer_zero_vg_v": "0.1", "transfer_zero_rcw_ohm_um": "9000",
            "transfer_fixed_source": "electric_field",
            "transfer_fixed_requested_vg_v": "-4.0",
            "transfer_fixed_requested_field_mv_cm": "-2.0",
            "transfer_fixed_field_oxide_thickness_nm": "20.0",
            "transfer_fixed_vg_v": "-4.0", "transfer_fixed_rcw_ohm_um": "6000",
            "transfer_max_vg_v": "-5.0", "transfer_max_rcw_ohm_um": "5000",
            "transfer_max_selection_reason": "largest |measured/common Vg| with a TLM fit",
        }
        _add_tlm_summary_metrics(row, tlm_row)
        self.assertEqual(row["rcw_ohm_um"], 12000.0)
        self.assertEqual(row["tlm_zero_rcw_kohm_um"], 9.0)
        self.assertEqual(row["tlm_fixed_source"], "electric_field")
        self.assertEqual(row["tlm_fixed_requested_vg_v"], -4.0)
        self.assertEqual(row["tlm_fixed_rcw_kohm_um"], 6.0)
        self.assertEqual(row["tlm_max_vg_v"], -5.0)
        self.assertEqual(row["tlm_max_rcw_kohm_um"], 5.0)
        self.assertEqual(row["tlm_max_analysis_level"], "master_aggregate")

    def test_width_normalized_currents_and_optional_dual_ion_reporting(self):
        result = {
            "identity": {"measurement_type": "transfer", "direction": "forward", "bias_value_v": 0.1},
            "on_off": {"ion_a": 20e-6, "ioff_a": 2e-9},
            "current_summary": {
                "id_max_a": 20e-6, "id_min_a": 1e-9,
                "ion_const_vg_a": 12e-6, "ion_const_vg_requested_v": 2.0,
                "ion_const_vg_v": 1.9, "ion_const_vov_a": 15e-6,
                "ion_const_vov_requested_v": 3.0, "ion_const_vov_v": 2.9,
                "ion_const_vov_actual_vg_v": 4.1,
            },
        }
        settings = {
            "channel_width_um": 10.0, "ion_method_configured": "maximum_measured",
            "ioff_method_configured": "minimum_above_ig",
            "report_ion_at_fixed_vg": True, "report_ion_at_fixed_overdrive": True,
        }
        values = transfer_sweep_metrics(result, settings)
        self.assertAlmostEqual(values["ion_configured_ua_per_um"], 2.0)
        self.assertAlmostEqual(values["ioff_configured_ua_per_um"], 2e-4)
        self.assertAlmostEqual(values["ion_const_vg_ua_per_um"], 1.2)
        self.assertAlmostEqual(values["ion_const_vov_ua_per_um"], 1.5)
        self.assertEqual(values["ion_const_vg_v"], 1.9)
        self.assertEqual(values["ion_const_vov_v"], 2.9)

        hidden = transfer_sweep_metrics(result, {
            **settings, "report_ion_at_fixed_vg": False,
            "report_ion_at_fixed_overdrive": False,
        })
        self.assertIsNone(hidden["ion_const_vg_a"])
        self.assertIsNone(hidden["ion_const_vov_a"])
        self.assertEqual(hidden["ion_configured_a"], 20e-6)

    def test_configured_ion_and_preferred_sweep_selection(self):
        ion, method, note = choose_ion(
            {"ion_max_a": 9e-6, "ion_const_vg_a": 4e-6}, "constant_vg"
        )
        self.assertEqual((ion, method, note), (4e-6, "constant_vg", None))
        sweeps = [
            {"sweep_id": "r01", "direction": "reverse", "bias_value_v": 0.1},
            {"sweep_id": "f10", "direction": "forward", "bias_value_v": 1.0},
            {"sweep_id": "f01", "direction": "forward", "bias_value_v": 0.1},
        ]
        selected, reason = select_preferred_sweep(sweeps, {"preferred_vd_v": 0.1})
        self.assertEqual(selected["sweep_id"], "f01")
        self.assertIn("preferred", reason)

    def test_includes_configured_ion_currents_and_ioff_above_ig(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            plots = root / "Sample_Device" / "plots"
            plots.mkdir(parents=True)
            payload = {
                "sample_id": "Sample",
                "measurement_type": "transfer",
                "analysis_settings": {
                    "ion_constant_vg_v": 2.0,
                    "ion_overdrive_v": 1.5,
                    "ioff_leakage_factor": 3.0,
                    "ioff_method_used": "minimum_abs_id_above_gate_leakage",
                },
                "segments": {
                    "0.1": {"ion_const_vg_a": 1e-5, "ion_const_vg_v": 2.0,
                            "ion_const_overdrive_a": 8e-6, "ion_const_overdrive_v": 1.5,
                            "ioff_a": 4e-11},
                    "1.0": {"ion_const_vg_a": 2e-5, "ion_const_vg_v": 2.0,
                            "ion_const_overdrive_a": 1.5e-5, "ion_const_overdrive_v": 1.5,
                            "ioff_a": 2e-11},
                },
            }
            (plots / "transfer_metrics.json").write_text(json.dumps(payload), encoding="utf-8")
            rows, by_sample = collect_metrics(root)
            self.assertEqual(rows[0]["ion_const_vov_a"], 1.5e-5)
            self.assertEqual(rows[0]["ion_const_vg_a"], 2e-5)
            self.assertEqual(rows[0]["ioff_above_ig_a"], 2e-11)
            self.assertEqual(rows[0]["ioff_leakage_factor"], 3.0)

            csv_path = root / "master_summary.csv"
            xlsx_path = root / "master_summary.xlsx"
            write_summary_csv(rows, by_sample, csv_path)
            write_summary_excel(rows, xlsx_path)
            self.assertIn("ioff_above_ig_a", csv_path.read_text(encoding="utf-8").splitlines()[0])
            import openpyxl
            workbook = openpyxl.load_workbook(xlsx_path, read_only=True)
            headers = next(workbook["FET Summary"].iter_rows(values_only=True))
            self.assertIn("Ioff above Ig (A)", headers)
            workbook.close()


class ClassificationTests(unittest.TestCase):
    def test_metadata_name_does_not_override_filename_or_data(self):
        result = classify_measurement(
            {"measurement_type": "Output", "device_id": "wrong-name"},
            ["Vg", "Vd", "Id"],
            {"Vg": [-1, 0, 1], "Vd": [0.1, 0.1, 0.1], "Id": [1e-9, 1e-7, 1e-5]},
            filename="IdVg__Sample_Device.csv",
            filename_patterns={
                "active": "default",
                "patterns": {
                    "default": {
                        "regex": r"(?P<measurement_type>[^_]+)__(?P<sample_label>[^_]+)"
                    }
                },
            },
        )
        self.assertEqual(result["type"], MeasurementType.TRANSFER)
        self.assertEqual(result["filename_info"]["sample_label"], "Sample")
        self.assertEqual(result["filename_info"]["naming_source"], "filename")

    def test_nonstandard_filename_recovers_identity_from_csv_header(self):
        result = classify_measurement(
            {"setup_title": "IdVg", "device_id": "K7a_SnO_5mtorr_TLM1_25um", "count": "1"},
            ["Vbg", "Vd", "Id"],
            {"Vbg": [-1, 0, 1], "Vd": [0.1, 0.1, 0.1], "Id": [1e-9, 1e-7, 1e-5]},
            filename="IdVg [K7a_SnO_5mtorr_TLM1_25um(1) ; date].csv",
            filename_patterns={
                "active": "default",
                "patterns": {"default": {"regex": r"(?P<measurement_type>[^_]+)__(?P<sample_label>[^_]+)"}},
            },
        )
        info = result["filename_info"]
        self.assertEqual(result["type"], MeasurementType.TRANSFER)
        self.assertEqual(info["measurement_type"], "IdVg")
        self.assertEqual(info["sample_label"], "K7a_SnO_5mtorr")
        self.assertEqual(info["device_type"], "TLM1")
        self.assertEqual(info["channel_length_um"], "25")
        self.assertEqual(info["measurement_count"], "1")
        self.assertEqual(info["naming_source"], "csv_header")


class CliHelpTests(unittest.TestCase):
    def test_analyze_help_includes_gui_launch_syntax(self):
        from fet_analyzer.cli import build_parser

        help_text = build_parser().format_help()
        self.assertIn("Interfaces:", help_text)
        self.assertIn("python -m fet_analyzer.dashboard --root .", help_text)
        self.assertIn("fet-dashboard --root .", help_text)
        self.assertIn("python -m fet_analyzer.native --root .", help_text)
        self.assertIn("fet-native --root .", help_text)
        self.assertIn("cd <measurement-folder>", help_text)
        self.assertIn('python -m pip install -e ".[native]"', help_text)
        self.assertNotIn("--analysis-mode", help_text)
        self.assertNotIn("C-V", help_text)
        self.assertIn("--save-plots", help_text)
        self.assertIn("--batch-plots", help_text)


class ReportingLayoutTests(unittest.TestCase):
    def test_flat_output_paths_are_named_from_input_and_plots_are_optional(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            paths = create_output_structure(root, "device one", root / "device one.csv")
            self.assertEqual(paths["analysis_xlsx"], root / "device_one.xlsx")
            self.assertEqual(paths["report"], root / "device_one_report.html")
            self.assertFalse(paths["plots"].exists())
            kept = create_output_structure(root, "device two", root / "device two.csv", keep_plots=True)
            self.assertTrue(kept["plots"].is_dir())

    def test_html_report_embeds_png_and_writes_no_markdown(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            png = root / "figure.png"
            png.write_bytes(b"\x89PNG\r\n\x1a\nexample")
            report = root / "report.html"
            write_html_report(report, "Test", ["# Test", "", "## Plots"], [png])
            rendered = report.read_text(encoding="utf-8")
            self.assertIn("data:image/png;base64,", rendered)
            self.assertNotIn("plot-preview", rendered)
            self.assertNotIn("data-plot-preview", rendered)
            self.assertFalse(report.with_suffix(".md").exists())

    def test_device_html_reports_requested_and_actual_ion_bias(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            report = root / "device_report.html"
            generate_device_report(
                device_name="device",
                classification={"type": MeasurementType.TRANSFER, "confidence": 1.0},
                metrics={"device_summary": {"segments": [{
                    "direction": "forward", "bias_value_v": 0.1,
                    "ion_const_vg_requested_v": 1.6,
                    "ion_const_vg_v": 2.0,
                    "ion_gate_field_requested_mv_cm": 0.8,
                    "ion_gate_field_actual_mv_cm": 1.0,
                }]}},
                parsed_metadata={},
                plot_dir=root / "plots",
                output_path=report,
            )
            rendered = report.read_text(encoding="utf-8")
            self.assertIn("Ion Read Conditions (Requested vs Actual)", rendered)
            self.assertIn("1.6", rendered)
            self.assertIn("2", rendered)


class ParameterPrecedenceTests(unittest.TestCase):
    def test_filename_overrides_metadata(self):
        values, sources = infer_geometry(
            Path("IdVg__S_TLM1_10um__1.csv"),
            {"channel_length_um": 99, "channel_width_um": 25},
        )
        self.assertEqual(values["channel_length_um"], 10)
        self.assertTrue(sources["channel_length_um"].startswith("filename:"))
        self.assertEqual(values["channel_width_um"], 25)
        self.assertTrue(sources["channel_width_um"].startswith("metadata:"))

    def test_device_parameter_table_is_highest_priority(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "device_parameters.txt"
            path.write_text(
                "sample_label\tdevice_pattern\tchannel_length_um\n"
                "S\t*\t42\n",
                encoding="utf-8",
            )
            rows = load_master_table(path)
            values, sources, _ = resolve_device_parameters(
                {"channel_length_um": 1},
                rows,
                "S",
                "IdVg__S_TLM1_10um__1.csv",
                inferred={"channel_length_um": 10},
                inferred_sources={"channel_length_um": "filename"},
            )
        self.assertEqual(values["channel_length_um"], 42)
        self.assertTrue(sources["channel_length_um"].startswith("device_parameters.txt"))


class SweepTests(unittest.TestCase):
    def test_transfer_sweep_records_field_resolved_overdrive(self):
        vg = [-8, -7, -6, -5, -4, -3, -2, -1, 0]
        current = [10 ** ((-value - 8) / 2 - 9) for value in vg]
        segment = SweepSegment(
            direction="forward",
            bias_level=-0.1,
            bias_variable="Vd",
            sweep_variable="Vg",
            sweep_values=vg,
            data={"Vg": vg, "Vd": [-0.1] * len(vg), "Vs": [0.0] * len(vg), "Id": current},
        )
        result = analyze_transfer_sweep(
            segment,
            {"drain_current_raw_column": "Id"},
            {"transfer": {
                "ion_overdrive_v": -2.0,
                "ion_overdrive_field_mv_cm": -2.0,
                "ss_require_above_gate_leakage": False,
            }},
            {
                "channel_width_um": 100,
                "channel_length_um": 10,
                "oxide_thickness_nm": 20,
                "dielectric_constant": 3.9,
                "polarity": "p",
            },
        )
        self.assertAlmostEqual(result.current_summary["ion_const_vov_requested_v"], -4.0)
        self.assertAlmostEqual(
            result.current_summary["ion_const_vov_target_vg_v"],
            result.vth["vth_v"] - 4.0,
        )
        self.assertEqual(result.current_summary["ion_const_vov_actual_vg_v"], -7.0)
        self.assertAlmostEqual(
            result.current_summary["ion_const_vov_v"],
            -7.0 - result.vth["vth_v"],
        )
        self.assertEqual(result.current_summary["ion_overdrive_source"], "electric_field")
        self.assertEqual(result.current_summary["ion_overdrive_input_v"], -2.0)

    def test_transfer_is_split_by_bias_and_direction(self):
        vg_one = [-2, -1, 0, 1, 2, 1, 0, -1, -2]
        vg = vg_one + vg_one
        parsed = {
            "columns": ["Vg", "Vd", "Vs", "Id"],
            "data": {
                "Vg": vg,
                "Vd": [0.1] * len(vg_one) + [1.0] * len(vg_one),
                "Vs": [0.0] * len(vg),
                "Id": [10 ** (value / 2 - 8) for value in vg],
            },
        }
        classification = {
            "sweep_variable": "Vg",
            "bias_variable": "Vd",
            "drain_current_raw_column": "Id",
            "device_polarity": "n",
            "source_filename": "IdVg__S_Device.csv",
        }
        segments = segment_sweeps(parsed, classification)
        self.assertEqual({round(s.bias_level, 1) for s in segments}, {0.1, 1.0})
        self.assertEqual({s.direction for s in segments}, {"forward", "reverse"})
        results = analyze_transfer_sweeps(
            segments,
            classification,
            {"transfer": {}},
            {
                "channel_width_um": 100,
                "channel_length_um": 10,
                "oxide_thickness_nm": 90,
                "dielectric_constant": 3.9,
                "polarity": "n",
            },
        )
        self.assertEqual(len(results), len(segments))
        self.assertEqual({r.identity.direction for r in results}, {"forward", "reverse"})
        self.assertEqual(len({r.identity.sweep_id for r in results}), len(results))

    def test_ss_requires_one_decade_and_average_requires_two(self):
        vg = np.linspace(0, 1, 21).tolist()
        under_one = (10 ** np.linspace(-9, -8.2, 21)).tolist()
        two_decades = (10 ** np.linspace(-10, -7.5, 21)).tolist()
        self.assertIsNone(compute_ss(vg, under_one)["ss_mv_dec"])
        result = compute_ss(vg, two_decades)
        self.assertIsNotNone(result["ss_mv_dec"])
        self.assertIsNotNone(result["ss_avg_mv_dec"])
        self.assertGreaterEqual(result["ss_decades"], 1.0)
        self.assertGreaterEqual(result["ss_avg_decades"], 2.0)

    def test_ss_uses_only_points_above_gate_leakage_and_reports_units(self):
        vg = [0, 1, 2, 3, 4, 5, 6, 7]
        drain = [1e-12, 1e-11, 1e-10, 1e-9, 1e-8, 1e-7, 1e-6, 1e-5]
        gate = [2e-12, 2e-11, 2e-10, 1e-10, 1e-9, 1e-8, 1e-7, 1e-6]
        result = compute_ss(
            vg, drain, gate_leakage_values=gate,
            require_above_gate_leakage=True,
        )
        self.assertAlmostEqual(result["ss_mv_dec"], 1000.0, places=6)
        self.assertAlmostEqual(result["ss_fit_slope_v_dec"], 1.0, places=6)
        self.assertEqual(result["ss_unit"], "mV/dec")
        self.assertEqual(result["ss_fit_indices"], [3, 4, 5])
        self.assertEqual(result["ss_excluded_below_ig"], 3)

    def test_ss_is_unavailable_when_gate_leakage_is_required_but_missing(self):
        result = compute_ss(
            [0, 1, 2, 3, 4], [1e-12, 1e-11, 1e-10, 1e-9, 1e-8],
            require_above_gate_leakage=True,
        )
        self.assertIsNone(result["ss_mv_dec"])
        self.assertTrue(result["warnings"])

    def test_ss_does_not_bridge_points_removed_by_leakage_filter(self):
        vg = list(range(10))
        drain = (10 ** np.array([0, 0.5, 1, 1.5, 2, 3, 2.5, 2, 1.5, 1])).tolist()
        gate = [value / 10 for value in drain]
        gate[4] = drain[4] * 2
        result = compute_ss(vg, drain, gate_leakage_values=gate)
        indices = result["ss_fit_indices"]
        self.assertNotIn(4, indices)
        self.assertTrue(all(b - a == 1 for a, b in zip(indices, indices[1:])))

    def test_ss_rejects_non_monotonic_fit_windows(self):
        current = [1e-10, 1e-9, 1e-8, 1e-9, 1e-10]
        result = compute_ss(
            [0, 1, 2, 3, 4],
            current,
            min_decades=1.5,
        )
        selected = np.log10(np.asarray(current)[result["ss_fit_indices"]])
        self.assertTrue(np.all(np.diff(selected) >= 0) or np.all(np.diff(selected) <= 0))
        self.assertLess(len(result["ss_fit_indices"]), len(current))

    def test_ss_discards_wrong_side_of_ioff_for_device_polarity(self):
        vg = [-2, -1, 0, 1, 2]
        current = [1e-7, 1e-8, 1e-10, 1e-9, 1e-8]
        pfet = compute_ss(vg, current, polarity="p")
        nfet = compute_ss(vg, current, polarity="n")
        self.assertTrue(all(vg[index] <= 0 for index in pfet["ss_fit_indices"]))
        self.assertTrue(all(vg[index] >= 0 for index in nfet["ss_fit_indices"]))
        self.assertEqual(pfet["ss_ioff_vg"], 0.0)
        self.assertEqual(nfet["ss_ioff_vg"], 0.0)

    def test_optimized_ss_matches_exhaustive_polyfit(self):
        rng = np.random.default_rng(42)
        vg = np.linspace(-4, 5, 73)
        log_current = -10 + 0.7 * vg + rng.normal(0, 0.025, len(vg))
        current = 10 ** log_current
        expected = None
        for start in range(len(vg) - 2):
            for end in range(start + 2, len(vg)):
                window = log_current[start:end + 1]
                decades = float(np.max(window) - np.min(window))
                if decades < 1.0 or np.std(window) <= 1e-15:
                    continue
                slope, _ = np.polyfit(window, vg[start:end + 1], 1)
                value = abs(float(slope)) * 1000.0
                if expected is None or value < expected:
                    expected = value
        actual = compute_ss(vg.tolist(), current.tolist())["ss_mv_dec"]
        self.assertIsNotNone(actual)
        self.assertAlmostEqual(actual, expected, places=8)

    def test_output_resistance_and_no_gate_bias(self):
        vd = [-0.1, -0.05, 0.0, 0.05, 0.1]
        resistance = 2000.0
        segment = SweepSegment(
            direction="forward",
            bias_level=None,
            bias_variable=None,
            sweep_variable="Vd",
            sweep_values=vd,
            data={"Vd": vd, "Id": [value / resistance for value in vd]},
        )
        result = analyze_output_sweep(
            segment,
            {"drain_current_raw_column": "Id"},
            {"tlm": {"tlm_output_vd_max": 0.1}},
            source_filename="Output__S_Device.csv",
        )
        self.assertTrue(result.no_gate_bias)
        self.assertIsNone(result.gate_bias_v)
        self.assertAlmostEqual(result.resistance_avg_ohm, resistance)
        self.assertAlmostEqual(result.resistance_linear_fit_ohm, resistance)
        reverse = analyze_output_sweep(
            SweepSegment(
                direction="reverse",
                bias_level=None,
                bias_variable=None,
                sweep_variable="Vd",
                sweep_values=list(reversed(vd)),
                data={
                    "Vd": list(reversed(vd)),
                    "Id": [value / resistance for value in reversed(vd)],
                },
            ),
            {"drain_current_raw_column": "Id"},
            {"tlm": {"tlm_output_vd_max": 0.1}},
            sweep_index=1,
        )
        summary = summarize_resistance_by_gate_bias([result, reverse])
        self.assertEqual(summary["no_gate_bias"]["n_sweeps"], 2)
        self.assertEqual(summary["no_gate_bias"]["directions"], ["forward", "reverse"])
        self.assertAlmostEqual(summary["no_gate_bias"]["resistance_avg_ohm"], resistance)

    def test_gate_leakage_ratio_is_ig_over_id(self):
        vg = [-2, -1, 0, 1, 2]
        segment = SweepSegment(
            direction="forward",
            bias_level=0.1,
            bias_variable="Vd",
            sweep_variable="Vg",
            sweep_values=vg,
            data={
                "Vg": vg,
                "Id": [1e-6] * len(vg),
                "Ig": [1e-8] * len(vg),
            },
        )
        result = analyze_transfer_sweep(
            segment,
            {
                "drain_current_raw_column": "Id",
                "gate_leakage_column": "Ig",
            },
            {"transfer": {}},
            {
                "channel_width_um": 100,
                "channel_length_um": 10,
                "oxide_thickness_nm": 90,
                "dielectric_constant": 3.9,
            },
        )
        self.assertAlmostEqual(result.current_summary["leakage_ratio_max"], 0.01)


class PipelineRegressionTests(unittest.TestCase):
    def test_ss_fit_window_is_bounded_in_log_current_space(self):
        from fet_analyzer.plotting.transfer_plots import _ss_fit_window_bounds

        bounds = _ss_fit_window_bounds(
            np.array([-1.0, 0.0, 1.0]),
            np.array([1e-10, 1e-9, 1e-8]),
        )
        self.assertIsNotNone(bounds)
        x_min, x_max, y_min, y_max = bounds
        self.assertEqual((x_min, x_max), (-1.0, 1.0))
        self.assertLess(y_min, 1e-10)
        self.assertGreater(y_max, 1e-8)
        self.assertGreater(y_min, 1e-11)
        self.assertLess(y_max, 1e-7)

    def test_transfer_plot_handles_empty_hysteresis_noise_window(self):
        from fet_analyzer.plotting.transfer_plots import plot_transfer_curves

        vg = [-2, -1, 0, 1, 2]
        segments = [
            SweepSegment(
                direction=direction,
                bias_level=0.1,
                bias_variable="Vd",
                sweep_variable="Vg",
                sweep_values=values,
                data={"Vg": values, "Id": [1e-15] * len(values)},
            )
            for direction, values in (("forward", vg), ("reverse", list(reversed(vg))))
        ]
        classification = {
            "type": MeasurementType.TRANSFER,
            "sweep_variable": "Vg",
            "bias_variable": "Vd",
            "drain_current_column": "Id",
            "drain_current_raw_column": "Id",
            "source_filename": "IdVg__S_Device.csv",
            "filename_info": {"sample_label": "S", "measurement_type": "IdVg"},
        }
        with tempfile.TemporaryDirectory() as folder:
            metrics = plot_transfer_curves(
                segments,
                classification,
                Path(folder),
                config={"transfer": {"normalize_by_width": False, "noise_floor_a": 1e-13}},
                device_params={"channel_width_um": 100, "channel_length_um": 10},
            )
        self.assertIn("segments", metrics)

    def test_device_excel_uses_shared_gm_function(self):
        from fet_analyzer.reports.device_report import write_device_excel

        vg = [-2, -1, 0, 1, 2]
        segment = SweepSegment(
            direction="forward",
            bias_level=0.1,
            bias_variable="Vd",
            sweep_variable="Vg",
            sweep_values=vg,
            data={"Vg": vg, "Id": [10 ** (value / 2 - 8) for value in vg]},
        )
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "analysis.xlsx"
            write_device_excel(
                [segment],
                {
                    "type": MeasurementType.TRANSFER,
                    "sweep_variable": "Vg",
                    "bias_variable": "Vd",
                    "drain_current_raw_column": "Id",
                },
                {
                    "segments": {"0.1": {"ion_a": 1e-6, "ioff_a": 1e-9}},
                    "summary": {},
                    "analysis_settings": {
                        "ion_method_used": "maximum_measured_abs_id",
                        "ion_method_configured": "max_common_overdrive",
                        "ioff_leakage_factor": 3.0,
                        "vth_method": "peak_gm_tangent",
                    },
                },
                output,
            )
            self.assertTrue(output.exists())
            import openpyxl
            workbook = openpyxl.load_workbook(output, read_only=True)
            self.assertIn("Analysis Settings", workbook.sheetnames)
            self.assertEqual(workbook.sheetnames[:2], ["Summary", "Charts"])
            for redundant in ("Intermediate Segments", "Intermediate GM", "Intermediate SS",
                              "Intermediate Summary", "Per-Bias Metrics", "Origin Import"):
                self.assertNotIn(redundant, workbook.sheetnames)
            self.assertEqual(workbook["Audit Metrics"].sheet_state, "hidden")
            rows = list(workbook["Audit Metrics"].iter_rows(values_only=True))
            ion_row = next(row for row in rows if row[0] == "segments.0.1.ion_a")
            self.assertIn("maximum measured |Id|", ion_row[1])
            self.assertIn("largest gate overdrive", ion_row[1])
            workbook.close()


class SampleDatabaseTests(unittest.TestCase):
    def test_discovers_metric_against_varying_length(self):
        records = []
        for length, vth in ((5.0, -1.0), (10.0, -0.5), (25.0, 0.1)):
            records.append({
                "source_file": Path(f"IdVg__S_TLM1_{length:g}um.csv"),
                "classification": {
                    "type": MeasurementType.TRANSFER,
                    "filename_info": {"sample_label": "S"},
                },
                "device_params": {
                    "channel_length_um": length,
                    "channel_width_um": 100.0,
                },
                "metrics": {"summary": {"vth_v_mean": vth}},
            })
        comparisons = discover_comparisons(build_observations(records))
        match = [
            item for item in comparisons
            if item["sample_id"] == "S"
            and item["parameter"] == "channel_length_um"
            and item["metric"] == "summary.vth_v_mean"
        ]
        self.assertEqual(len(match), 1)
        self.assertEqual(match[0]["n_points"], 3)


class RunManifestTests(unittest.TestCase):
    def test_manifest_records_stage_completion_and_embedded_errors(self):
        from fet_analyzer.run_manifest import mark_manifest_index_complete, write_run_manifest
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "measurement.csv"
            source.write_text("Vg,Id\n0,0\n", encoding="utf-8")
            config = deep_merge(DEFAULT_CONFIG, {"general": {"input_folder": str(root)}})
            errors = [{"stage": "batch_summary", "message": "failed"}]
            path = write_run_manifest(
                root, [source], config,
                {"transfer": 1, "output": 0, "general_iv": 0, "tlm": 0, "error": 1},
                [], errors,
            )
            payload = json.loads(path.read_text(encoding="utf-8"))
            self.assertEqual(payload["schema_version"], "2.0")
            self.assertEqual(payload["processed_measurements"], 1)
            self.assertFalse(payload["completion"]["successful"])
            self.assertFalse(payload["completion"]["html_index"])
            self.assertEqual(payload["errors"], errors)
            mark_manifest_index_complete(path)
            self.assertTrue(json.loads(path.read_text(encoding="utf-8"))["completion"]["html_index"])

    def test_unavailable_cloud_file_does_not_abort_manifest(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "cloud-placeholder.ztr"
            path.write_bytes(b"placeholder")
            with patch(
                "fet_analyzer.run_manifest.file_sha256",
                side_effect=OSError(22, "Invalid argument"),
            ):
                entry = input_manifest_entry(path)
        self.assertEqual(entry["hash_status"], "unavailable")
        self.assertIsNone(entry["sha256"])
        self.assertIn("Invalid argument", entry["hash_error"])


class FileDiscoveryTests(unittest.TestCase):
    def test_hidden_directories_are_not_measurement_sources(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "device.csv").write_text("Vg,Id\n0,0\n", encoding="utf-8")
            hidden = root / ".tool-cache"
            hidden.mkdir()
            (hidden / "fixture.csv").write_text("x,y\n1,2\n", encoding="utf-8")
            files, _ = discover_files(root, True, [], ["csv", "xlsx", "xtr", "ztr"])
        self.assertEqual([path.name for path in files], ["device.csv"])

    def test_tool_build_directories_are_not_measurement_sources(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "device.csv").write_text("Vg,Id\n0,0\n", encoding="utf-8")
            for name in ("build", "dist", "node_modules", "__pycache__"):
                generated = root / name
                generated.mkdir()
                (generated / "fixture.csv").write_text("x,y\n1,2\n", encoding="utf-8")
            files, _ = discover_files(root, True, [], ["csv", "xlsx", "xtr", "ztr"])
        self.assertEqual([path.name for path in files], ["device.csv"])


class TLMTests(unittest.TestCase):
    def test_index_prefers_accepted_ungated_then_accepted_near_zero_gated(self):
        base = {
            "rsh_ohm_sq": 1000.0, "rc_ohm": 10.0, "rcw_ohm_um": 1000.0,
            "lt_um": 1.0, "rho_film_ohm_cm": None, "rhoc_ohm_cm2": 1e-6,
            "acceptance_reasons": [],
        }
        ungated = [{**base, "r2": 0.95, "acceptance_status": "accepted", "gate_bias": "ungated"}]
        gated = [{**base, "r2": 0.99, "acceptance_status": "accepted", "vg_v": 0.2, "vds_v": -0.1}]
        selected = _index_fields(ungated, gated, DEFAULT_CONFIG)
        self.assertEqual(selected["index_source"], "ungated_ltlm")
        self.assertIsNone(selected["index_actual_vg_v"])
        self.assertIsNone(selected["index_rho_film_ohm_cm"])

        ungated[0]["acceptance_status"] = "review"
        selected = _index_fields(ungated, gated, DEFAULT_CONFIG)
        self.assertEqual(selected["index_source"], "gated_transfer_vg0")
        self.assertEqual(selected["index_actual_vg_v"], 0.2)

        gated[0]["acceptance_status"] = "review"
        selected = _index_fields(ungated, gated, DEFAULT_CONFIG)
        self.assertEqual(selected["index_source"], "unavailable")

    def test_summary_keeps_ungated_and_output_sanity_separate(self):
        base = {
            "rsh_ohm_sq": 1000.0, "rc_ohm": 10.0, "rcw_ohm_um": 1000.0,
            "lt_um": 1.0, "rho_film_ohm_cm": 2e-3, "rhoc_ohm_cm2": 1e-6,
            "r2": 0.99, "acceptance_status": "accepted", "acceptance_reasons": [],
            "warnings": [],
        }
        ungated = [{**base, "gate_bias": "ungated", "rsh_ohm_sq": 1100.0}]
        output = [{**base, "gate_bias": 2.0, "rsh_ohm_sq": 2200.0}]
        records = [
            (5.0, {
                "source_file": Path("LTLM__S_TLM1_5um.csv"),
                "device_params": {"channel_width_um": 100.0},
                "classification": {"tlm_role": "ungated_ltlm"},
            }),
            (5.0, {
                "source_file": Path("IdVd__S_TLM1_5um.csv"),
                "device_params": {"channel_width_um": 100.0},
                "classification": {"tlm_role": "output_sanity"},
            }),
        ]
        row = _summary_row(
            "S", "TLM1", records, DEFAULT_CONFIG,
            ungated_fits=ungated, output_fits=output, transfer_fits=[], length_rows=[],
        )
        self.assertEqual(row["ungated_rsh_ohm_sq"], 1100.0)
        self.assertEqual(row["output_rsh_ohm_sq"], 2200.0)
        self.assertEqual(row["index_source"], "ungated_ltlm")
        self.assertEqual(row["index_rsh_ohm_sq"], 1100.0)
        self.assertNotIn("IdVd", row["index_source_files"])

    def test_summary_reports_zero_fixed_field_and_max_fittable_tlm_points(self):
        def fit(vg, vds, rcw, r2=0.99):
            return {
                "vg_v": vg, "vds_v": vds, "rcw_ohm_um": rcw,
                "rc_ohm": rcw / 200.0, "rsh_ohm_sq": rcw * 2,
                "lt_um": 1.0, "rhoc_ohm_cm2": 1e-6,
                "rho_film_ohm_cm": 2e-3, "r2": r2,
                "acceptance_status": "accepted", "acceptance_reasons": [],
            }
        fits = [fit(-5.0, -0.1, 5000), fit(-4.0, -0.1, 6000), fit(0.2, -0.1, 9000)]
        record = {
            "device_params": {
                "channel_width_um": 100.0, "oxide_thickness_nm": 20.0,
                "film_thickness_nm": 20.0,
            },
            "classification": {"sweep_variable": "Vg"},
        }
        config = deep_merge(DEFAULT_CONFIG, {"tlm": {
            "transfer_read_vg_v": -3.0,
            "transfer_gate_field_mv_cm": -2.0,
        }})
        row = _summary_row(
            "S", "TLM1", [(5.0, record)], config,
            output_fits=[], transfer_fits=fits, length_rows=[],
        )
        self.assertEqual(row["transfer_zero_vg_v"], 0.2)
        self.assertEqual(row["transfer_fixed_source"], "electric_field")
        self.assertEqual(row["transfer_fixed_requested_field_mv_cm"], -2.0)
        self.assertEqual(row["transfer_fixed_requested_vg_v"], -4.0)
        self.assertEqual(row["transfer_fixed_vg_v"], -4.0)
        self.assertEqual(row["transfer_fixed_vg_delta_v"], 0.0)
        self.assertEqual(row["transfer_fixed_actual_field_mv_cm"], -2.0)
        self.assertEqual(row["transfer_fixed_rcw_ohm_um"], 6000)
        self.assertEqual(row["transfer_max_vg_v"], -5.0)
        self.assertEqual(row["transfer_max_rcw_ohm_um"], 5000)
        self.assertIn("largest |measured/common Vg|", row["transfer_max_selection_reason"])

    def test_tlm_operating_point_ties_use_preferred_vds(self):
        fits = [
            {"vg_v": -0.2, "vds_v": -1.0, "r2": 0.99, "acceptance_status": "accepted"},
            {"vg_v": 0.2, "vds_v": -0.1, "r2": 0.80, "acceptance_status": "review"},
        ]
        selected = _near_zero_transfer_fit(fits, {"summary": {"preferred_vd_v": 0.1}})
        self.assertEqual(selected["vds_v"], -0.1)
        maximum = _maximum_fittable_transfer_fit(fits, {"summary": {"preferred_vd_v": 0.1}})
        self.assertEqual(maximum["vg_v"], 0.2)

    def test_ltlm_is_output_analysis_with_ungated_tlm_role(self):
        classification = classify_measurement(
            {"measurement_type": "LTLM"},
            ["Vd", "Vs", "Id"],
            {"Vd": [-0.1, 0.0, 0.1], "Vs": [0.0, 0.0, 0.0], "Id": [-1e-3, 0.0, 1e-3]},
            filename="LTLM__S_TLM1_10um__1.ztr",
            filename_patterns=DEFAULT_CONFIG["filename_patterns"],
        )
        self.assertEqual(classification["type"], MeasurementType.OUTPUT)
        self.assertTrue(classification["is_tlm"])
        self.assertTrue(classification["is_ungated_tlm"])
        self.assertEqual(classification["tlm_analysis_mode"], "ungated")
        self.assertEqual(classification["tlm_role"], "ungated_ltlm")

    def test_ltlm_filename_is_authoritative_even_with_incidental_gate_column(self):
        classification = classify_measurement(
            {"measurement_type": "Output"},
            ["Vd", "Vg", "Id"],
            {"Vd": [-0.1, 0.0, 0.1], "Vg": [0.0, 0.0, 0.0], "Id": [-1e-3, 0.0, 1e-3]},
            filename="LTLM__S_TLM1_10um__1.csv",
            filename_patterns=DEFAULT_CONFIG["filename_patterns"],
        )
        self.assertTrue(classification["is_explicit_ltlm"])
        self.assertTrue(classification["is_ungated_tlm"])
        self.assertEqual(classification["tlm_role"], "ungated_ltlm")

    def test_near_ioff_walks_toward_on_and_stops_at_first_all_length_accepted_fit(self):
        records = []
        for length, ioff_vg in ((5.0, -1.0), (10.0, -2.0), (20.0, -3.0)):
            records.append((length, {
                "source_file": Path(f"IdVg__S_TLM1_{length:g}um.csv"),
                "device_params": {"polarity": "p", "channel_width_um": 100.0},
                "metrics": {"sweep_results": [{
                    "identity": {"direction": "forward", "measured_vds_v": -0.1},
                    "subthreshold_swing": {"ss_ioff_vg": ioff_vg},
                    "vth": {"vth_v": -0.5},
                }]},
            }))
        base = {
            "vds_v": -0.1, "n_lengths": 3, "expected_lengths": 3, "excluded": 0,
            "slope_ohm_per_um": 10.0, "intercept_ohm": 100.0, "r2": 0.99,
            "points": [(5.0, 1.5), (10.0, 2.0), (20.0, 3.0)],
        }
        fits = [
            {**base, "vg_v": -3.0, "acceptance_status": "review"},
            {**base, "vg_v": -4.0, "acceptance_status": "accepted"},
            {**base, "vg_v": -5.0, "acceptance_status": "accepted"},
        ]
        selected = next(
            row for row in _gated_operating_points(records, fits, deep_merge(DEFAULT_CONFIG, {
                "transfer": {"ion_method": "maximum_measured"},
            })) if row["condition"] == "near_ioff_first_accepted"
        )
        self.assertEqual(selected["requested_vg_v"], -3.0)
        self.assertEqual(selected["fit"]["vg_v"], -4.0)
        self.assertIn("Vg=-3: review", selected["attempted_candidates"])
        self.assertIn("Vg=-4: accepted", selected["attempted_candidates"])

    def test_separate_master_tlm_workbook_has_mode_sheets_and_ungated_fit(self):
        records = []
        for length in (5.0, 10.0, 20.0):
            resistance = 1000.0 + 100.0 * length
            segment = SweepSegment(
                direction="forward", bias_level=None, bias_variable=None,
                sweep_variable="Vd", sweep_values=[-0.1, 0.0, 0.1],
                data={"Vd": [-0.1, 0.0, 0.1], "Id": [-0.1 / resistance, 0.0, 0.1 / resistance]},
            )
            records.append((length, {
                "source_file": Path(f"LTLM__S_TLM1_{length:g}um.csv"),
                "classification": {
                    "sweep_variable": "Vd", "drain_current_raw_column": "Id",
                    "is_tlm": True, "is_ungated_tlm": True, "tlm_role": "ungated_ltlm",
                },
                "segments": [segment],
                "device_params": {"channel_width_um": 100.0, "film_thickness_nm": 20.0, "polarity": "p"},
            }))
        with tempfile.TemporaryDirectory() as folder:
            path = _write_extended_master_tlm(Path(folder), {("S", "TLM1"): records}, DEFAULT_CONFIG)
            self.assertIsNotNone(path)
            import openpyxl
            workbook = openpyxl.load_workbook(path, read_only=True)
            try:
                self.assertEqual(
                    workbook.sheetnames,
                    ["README", "Gated TLM", "Ungated LTLM", "Output Sanity Check", "Fit Audit"],
                )
                rows = list(workbook["Ungated LTLM"].iter_rows(values_only=True))
                self.assertEqual(rows[1][2], "ungated_ltlm")
                self.assertIsNotNone(rows[1][22])  # film resistivity
                self.assertIsNotNone(rows[1][28])  # R2
            finally:
                workbook.close()

    def test_tlm_continuation_replaces_only_tlm_artifacts(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            output = root / "output"
            old_tlm = output / "TLM"
            batch = output / "batch_summary"
            old_tlm.mkdir(parents=True)
            batch.mkdir()
            (old_tlm / "old.txt").write_text("old", encoding="utf-8")
            sentinel = batch / "master_summary.csv"
            sentinel.write_text("device,value\nA,1\n", encoding="utf-8")
            before = sentinel.read_bytes()
            files = []
            for length in (5, 10, 20, 40, 80):
                resistance = 1000 + 100 * length
                path = root / f"LTLM__S_TLM1_{length}um__1.csv"
                path.write_text(
                    "Vd,Vs,Id\n" + "\n".join(
                        f"{vd},0,{vd / resistance}" for vd in (-0.1, -0.05, 0.0, 0.05, 0.1)
                    ) + "\n",
                    encoding="utf-8",
                )
                files.append(path)
            manifest = run_tlm_continuation(files, root, output, DEFAULT_CONFIG)
            self.assertTrue(manifest.exists())
            self.assertTrue((output / "TLM" / "master_tlm_summary.xlsx").exists())
            self.assertFalse((output / "TLM" / "old.txt").exists())
            self.assertEqual(sentinel.read_bytes(), before)
            import openpyxl
            individual_path = next((output / "TLM").glob("*/TLM1/TLM_*_TLM1.xlsx"))
            individual = openpyxl.load_workbook(individual_path, read_only=True)
            summary = openpyxl.load_workbook(output / "TLM" / "tlm_master_summary.xlsx", read_only=True)
            try:
                self.assertIn("Ungated LTLM Fits", individual.sheetnames)
                self.assertEqual(individual.sheetnames[-1], "Output Sanity Check")
                headers = [cell.value for cell in summary["TLM Summary"][1]]
                values = [cell.value for cell in summary["TLM Summary"][2]]
                row = dict(zip(headers, values))
                self.assertEqual(row["ungated_fit_count"], 1)
                self.assertEqual(row["output_fit_count"], 0)
                self.assertEqual(row["index_source"], "ungated_ltlm")
                self.assertEqual(row["index_acceptance_status"], "accepted")
            finally:
                individual.close()
                summary.close()
            report_text = individual_path.with_name(f"{individual_path.stem}_report.html").read_text(encoding="utf-8")
            self.assertIn("Ungated LTLM", report_text)
            self.assertNotIn("output TLM", report_text)

    def test_tlm_continuation_failure_preserves_existing_tlm(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            output = root / "output"
            old_tlm = output / "TLM"
            old_tlm.mkdir(parents=True)
            sentinel = old_tlm / "old.txt"
            sentinel.write_text("preserve", encoding="utf-8")
            source = root / "LTLM__S_TLM1_5um__1.csv"
            source.write_text("Vd,Vs,Id\n-0.1,0,-0.001\n0,0,0\n0.1,0,0.001\n", encoding="utf-8")
            with patch("fet_analyzer.analysis.tlm_workflow.generate_tlm_workflow", side_effect=RuntimeError("boom")):
                with self.assertRaises(RuntimeError):
                    run_tlm_continuation([source], root, output, DEFAULT_CONFIG)
            self.assertEqual(sentinel.read_text(encoding="utf-8"), "preserve")

    def test_ungated_ltlm_uses_sole_reverse_labelled_iv_sweep(self):
        segment = SweepSegment(
            direction="reverse", bias_level=None, bias_variable=None,
            sweep_variable="Vd", sweep_values=[-0.1, 0.0, 0.1],
            data={"Vd": [-0.1, 0.0, 0.1], "Id": [-1e-3, 0.0, 1e-3]},
        )
        classification = {
            "sweep_variable": "Vd", "drain_current_raw_column": "Id",
            "is_tlm": True,
        }
        extracted = extract_rtotal([segment], classification, {"tlm": {}})
        self.assertAlmostEqual(extracted["rtotal_ohm"], 100.0)
        record = {"classification": classification, "segments": [segment]}
        self.assertAlmostEqual(_output_resistance(record, 0.1)["floating_gate"], 100.0)

    def test_negative_intercept_is_retained_with_warning(self):
        result = extract_tlm([(5.0, 0.4), (10.0, 0.9), (20.0, 1.9)], width_um=100.0)
        self.assertEqual(result["status"], "rejected_nonphysical")
        self.assertIn("negative_intercept", result["acceptance_reasons"])
        self.assertLess(result["intercept_ohm"], 0)
        self.assertLess(result["rc_ohm"], 0)
        self.assertTrue(any("Negative TLM intercept" in warning for warning in result["warnings"]))

    def test_film_resistivity_uses_sheet_resistance_and_thickness(self):
        result = extract_tlm(
            [(5.0, 150.0), (10.0, 200.0), (20.0, 300.0)],
            width_um=100.0,
            film_thickness_nm=20.0,
        )
        self.assertAlmostEqual(result["rho_film_ohm_cm"], result["rsh_ohm_sq"] * 20e-7)

    def test_tlm_gate_field_uses_nearest_measured_vg_and_records_both(self):
        segment = SweepSegment(
            direction="forward", bias_level=-0.1, bias_variable="Vd",
            sweep_variable="Vg", sweep_values=[-4.0, -2.0, 0.0],
            data={"Vg": [-4.0, -2.0, 0.0], "Id": [-2e-3, -1e-3, -1e-6]},
        )
        result = extract_rtotal(
            [segment], {"sweep_variable": "Vg", "drain_current_raw_column": "Id"},
            {"tlm": {"transfer_read_mode": "constant_gate_field", "transfer_gate_field_mv_cm": -0.9}},
            {"oxide_thickness_nm": 20.0},
        )
        point = next(iter(result["rtotal_per_segment"].values()))
        self.assertAlmostEqual(point["requested_vg_v"], -1.8)
        self.assertEqual(point["actual_vg_v"], -2.0)
        self.assertAlmostEqual(point["actual_gate_field_mv_cm"], -1.0)

    def test_tlm_maximum_current_records_actual_measured_vg(self):
        segment = SweepSegment(
            direction="forward", bias_level=-0.1, bias_variable="Vd",
            sweep_variable="Vg", sweep_values=[-4.0, -2.0, 0.0],
            data={"Vg": [-4.0, -2.0, 0.0], "Id": [-2e-3, -3e-3, -1e-6]},
        )
        result = extract_rtotal(
            [segment], {"sweep_variable": "Vg", "drain_current_raw_column": "Id"},
            {"tlm": {"transfer_read_mode": "maximum_current"}},
        )
        point = next(iter(result["rtotal_per_segment"].values()))
        self.assertEqual(point["actual_vg_v"], -2.0)
        self.assertEqual(point["requested_vg_v"], -2.0)
        self.assertIn("maximum_current", point["method"])
        self.assertTrue(any("may differ" in warning for warning in result["warnings"]))

    def test_master_tlm_aggregates_structures_at_each_length(self):
        records = []
        for tlm_id, scale in (("TLM1", 1.0), ("TLM2", 1.08)):
            for length in (5.0, 10.0, 20.0):
                resistances = [100.0 + 10.0 * length, 120.0 + 12.0 * length, 150.0 + 15.0 * length]
                segment = SweepSegment(
                    direction="forward", bias_level=-0.1, bias_variable="Vd", sweep_variable="Vg",
                    sweep_values=[-2.0, -1.0, 0.0],
                    data={"Vg": [-2.0, -1.0, 0.0], "Id": [-0.1 / (value * scale) for value in resistances]},
                )
                records.append((tlm_id, length, {
                    "source_file": Path(f"IdVg__S_{tlm_id}_{length:g}um.csv"),
                    "classification": {"sweep_variable": "Vg", "drain_current_raw_column": "Id"},
                    "segments": [segment],
                    "device_params": {"channel_width_um": 100.0, "film_thickness_nm": 20.0},
                }))
        fits = _aggregate_transfer_rows(records, {}, 100.0, 20.0)
        self.assertEqual(len(fits), 3)
        fit = fits[0]
        self.assertEqual(fit["n_tlm_structures"], 2)
        self.assertEqual(fit["n_lengths"], 3)
        self.assertTrue(all(item["n_observations"] == 2 for item in fit["length_stats"]))
        self.assertTrue(all(item["id_std_a"] > 0 for item in fit["length_stats"]))
        self.assertAlmostEqual(fit["rho_film_ohm_cm"], fit["rsh_ohm_sq"] * 20e-7)


class TemporaryWorkspaceTests(unittest.TestCase):
    def test_run_workspace_contains_and_removes_all_temporary_directories(self):
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder)
            workspace = RunWorkspace(output)
            device = workspace.temporary_directory("device_")
            tlm = temporary_directory(
                {"execution": {"temporary_root": str(workspace.path)}},
                "tlm_", output / "TLM",
            )
            self.assertEqual(device.parent, workspace.path)
            self.assertEqual(tlm.parent, workspace.path)
            self.assertTrue(output.joinpath(".fw").exists())
            workspace.cleanup()
            self.assertFalse(output.joinpath(".fw").exists())

    def test_new_workspace_removes_only_stale_run_directories(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder) / ".fet_work"
            stale = root / "run_111_stale"
            active = root / "run_222_active"
            stale.mkdir(parents=True)
            active.mkdir()
            with patch(
                "fet_analyzer.temp_workspace._process_is_running",
                side_effect=lambda pid: pid == 222,
            ):
                workspace = RunWorkspace(Path(folder))
                self.assertFalse(stale.exists())
                self.assertTrue(active.exists())
                workspace.cleanup()
                self.assertTrue(root.exists())


class HtmlIndexTests(unittest.TestCase):
    def test_index_links_reports_and_exposes_dashboard_and_master_tlm(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "batch_summary").mkdir()
            (root / "TLM" / "S" / "MASTER").mkdir(parents=True)
            (root / "errors").mkdir()
            (root / "batch_summary" / "master_summary.csv").write_text(
                "device,sample_id,type,ion_ioff_log10,ss_min_mv_dec,quality_status\nD1,S,transfer,6,80,pass\n",
                encoding="utf-8",
            )
            (root / "D1_report.html").write_text("<html></html>", encoding="utf-8")
            (root / "D1.xlsx").write_bytes(b"xlsx")
            (root / "TLM" / "S" / "MASTER" / "TLM_S_MASTER_report.html").write_text("<html></html>", encoding="utf-8")
            (root / "TLM" / "S" / "MASTER" / "TLM_S_MASTER.xlsx").write_bytes(b"xlsx")
            (root / "TLM" / "tlm_master_summary.csv").write_text(
                "sample_id,tlm_id,analysis_level,n_tlm_structures,index_source,index_actual_vg_v,index_rsh_ohm_sq,index_rcw_ohm_um,index_rho_film_ohm_cm,index_rhoc_ohm_cm2,index_acceptance_status\nS,MASTER,master_aggregate,2,ungated_ltlm,,1000,25000,0.002,1e-5,accepted\n",
                encoding="utf-8",
            )
            (root / "run_manifest.json").write_text('{"statistics":{"transfer":1}}', encoding="utf-8")
            (root / "errors" / "error_report.json").write_text('{"error_count":0}', encoding="utf-8")
            page = generate_html_index(root)
            rendered = page.read_text(encoding="utf-8")
            self.assertIn("Device dashboard", rendered)
            self.assertIn("TLM dashboard", rendered)
            self.assertIn("master_aggregate", rendered)
            self.assertIn("Contact resistance (kΩ·µm)", rendered)
            self.assertIn("<td>25</td>", rendered)
            self.assertIn("Ungated LTLM", rendered)
            self.assertIn("D1_report.html", rendered)


if __name__ == "__main__":
    unittest.main()
