from __future__ import annotations

import csv
import base64
import json
import tempfile
import unittest
import logging
import sys
from pathlib import Path
from unittest.mock import MagicMock, patch

from openpyxl import Workbook

from fet_analyzer.dashboard.adapters import ReportCatalog
from fet_analyzer.dashboard.results_api import resolve_artifact
from fet_analyzer.dashboard.server import DashboardState, create_server
from fet_analyzer.utils.logging import ConsoleSafeFormatter
from fet_analyzer.reports.error_report import error_entry, write_error_reports


class DashboardAdapterTests(unittest.TestCase):
    @staticmethod
    def _measurement(root: Path) -> None:
        (root / "IdVg__Sample_Detail_FET_Device__1.csv").write_text(
            "Vg,Vd,Id\n-1,0.1,1e-8\n0,0.1,1e-7\n1,0.1,1e-6\n",
            encoding="utf-8",
        )

    @staticmethod
    def _ready_config(state: DashboardState) -> None:
        state.config_payload(create=True)
        state.save_config({"transfer.ion_fixed_vg_v": 1.0})

    def test_console_formatter_is_ascii_safe(self):
        record = logging.LogRecord("test", logging.INFO, "", 0, "→ Classified — 10 μm ✓", (), None)
        rendered = ConsoleSafeFormatter("%(message)s").format(record)
        self.assertTrue(rendered.isascii())
        self.assertIn("-> Classified - 10 um OK", rendered)

    def test_catalog_reads_csv_excel_and_nested_json(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            with (root / "summary.csv").open("w", newline="", encoding="utf-8") as handle:
                csv.writer(handle).writerows([["device", "ion_a"], ["A", 1e-6]])
            workbook = Workbook()
            workbook.active.title = "Sweep Results"
            workbook.active.append(["vg_v", "id_a"])
            workbook.active.append([0, 1e-9])
            workbook.save(root / "analysis.xlsx")
            (root / "metrics.json").write_text(json.dumps({"summary": {"vth_v": 2.5}}), encoding="utf-8")

            catalog = ReportCatalog()
            refs = catalog.scan(root)
            self.assertEqual({ref.name for ref in refs}, {"summary", "Sweep Results", "metrics"})
            csv_ref = next(ref for ref in refs if ref.name == "summary")
            table = catalog.read(csv_ref)
            self.assertEqual(table["columns"], ["device", "ion_a"])
            self.assertEqual(table["rows"][0]["device"], "A")
            json_ref = next(ref for ref in refs if ref.name == "metrics")
            self.assertEqual(catalog.read(json_ref)["rows"][0]["field"], "summary.vth_v")

    def test_state_discovers_existing_pipeline_output(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            output = root / "output"
            output.mkdir()
            (output / "master_summary.csv").write_text("device,value\nA,1\n", encoding="utf-8")
            (output / "run_manifest.json").write_text(json.dumps({"status": "success", "statistics": {"transfer": 1}}), encoding="utf-8")
            state = DashboardState(root)
            self.assertTrue(state.datasets)
            self.assertEqual(state.summary()["manifest"]["status"], "success")

    def test_dashboard_progress_parsing_uses_utf8_safe_subprocess(self):
        class FakeProcess:
            stdout = iter(["[PROGRESS] 1/2 50% Processing first.ztr\n", "[PROGRESS] 2/2 100% Completed all outputs\n"])
            def poll(self): return 0
            def wait(self, timeout=None): return 0

        with tempfile.TemporaryDirectory() as tmp, patch(
            "fet_analyzer.dashboard.server.subprocess.Popen", return_value=FakeProcess()
        ) as popen:
            self._measurement(Path(tmp))
            state = DashboardState(Path(tmp))
            self._ready_config(state)
            state.start_run()
            for _ in range(100):
                if state.progress_current == 2:
                    break
                import time
                time.sleep(0.005)
            self.assertEqual((state.progress_current, state.progress_total), (2, 2))
            self.assertEqual(state.progress_label, "Analysis complete")
            self.assertTrue(state.progress_complete)
            self.assertEqual(popen.call_args.kwargs["encoding"], "utf-8")
            self.assertEqual(popen.call_args.kwargs["errors"], "replace")

    def test_stop_run_terminates_active_process(self):
        class ActiveProcess:
            terminated = False
            def poll(self): return None if not self.terminated else 1
            def terminate(self): self.terminated = True
            def wait(self, timeout=None): return 1

        with tempfile.TemporaryDirectory() as tmp:
            state = DashboardState(Path(tmp))
            process = ActiveProcess()
            state.process = process
            state.stop_run()
            self.assertTrue(process.terminated)
            self.assertTrue(state.cancelled)

    def test_dashboard_execution_controls_map_to_cli_flags(self):
        class FakeProcess:
            stdout = iter(())
            def poll(self): return 0
            def wait(self, timeout=None): return 0

        with tempfile.TemporaryDirectory() as tmp, patch(
            "fet_analyzer.dashboard.server.subprocess.Popen", return_value=FakeProcess()
        ) as popen:
            self._measurement(Path(tmp))
            state = DashboardState(Path(tmp))
            self._ready_config(state)
            state.start_run(device_excel=False, plot_copies=False, batch_plots=False, overwrite=True, workers=3)
            command = popen.call_args.args[0]
            self.assertIn("--no-device-excel", command)
            self.assertNotIn("--save-plots", command)
            self.assertNotIn("--batch-plots", command)
            self.assertNotIn("--analysis-mode", command)
            self.assertIn("--overwrite", command)
            self.assertEqual(command[command.index("--workers") + 1], "3")

    def test_dashboard_tlm_continuation_maps_to_scoped_cli_flags(self):
        class FakeProcess:
            stdout = iter(())
            def poll(self): return 0
            def wait(self, timeout=None): return 0

        with tempfile.TemporaryDirectory() as tmp, patch(
            "fet_analyzer.dashboard.server.subprocess.Popen", return_value=FakeProcess()
        ) as popen:
            state = DashboardState(Path(tmp))
            with patch.object(state, "tlm_preflight_payload", return_value={"status": "pass", "issues": []}):
                state.start_run(tlm_continue=True)
            command = popen.call_args.args[0]
            self.assertIn("--tlm-continue", command)
            self.assertEqual(command[command.index("--only") + 1], "tlm")

    def test_dashboard_reuses_completed_tlm_preflight_once(self):
        with tempfile.TemporaryDirectory() as tmp:
            state = DashboardState(Path(tmp))
            payload = {"status": "pass", "issues": [], "eligible_groups": []}
            with patch(
                "fet_analyzer.file_discovery.discover_files", return_value=([], []),
            ), patch(
                "fet_analyzer.tlm_continue.tlm_preflight", return_value=payload,
            ) as preflight:
                self.assertEqual(state.tlm_preflight_payload(), payload)
                self.assertEqual(state.tlm_preflight_payload(consume_cached=True), payload)
                self.assertEqual(preflight.call_count, 1)
                state.tlm_preflight_payload(consume_cached=True)
                self.assertEqual(preflight.call_count, 2)

    def test_dashboard_server_can_bind_an_ephemeral_loopback_port(self):
        with tempfile.TemporaryDirectory() as tmp:
            server, url = create_server(Path(tmp), port=0)
            try:
                self.assertTrue(url.startswith("http://127.0.0.1:"))
                self.assertNotEqual(server.server_address[1], 0)
            finally:
                server.server_close()
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaisesRegex(ValueError, "loopback"):
                create_server(Path(tmp), host="0.0.0.0", port=0)

    def test_native_host_reuses_dashboard_server_and_closes_cleanly(self):
        from fet_analyzer import native

        fake_server = MagicMock()
        fake_webview = MagicMock()
        fake_webview.FileDialog.FOLDER = 20
        fake_window = fake_webview.create_window.return_value
        with patch.dict(sys.modules, {"webview": fake_webview}), patch.object(
            native, "create_server", return_value=(fake_server, "http://127.0.0.1:54321")
        ), patch.object(native.threading, "Thread") as thread:
            self.assertEqual(native.main(["--root", "."]), 0)
        thread.return_value.start.assert_called_once()
        fake_webview.create_window.assert_called_once()
        fake_webview.start.assert_called_once_with(debug=False)
        picker = fake_server.dashboard_state.folder_picker
        fake_window.create_file_dialog.return_value = (r"D:\Measurements",)
        self.assertEqual(picker(Path("D:/Initial")), (r"D:\Measurements",))
        fake_window.create_file_dialog.assert_called_once_with(20, directory="D:\\Initial")
        fake_server.shutdown.assert_called_once()
        fake_server.server_close.assert_called_once()

    def test_dashboard_native_folder_picker_normalizes_pywebview_result(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            selected = root / "selected"
            selected.mkdir()
            state = DashboardState(root)
            state.folder_picker = lambda initial: (str(selected),)
            self.assertEqual(state.browse_for_folder(), str(selected))
            state.folder_picker = lambda initial: None
            self.assertIsNone(state.browse_for_folder())

    def test_project_switch_updates_paths_immediately_and_reindexes_new_results(self):
        import time
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            selected = root / "new project"
            output = selected / "output"
            output.mkdir(parents=True)
            (output / "new_summary.csv").write_text("device,value\nnew,2\n", encoding="utf-8")
            state = DashboardState(root)
            old_revision = state.project_revision
            state.switch_project(selected)
            immediate = state.summary()
            self.assertEqual(Path(immediate["input_dir"]).resolve(), selected.resolve())
            self.assertEqual(Path(immediate["output_dir"]).resolve(), output.resolve())
            self.assertGreater(immediate["project_revision"], old_revision)
            for _ in range(200):
                if not state.summary()["catalog_loading"]:
                    break
                time.sleep(0.005)
            self.assertFalse(state.summary()["catalog_loading"])
            self.assertTrue(any(ref.name == "new_summary" for ref in state.datasets.values()))

    def test_project_switch_accepts_independent_output_folder(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            selected = root / "measurements"
            custom_output = root / "results elsewhere"
            selected.mkdir()
            state = DashboardState(root)
            state.switch_project(selected, custom_output)
            self.assertEqual(state.input_dir.resolve(), selected.resolve())
            self.assertEqual(state.output_dir.resolve(), custom_output.resolve())

    def test_dashboard_disables_overwrite_explicitly_when_unchecked(self):
        class FakeProcess:
            stdout = iter(())
            def poll(self): return 0
            def wait(self, timeout=None): return 0

        with tempfile.TemporaryDirectory() as tmp, patch(
            "fet_analyzer.dashboard.server.subprocess.Popen", return_value=FakeProcess()
        ) as popen:
            self._measurement(Path(tmp))
            state = DashboardState(Path(tmp))
            self._ready_config(state)
            state.start_run(overwrite=False)
            self.assertIn("--no-overwrite", popen.call_args.args[0])

    def test_dashboard_creates_and_edits_project_local_config(self):
        with tempfile.TemporaryDirectory() as tmp:
            state = DashboardState(Path(tmp))
            missing = state.config_payload()
            self.assertFalse(missing["exists"])

            created = state.config_payload(create=True)
            self.assertTrue(created["exists"])
            self.assertEqual(Path(created["path"]).name, "fet_analyzer_config.yaml")
            self.assertTrue(any(item["id"] == "generic_nfet" for item in created["presets"]))
            fields = {
                field["path"]: field
                for section in created["sections"]
                for field in section["fields"]
            }
            self.assertEqual(fields["transfer.normalize_by_width"]["type"], "boolean")
            self.assertEqual(fields["transfer.noise_floor_a"]["type"], "number")
            self.assertEqual(fields["transfer.ss_vg_range_v"]["type"], "nullable")

            saved = state.save_config({
                "transfer.noise_floor_a": 2e-13,
                "transfer.normalize_by_width": False,
                "transfer.ss_vg_range_v": [-5, 5],
            })
            saved_fields = {
                field["path"]: field["value"]
                for section in saved["sections"]
                for field in section["fields"]
            }
            self.assertEqual(saved_fields["transfer.noise_floor_a"], 2e-13)
            self.assertFalse(saved_fields["transfer.normalize_by_width"])
            self.assertEqual(saved_fields["transfer.ss_vg_range_v"], [-5, 5])

    def test_dashboard_creates_and_edits_device_parameters(self):
        with tempfile.TemporaryDirectory() as tmp:
            state = DashboardState(Path(tmp))
            missing = state.device_parameters_payload()
            self.assertFalse(missing["exists"])

            created = state.device_parameters_payload(create=True)
            self.assertTrue(created["exists"])
            self.assertIn("film_thickness_nm", created["columns"])
            self.assertIn("contact_length_um", created["columns"])

            saved = state.save_device_parameters([
                {
                    "sample_label": "SampleA",
                    "device_pattern": "*TLM1*",
                    "polarity": "N",
                    "channel_width_um": "100.000",
                    "film_thickness_nm": 20,
                    "top_gated": "yes",
                    "notes": "confirmed geometry",
                }
            ])
            self.assertEqual(saved["rows"][0]["polarity"], "n")
            self.assertEqual(saved["rows"][0]["channel_width_um"], "100")
            self.assertEqual(saved["rows"][0]["film_thickness_nm"], "20")
            self.assertEqual(saved["rows"][0]["top_gated"], "true")
            self.assertIn("SampleA\t*TLM1*", state.device_parameters_path.read_text(encoding="utf-8"))

            layered = state.save_device_parameters([
                {"sample_label": "SampleA", "device_pattern": "*", "channel_width_um": 100},
                {"sample_label": "SampleA", "device_pattern": "*", "film_thickness_nm": 20},
            ])
            self.assertEqual(len(layered["rows"]), 2)

    def test_dashboard_rejects_invalid_device_parameters_without_overwriting(self):
        with tempfile.TemporaryDirectory() as tmp:
            state = DashboardState(Path(tmp))
            state.device_parameters_payload(create=True)
            original = state.device_parameters_path.read_text(encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "greater than zero"):
                state.save_device_parameters([
                    {"sample_label": "A", "device_pattern": "*", "film_thickness_nm": 0}
                ])
            self.assertEqual(state.device_parameters_path.read_text(encoding="utf-8"), original)
            with self.assertRaisesRegex(ValueError, "equally specific"):
                state.save_device_parameters([
                    {"sample_label": "A", "device_pattern": "*", "channel_width_um": 100},
                    {"sample_label": "A", "device_pattern": "*", "channel_width_um": 200},
                ])

    def test_ion_preflight_resolves_each_device_without_running_analysis(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "IdVg__S_Detail_FET_Device__1.csv").write_text(
                "Vg,Vd,Id\n-2,0.1,1e-6\n0,0.1,1e-8\n2,0.1,1e-9\n",
                encoding="utf-8",
            )
            state = DashboardState(root)
            state.config_payload(create=True)
            state.save_config({
                "transfer.ion_method": "fixed_overdrive_field",
                "transfer.ion_overdrive_field_mv_cm": -2.0,
            })
            payload = state.ion_preflight()
            self.assertEqual(len(payload["rows"]), 1)
            self.assertAlmostEqual(payload["rows"][0]["resolved_overdrive_v"], -18.0)
            self.assertEqual(payload["rows"][0]["status"], "Review")

    def test_ztr_preflight_uses_and_removes_system_temporary_cache(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "IdVg__S_Detail_FET_Device__1.ztr"
            source.write_bytes(b"placeholder")
            cache_paths: list[Path] = []

            def fake_decompress(_source, cache_dir):
                cached = Path(cache_dir) / "measurement.xtr"
                cached.write_text("temporary", encoding="utf-8")
                cache_paths.append(cached)
                return cached

            parsed = {
                "metadata": {}, "columns": ["Vg", "Vd", "Id"],
                "data": {"Vg": [-1, 0, 1], "Vd": [0.1] * 3, "Id": [1e-9, 1e-8, 1e-7]},
                "num_rows": 3,
            }
            with patch("fet_analyzer.parsers.ztr_parser.decompress_ztr_to_cache", side_effect=fake_decompress), patch(
                "fet_analyzer.registry.parse_measurement", return_value=parsed
            ):
                payload = DashboardState(root).ion_preflight(recursive=False)
            self.assertEqual(len(payload["rows"]), 1)
            self.assertTrue(cache_paths)
            self.assertFalse(cache_paths[0].parent.exists())
            self.assertFalse((root / "decompressed_ztr").exists())

    def test_run_preflight_does_not_parse_every_measurement(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            self._measurement(root)
            with patch("fet_analyzer.registry.parse_measurement") as parser:
                payload = DashboardState(root).preflight_payload(overwrite=True)
            parser.assert_not_called()
            self.assertEqual(payload["input_count"], 1)
            self.assertFalse(payload["measurement_inspection"])

    def test_config_and_device_enums_are_exposed_as_dropdown_choices(self):
        with tempfile.TemporaryDirectory() as tmp:
            state = DashboardState(Path(tmp))
            payload = state.config_payload(create=True)
            fields = {
                field["path"]: field
                for section in payload["sections"] for field in section["fields"]
            }
            self.assertIn("fixed_gate_field", fields["transfer.ion_method"]["options"])
            self.assertIn("maximum_current", fields["tlm.transfer_read_mode"]["options"])
            self.assertIn("constant_overdrive_field", fields["tlm.transfer_read_mode"]["options"])
            self.assertNotIn("gaussian", fields["transfer.smooth_method"]["options"])
            self.assertEqual(fields["transfer.vth_method"]["options"], ["peak_gm_tangent", "constant_current"])
            self.assertEqual(fields["transfer.ss_method"]["options"], ["minimum"])
            self.assertEqual(fields["tlm.tlm_ss_method"]["options"], ["minimum"])
            self.assertNotIn("tlm.tlm_ion_method", fields)
            self.assertEqual(fields["device_defaults.polarity"]["options"], ["p", "n"])
            parameters = state.device_parameters_payload(create=True)
            polarity = next(field for field in parameters["fields"] if field["name"] == "polarity")
            self.assertEqual(polarity["options"], ["p", "n"])

    def test_dashboard_automatically_uses_local_config(self):
        class FakeProcess:
            stdout = iter(())
            def poll(self): return 0
            def wait(self, timeout=None): return 0

        with tempfile.TemporaryDirectory() as tmp, patch(
            "fet_analyzer.dashboard.server.subprocess.Popen", return_value=FakeProcess()
        ) as popen:
            self._measurement(Path(tmp))
            state = DashboardState(Path(tmp))
            state.config_payload(create=True)
            state.save_config({"transfer.ion_fixed_vg_v": 1.0})
            state.start_run()
            command = popen.call_args.args[0]
            self.assertIn("--config", command)
            self.assertIn(str(state.local_config_path), command)

    def test_frozen_dashboard_resolves_primary_self_worker(self):
        from fet_analyzer import runtime
        with tempfile.TemporaryDirectory() as tmp, patch.object(runtime, "is_frozen", return_value=True), patch.object(
            runtime, "frozen_distribution_dir", return_value=Path(tmp)
        ):
            with self.assertRaisesRegex(FileNotFoundError, "executable is missing"):
                runtime.resolve_worker_command()
            primary = Path(tmp) / "FET-Analyzer-v2.exe"
            primary.touch()
            self.assertEqual(runtime.resolve_worker_command(), [str(primary), "--worker"])

    def test_worker_probe_reports_response_and_access_denial(self):
        from fet_analyzer import runtime
        with patch.object(runtime, "is_frozen", return_value=True), patch.object(
            runtime, "resolve_worker_command", return_value=[r"C:\FET Analyzer\FET-Analyzer-v2.exe", "--worker"]
        ), patch.object(runtime.subprocess, "run", return_value=runtime.subprocess.CompletedProcess(
            [], 0, stdout="worker-ready\n", stderr=""
        )):
            self.assertEqual(runtime.probe_worker()["status"], "responded")
        denied = PermissionError(13, "Access is denied")
        with patch.object(runtime, "is_frozen", return_value=True), patch.object(
            runtime, "resolve_worker_command", return_value=[r"C:\FET Analyzer\FET-Analyzer-v2.exe", "--worker"]
        ), patch.object(runtime.subprocess, "run", side_effect=denied):
            payload = runtime.probe_worker()
        self.assertEqual(payload["status"], "blocked")
        self.assertIn("Access is denied", payload["error"])
        self.assertIn("approved local folder", payload["error"])

    def test_preflight_reports_all_blockers_together(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            output = root / "output"
            output.mkdir()
            (output / "existing.txt").write_text("old", encoding="utf-8")
            payload = DashboardState(root).preflight_payload(overwrite=False)
            codes = {item["code"] for item in payload["issues"]}
            self.assertEqual(payload["status"], "error")
            self.assertIn("no_inputs", codes)
            self.assertIn("output_exists", codes)

    def test_comparison_payload_reads_canonical_summary(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "output"
            (output / "batch_summary").mkdir(parents=True)
            (output / "batch_summary" / "master_summary.csv").write_text(
                "sample_id,vth_v_mean\nA,1.2\n", encoding="utf-8"
            )
            payload = DashboardState.comparison_payload(Path(tmp))
            self.assertEqual(payload["rows"][0]["sample_id"], "A")
            self.assertEqual(payload["row_count"], 1)

    def test_batch_payload_returns_typed_summary_and_definitions(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            batch = root / "output" / "batch_summary"
            batch.mkdir(parents=True)
            (batch / "master_summary.csv").write_text(
                "sample_id,device,ion_ioff_log10\nA,D1,5.2\n", encoding="utf-8"
            )
            (batch / "metric_statistics.csv").write_text(
                "sample_id,metric,n_valid,median\nA,ion_ioff_log10,1,5.2\n", encoding="utf-8"
            )
            payload = DashboardState(root).batch_payload()
            self.assertTrue(payload["available"])
            self.assertEqual(payload["rows"][0]["ion_ioff_log10"], 5.2)
            self.assertTrue(any(item["key"] == "ion_configured_a" for item in payload["definitions"]))

    def test_results_payload_preserves_canonical_values_and_safe_artifacts(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            output = root / "output"
            batch = output / "batch_summary"
            batch.mkdir(parents=True)
            workbook = output / "device_a.xlsx"
            report = output / "device_a_report.html"
            Workbook().save(workbook)
            report.write_text("<html>device</html>", encoding="utf-8")
            (batch / "master_summary.csv").write_text(
                "sample_id,device,source_file,type,ion_configured_ua_per_um,quality_status\n"
                "Sample A,device_a,device_a.csv,transfer,12.5,accepted\n",
                encoding="utf-8",
            )
            (batch / "metric_statistics.csv").write_text(
                "sample_id,metric,n_files,n_valid,missing,mean,std_dev,ci95_half_width\n"
                "Sample A,ion_configured_ua_per_um,1,1,0,12.5,,\n",
                encoding="utf-8",
            )
            payload = DashboardState(root).results_payload()
            self.assertEqual(payload["rows"][0]["ion_configured_ua_per_um"], 12.5)
            self.assertEqual(resolve_artifact(output, payload["rows"][0]["report_id"]), report.resolve())
            self.assertEqual(resolve_artifact(output, payload["rows"][0]["workbook_id"]), workbook.resolve())
            outside = root / "outside.html"
            outside.write_text("outside", encoding="utf-8")
            identifier = base64.urlsafe_b64encode(b"../outside.html").decode("ascii").rstrip("=")
            with self.assertRaisesRegex(ValueError, "outside"):
                resolve_artifact(output, identifier)

    def test_tlm_results_payload_reads_points_fits_and_joins_device_metrics(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            output = root / "output"
            batch = output / "batch_summary"
            tlm_dir = output / "TLM" / "Sample A" / "TLM1"
            batch.mkdir(parents=True)
            tlm_dir.mkdir(parents=True)
            (batch / "master_summary.csv").write_text(
                "sample_id,device,source_file,dibl_mv_v,quality_status\n"
                "Sample A,device_a,device_a.csv,18.0,accepted\n",
                encoding="utf-8",
            )
            workbook = Workbook()
            metrics = workbook.active
            metrics.title = "Metrics vs Lch"
            metrics.append(["source_file", "lch_um", "mobility_cm2_vs"])
            metrics.append(["device_a.csv", 2.0, 4.5])
            fits = workbook.create_sheet("TLM Parameters vs Vg")
            fits.append(["vg_v", "rcw_ohm_um", "rho_film_ohm_cm", "r2"])
            fits.append([1.0, 2500.0, 3.2, 0.99])
            diagnostics = workbook.create_sheet("Transfer Fit Diagnostics")
            diagnostics.append(["vg_v", "lch_um", "measured_rtotal_w_ohm_um", "predicted_rtotal_w_ohm_um", "residual_ohm_um"])
            diagnostics.append([1.0, 2.0, 5000.0, 4950.0, 50.0])
            path = tlm_dir / "Sample_A_TLM1.xlsx"
            workbook.save(path)
            payload = DashboardState(root).tlm_results_payload()
            structure = payload["structures"][0]
            self.assertEqual(structure["metrics_vs_lch"][0]["dibl_mv_v"], 18.0)
            self.assertEqual(structure["fits"][0]["rcw_kohm_um"], 2.5)
            self.assertEqual(structure["fits"][0]["mode"], "transfer")
            self.assertEqual(structure["diagnostics"][0]["residual_ohm_um"], 50.0)

    def test_results_ui_uses_local_plotly_and_has_no_presentation_feature(self):
        static = Path(__file__).parents[1] / "fet_analyzer" / "dashboard" / "static"
        html = (static / "index.html").read_text(encoding="utf-8")
        analytics = (static / "analytics.js").read_text(encoding="utf-8")
        self.assertIn('src="plotly-3.7.0.min.js"', html)
        self.assertIn("Sample overview", html)
        self.assertIn("TLM explorer", html)
        self.assertIn("applyProjectChange(status)", html)
        self.assertIn("await refresh()", html)
        self.assertNotIn('id="outputPath" readonly', html)
        self.assertIn('id="browseOutputBtn"', html)
        self.assertIn("payload.output_path=$('outputPath').value", html)
        self.assertIn("Run TLM only - continue", html)
        self.assertIn("Checking TLM inputs", html)
        self.assertIn("Confirm TLM refresh", html)
        self.assertIn("pendingTlmOptions", html)
        self.assertIn("/api/tlm-preflight", html)
        self.assertIn("/api/run-tlm", html)
        self.assertIn(".notice[hidden]{display:none}", html)
        self.assertIn(".notice:not([hidden]){display:block}", html)
        self.assertNotIn("margin-bottom:12px;display:none}.gallery", html)
        self.assertNotIn("presentation", (html + analytics).lower())
        self.assertNotIn("sooraj.sanjay@gmail.com", html.lower())
        self.assertNotIn("mailto:", html.lower())
        self.assertTrue((static / "plotly-3.7.0.min.js").stat().st_size > 1_000_000)
        self.assertTrue((static / "PLOTLY-LICENSE.txt").is_file())

    def test_error_report_writes_clean_and_failure_artifacts(self):
        with tempfile.TemporaryDirectory() as tmp:
            try:
                raise ValueError("bad measurement")
            except ValueError as exc:
                errors = [error_entry("file_processing", exc, "device.ztr")]
            paths = write_error_reports(Path(tmp), errors)
            payload = json.loads(paths["json"].read_text(encoding="utf-8"))
            self.assertEqual(payload["error_count"], 1)
            self.assertIn("bad measurement", paths["html"].read_text(encoding="utf-8"))
            self.assertFalse((Path(tmp) / "errors" / "error_report.md").exists())
            self.assertTrue(paths["csv"].exists())


if __name__ == "__main__":
    unittest.main()
